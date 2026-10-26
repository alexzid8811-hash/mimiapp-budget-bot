#!/usr/bin/env python3
"""Пересоздаёт все скриншоты инструкции docs/guide/images.

Запуск из корня репозитория:

    python scripts/make_guide_screenshots.py

Скрипт поднимает приложение локально на временной базе (DEV_MODE, без
Telegram), наполняет её демо-данными через API приложения, снимает экраны
в Chromium (390×844, светлая тема, device_scale_factor 2) и рисует на них
пронумерованные отметки.  Дата в приложении фиксирована: на сервере
подменяются app.clock.today/now (как в тестах), в браузере — page.clock.

Кроме картинок скрипт пишет docs/guide/generated/numbers.json (числа,
которые видны на скриншотах) и docs/guide/generated/bot-messages.md
(сообщения бота с теми же демо-данными).  По ним сверяют текст глав.

Подробности — docs/guide/HOWTO-UPDATE.md.
"""
from __future__ import annotations

import argparse
import json
import os
import random
import shutil
import socket
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.request
from datetime import date, datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

ROOT = Path(__file__).resolve().parents[1]
IMAGES = ROOT / "docs" / "guide" / "images"
GENERATED = ROOT / "docs" / "guide" / "generated"
TZ = "Europe/Moscow"
# Один цвет отметок на всю инструкцию: такого нет ни в интерфейсе, ни в цветах категорий.
MARK = "#E0009A"
VIEWPORT = {"width": 390, "height": 844}

# Заглушка Telegram Mini App SDK: приложение работает в DEV_MODE без initData.
TELEGRAM_STUB = """
window.Telegram = {WebApp: {
  initData: '', initDataUnsafe: {}, colorScheme: 'light', themeParams: {},
  platform: 'ios', version: '8.0', isFullscreen: false,
  safeAreaInset: {top: 0}, contentSafeAreaInset: {top: 0},
  ready() {}, expand() {}, onEvent() {}, offEvent() {},
}};
"""


# ---------------------------------------------------------------------------
# Режимы, которые запускаются в отдельном процессе с подменённой датой
# ---------------------------------------------------------------------------

CLOCK_OVERRIDE: dict = {}


def patch_clock(date_file: Path) -> None:
    """Дата приложения берётся из файла: её можно менять между сценами."""
    sys.path.insert(0, str(ROOT))
    from app import clock

    def now():
        day = CLOCK_OVERRIDE.get("day") or date.fromisoformat(date_file.read_text().strip())
        return datetime(day.year, day.month, day.day, 12, 0, tzinfo=ZoneInfo(clock.budget_timezone()))

    clock.now = now
    clock.today = lambda: now().date()


def serve(port: int, date_file: Path) -> None:
    patch_clock(date_file)
    import uvicorn

    uvicorn.run("app.cashflow_app:app", host="127.0.0.1", port=port, log_level="warning")


def bot_texts(date_file: Path) -> None:
    """Сообщения бота дословно по шаблонам bot.py на текущих демо-данных."""
    patch_clock(date_file)
    os.environ.setdefault("MINI_APP_URL", "https://example.invalid")
    import bot
    from app import clock
    from app.db import user_scope
    from app.morning_reports import pending_morning_reports
    from app.quick_stats import quick_numbers

    with user_scope(1):
        numbers = quick_numbers(1)
    result = {
        "today": bot.menu_reply_text(bot.BTN_TODAY, numbers),
        "spent": bot.menu_reply_text(bot.BTN_SPENT, numbers),
        "card": bot.menu_reply_text(bot.BTN_FREE if numbers and numbers.get("mode") == "irregular" else bot.BTN_CARD, numbers),
    }
    # Вчерашний отчёт уже был отправлен: от него считается «Изменение дневной суммы».
    from app.morning_reports import record_morning_report

    CLOCK_OVERRIDE["day"] = clock.today() - timedelta(days=1)
    for report in pending_morning_reports(clock.now()):
        record_morning_report(report)
    CLOCK_OVERRIDE.clear()
    reports = pending_morning_reports(clock.now())
    if reports:
        result["morning"] = bot.morning_report_text(reports[0])
    print(json.dumps(result, ensure_ascii=False))


# ---------------------------------------------------------------------------
# Сервер приложения и снимки базы
# ---------------------------------------------------------------------------

def free_port() -> int:
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


class App:
    """Приложение на временной базе.  Снимок базы (save/load) позволяет
    начинать каждую главу с одних и тех же данных."""

    def __init__(self, workdir: Path):
        self.workdir = workdir
        self.data = workdir / "data"
        self.snapshots = workdir / "snapshots"
        self.date_file = workdir / "today.txt"
        self.port = free_port()
        self.proc: subprocess.Popen | None = None
        self.today = date(2026, 1, 1)
        self.snapshots.mkdir(parents=True, exist_ok=True)

    # --- процесс ---
    def env(self) -> dict:
        env = dict(os.environ)
        env.update({
            "DEV_MODE": "true", "BOT_TOKEN": "", "BUDGET_TIMEZONE": TZ,
            "DATABASE_PATH": str(self.data / "budget.sqlite3"), "PYTHONPATH": str(ROOT),
        })
        return env

    def start(self) -> None:
        self.data.mkdir(parents=True, exist_ok=True)
        log = open(self.workdir / "server.log", "a")
        self.proc = subprocess.Popen(
            [sys.executable, __file__, "--serve", str(self.port), str(self.date_file)],
            cwd=ROOT, env=self.env(), stdout=log, stderr=subprocess.STDOUT,
        )
        for _ in range(200):
            try:
                urllib.request.urlopen(f"http://127.0.0.1:{self.port}/", timeout=1)
                return
            except Exception:
                time.sleep(0.05)
        raise RuntimeError("Приложение не запустилось, см. server.log")

    def stop(self) -> None:
        if self.proc:
            self.proc.terminate()
            self.proc.wait(10)
            self.proc = None

    def reset(self, day: date) -> None:
        """Пустая база: как у человека, который впервые открыл приложение."""
        self.stop()
        shutil.rmtree(self.data, ignore_errors=True)
        self.set_date(day)
        self.start()

    def save(self, name: str) -> None:
        self.stop()
        target = self.snapshots / name
        shutil.rmtree(target, ignore_errors=True)
        shutil.copytree(self.data, target)
        (target / "today.txt").write_text(self.today.isoformat())
        self.start()

    def load(self, name: str, day: date | None = None) -> None:
        self.stop()
        shutil.rmtree(self.data, ignore_errors=True)
        shutil.copytree(self.snapshots / name, self.data)
        saved = date.fromisoformat((self.data / "today.txt").read_text())
        self.set_date(day or saved)
        self.start()

    # --- дата ---
    def set_date(self, day: date) -> None:
        self.today = day
        self.date_file.write_text(day.isoformat())

    # --- API ---
    def api(self, method: str, path: str, body: dict | None = None):
        data = None if body is None else json.dumps(body).encode()
        req = urllib.request.Request(f"http://127.0.0.1:{self.port}{path}", data=data, method=method,
                                     headers={"Content-Type": "application/json"})
        try:
            with urllib.request.urlopen(req, timeout=60) as res:
                raw = res.read()
                return json.loads(raw) if raw else None
        except urllib.error.HTTPError as exc:
            raise RuntimeError(f"{method} {path}: {exc.code} {exc.read().decode()}") from None

    def get(self, path):
        return self.api("GET", path)

    def post(self, path, body=None):
        return self.api("POST", path, body or {})

    def put(self, path, body=None):
        return self.api("PUT", path, body or {})

    def bot_texts(self) -> dict:
        out = subprocess.run([sys.executable, __file__, "--bot-texts", str(self.date_file)],
                             cwd=ROOT, env=self.env(), capture_output=True, text=True, check=True)
        return json.loads(out.stdout.strip().splitlines()[-1])

    @property
    def url(self) -> str:
        return f"http://127.0.0.1:{self.port}/"


# ---------------------------------------------------------------------------
# Браузер, отметки и снимки
# ---------------------------------------------------------------------------

MARKS_JS = """
({marks, color}) => {
  document.getElementById('guide-marks')?.remove();
  const vw = window.innerWidth, vh = window.innerHeight, R = 13;
  // Текст на экране: кружок не должен его закрывать.  Пока открыто окно,
  // важен только текст окна.
  const root = document.querySelector('dialog[open]') || document.body;
  const busy = [];
  const walker = document.createTreeWalker(root, NodeFilter.SHOW_TEXT);
  const range = document.createRange();
  while (walker.nextNode()) {
    const node = walker.currentNode;
    if (!node.nodeValue.trim()) continue;
    range.selectNodeContents(node);
    for (const r of range.getClientRects()) {
      if (r.width > 0 && r.height > 0 && r.bottom > 0 && r.top < vh) busy.push(r);
    }
  }
  root.querySelectorAll('input:not([type=hidden]),select,textarea,svg,img').forEach(el => {
    const r = el.getBoundingClientRect();
    if (!r.width || !r.height) return;
    if (el.matches('input[type=checkbox],svg,img')) { busy.push(r); return; }
    // Значение поля пишется слева: его область и считаем занятой.
    busy.push({left: r.left + 6, right: r.left + Math.min(r.width - 6, 170), top: r.top + 8, bottom: r.bottom - 8});
  });
  const placed = [];
  const overlap = (a, b) => Math.max(0, Math.min(a.right, b.right) - Math.max(a.left, b.left)) *
                            Math.max(0, Math.min(a.bottom, b.bottom) - Math.max(a.top, b.top));
  const layer = document.createElement('div');
  layer.id = 'guide-marks';
  // Слой в top layer: он рисуется поверх открытых окон <dialog>.
  layer.setAttribute('popover', 'manual');
  layer.style.cssText = 'position:fixed;inset:0;width:100vw;height:100vh;max-width:none;max-height:none;' +
    'margin:0;padding:0;border:0;background:transparent;pointer-events:none;overflow:visible;';
  const frames = marks.map(m => {
    const pad = m.pad;
    return {x: m.box.x - pad, y: m.box.y - pad, w: m.box.width + 2 * pad, h: m.box.height + 2 * pad};
  });
  marks.forEach((m, i) => {
    const {x, y, w, h} = frames[i];
    if (m.frame) {
      const frame = document.createElement('div');
      frame.style.cssText = `position:absolute;left:${x}px;top:${y}px;width:${w}px;height:${h}px;` +
        `border:2px solid ${color};border-radius:10px;box-sizing:border-box;`;
      layer.appendChild(frame);
    }
    const spots = {
      tr: [x + w, y], br: [x + w, y + h], tl: [x, y], bl: [x, y + h],
      mr: [x + w, y + h / 2], ml: [x, y + h / 2], tc: [x + w / 2, y], bc: [x + w / 2, y + h],
      or: [x + w + R + 3, y + h / 2], ol: [x - R - 3, y + h / 2],
      ar: [x + w - R - 4, y - R - 2], al: [x + R + 4, y - R - 2],
      ur: [x + w - R - 4, y + h + R + 2], ul: [x + R + 4, y + h + R + 2],
    };
    const order = m.pos ? [m.pos] : Object.keys(spots);
    let best = null;
    order.forEach((key, rank) => {
      let [cx, cy] = spots[key];
      const shift = Math.abs(cx - Math.min(vw - R - 2, Math.max(R + 2, cx))) +
                    Math.abs(cy - Math.min(vh - R - 2, Math.max(R + 2, cy)));
      cx = Math.min(vw - R - 2, Math.max(R + 2, cx));
      cy = Math.min(vh - R - 2, Math.max(R + 2, cy));
      const dot = {left: cx - R, right: cx + R, top: cy - R, bottom: cy + R};
      let score = rank * 6 + shift * 8;
      for (const r of busy) score += overlap(dot, r);
      for (const r of placed) score += 6 * overlap(dot, r);
      // Кружок не должен лечь на рамку другой отметки.
      frames.forEach((f, j) => {
        if (j === i) return;
        const edge = {left: f.x - 2, right: f.x + f.w + 2, top: f.y - 2, bottom: f.y + f.h + 2};
        const inner = {left: f.x + 2, right: f.x + f.w - 2, top: f.y + 2, bottom: f.y + f.h - 2};
        score += 2 * (overlap(dot, edge) - overlap(dot, inner));
      });
      if (!best || score < best.score) best = {score, cx, cy, dot};
    });
    placed.push(best.dot);
    const dot = document.createElement('div');
    dot.textContent = String(m.n);
    dot.style.cssText = `position:absolute;left:${best.cx - R}px;top:${best.cy - R}px;width:${2 * R}px;height:${2 * R}px;` +
      `border-radius:50%;background:${color};color:#fff;font:700 14px/26px system-ui,-apple-system,sans-serif;` +
      `text-align:center;box-shadow:0 0 0 2px #fff,0 1px 4px rgba(0,0,0,.35);`;
    layer.appendChild(dot);
  });
  document.body.appendChild(layer);
  layer.showPopover();
}
"""

SCREEN_TEXT_JS = """() => {
  const clean = t => (t || '').replace(/\\s+/g, ' ').trim();
  const lines = document.body.innerText.split('\\n').map(clean);
  document.querySelectorAll('select option, input, textarea').forEach(el => {
    lines.push(clean(el.textContent), clean(el.value), clean(el.placeholder));
  });
  return lines.filter(Boolean);
}"""

RECT_JS = "e => { const r = e.getBoundingClientRect(); return {x: r.x, y: r.y, width: r.width, height: r.height}; }"


def compress(path: Path) -> None:
    """Палитра из 256 цветов: интерфейс плоский, текст остаётся чётким, а файл
    становится в 3–4 раза меньше."""
    from PIL import Image

    with Image.open(path) as image:
        palette = image.convert("RGB").quantize(colors=256, method=Image.Quantize.FASTOCTREE,
                                                dither=Image.Dither.NONE)
    palette.save(path, optimize=True)


class M:
    """Отметка: номер у элемента (или нескольких) и рамка вокруг.  Место
    кружка выбирается само: там, где он меньше всего закрывает текст.  pos
    задаёт место явно (ключ из spots в MARKS_JS)."""

    def __init__(self, n: int, target, pos: str | None = None, pad: int = 4, frame: bool = True):
        self.n, self.target, self.pos, self.pad, self.frame = n, target, pos, pad, frame


class Shooter:
    def __init__(self, browser, app: App):
        self.browser = browser
        self.app = app
        self.context = None
        self.page = None
        self.facts: dict = {}
        self.names: list[str] = []
        self.screen_texts: set[str] = set()

    # --- страница ---
    def open(self, wait: str = "#leftToday") -> None:
        """Новая вкладка на дату приложения, светлая тема, телефон 390×844."""
        self.close()
        self.context = self.browser.new_context(
            viewport=VIEWPORT, device_scale_factor=2, color_scheme="light", locale="ru-RU",
            timezone_id=TZ, reduced_motion="reduce", is_mobile=True, has_touch=True,
        )
        page = self.context.new_page()
        page.route("**/telegram-web-app.js*", lambda route: route.fulfill(
            status=200, content_type="application/javascript", body=TELEGRAM_STUB))
        page.add_init_script(TELEGRAM_STUB)
        page.add_init_script("try { localStorage.setItem('budget-theme-mode', 'light'); } catch (e) {}")
        day = self.app.today
        page.clock.set_fixed_time(datetime(day.year, day.month, day.day, 12, 0, tzinfo=ZoneInfo(TZ)))
        page.on("dialog", lambda dialog: dialog.accept())
        self.page = page
        page.goto(self.app.url)
        self.settle()

    def close(self) -> None:
        if self.context:
            self.context.close()
            self.context = None
            self.page = None

    def settle(self, ms: int = 400) -> None:
        self.page.wait_for_load_state("networkidle")
        self.page.evaluate("document.fonts.ready")
        self.page.wait_for_timeout(ms)

    def reload(self) -> None:
        self.page.click("#refreshBtn")
        self.settle()

    # --- навигация ---
    def nav(self, name: str) -> None:
        """home / operations / analytics / buffer (Резервы) / settings."""
        if name == "settings":
            self.page.click(".topbar .settings-btn")
        else:
            self.page.click(f'.bottom-nav [data-nav="{name}"]')
        # Приложение прокручивает страницу вверх плавно: дожидаемся конца.
        self.page.evaluate("window.scrollTo({top: 0, behavior: 'instant'})")
        self.settle()

    def tab(self, seg: str) -> None:
        """Вкладка раздела «Резервы»: buffer, piggy, plan, reserve, obligations."""
        self.page.locator(f'.page.active [data-seg="{seg}"]').click()
        self.settle()

    def loc(self, target):
        if isinstance(target, str):
            return self.page.locator(target).first
        return target

    def scroll(self, target, block: str = "center") -> None:
        """Прокрутить к элементу.  start — элемент вверху экрана с небольшим
        отступом (центр — чтобы он не ушёл под нижнее меню)."""
        self.loc(target).evaluate(
            """(e, block) => {
                e.scrollIntoView({block, behavior: 'instant'});
                if (block === 'start' && !e.closest('dialog')) window.scrollBy({top: -12, behavior: 'instant'});
            }""", block)
        self.page.wait_for_timeout(150)

    def scroll_top(self) -> None:
        self.page.evaluate("window.scrollTo(0, 0)")
        self.page.wait_for_timeout(150)

    def box(self, target) -> dict:
        targets = target if isinstance(target, (list, tuple)) else [target]
        boxes = [self.loc(t).evaluate(RECT_JS) for t in targets]
        x1 = min(b["x"] for b in boxes)
        y1 = min(b["y"] for b in boxes)
        x2 = max(b["x"] + b["width"] for b in boxes)
        y2 = max(b["y"] + b["height"] for b in boxes)
        return {"x": x1, "y": y1, "width": x2 - x1, "height": y2 - y1}

    def text(self, target) -> str:
        return " ".join(self.loc(target).inner_text().split())

    def fact(self, value):
        """Текст элемента для numbers.json: селектор или локатор."""
        if isinstance(value, str) and value.startswith(("#", ".", "[")) or hasattr(value, "inner_text"):
            return self.text(value)
        return value

    # --- снимок ---
    def shot(self, name: str, marks: list[M] = (), facts: dict | None = None) -> None:
        page = self.page
        page.evaluate("document.getElementById('toast')?.classList.remove('show')")
        page.evaluate("document.activeElement && document.activeElement.blur && document.activeElement.blur()")
        page.wait_for_timeout(120)
        # Рамки ставятся только когда экран перестал двигаться.
        boxes = None
        for _ in range(20):
            fresh = [self.box(m.target) for m in marks]
            if fresh == boxes:
                break
            boxes = fresh
            page.wait_for_timeout(120)
        payload = [{"n": m.n, "box": box, "pos": m.pos, "pad": m.pad, "frame": m.frame}
                   for m, box in zip(marks, boxes or [])]
        page.evaluate(MARKS_JS, {"marks": payload, "color": MARK})
        page.wait_for_timeout(80)
        page.screenshot(path=str(IMAGES / name))
        compress(IMAGES / name)
        page.evaluate("document.getElementById('guide-marks')?.remove()")
        # Весь видимый текст экрана: по нему --check сверяет названия в тексте глав.
        self.screen_texts.update(page.evaluate(SCREEN_TEXT_JS))
        self.names.append(name)
        if facts:
            self.facts[name] = {key: self.fact(value) for key, value in facts.items()}
        print("  ✓", name)

    def element_shot(self, name: str, selector: str) -> None:
        self.page.locator(selector).screenshot(path=str(IMAGES / name))
        compress(IMAGES / name)
        self.names.append(name)
        print("  ✓", name)


HOME_FACTS = {
    "Осталось на сегодня": "#leftToday", "Лимит на день": "#dailyAvailable", "Потрачено": "#spentToday",
    "На карте": "#cardCash", "Буфер": "#cardBuffer", "Копилка": "#cardPiggy", "Буфер + копилка": "#cardReservesTotal",
    "Период": "#periodCaption",
}


# ---------------------------------------------------------------------------
# Демо-данные (только вымышленные суммы)
# ---------------------------------------------------------------------------

def days(start: date, end: date):
    day = start
    while day <= end:
        yield day
        day += timedelta(days=1)


def category_ids(app: App) -> dict:
    return {c["title"]: c["id"] for c in app.get("/api/bootstrap")["categories"]}


NOTES = {
    "Продукты": ["Продукты", "Овощи и фрукты", "Продукты на неделю", "Молоко и хлеб"],
    "Транспорт": ["Метро", "Такси", "Автобус"],
    "Кафе": ["Обед", "Кофе", "Ужин с друзьями"],
    "Дом": ["Бытовая химия", "Лампочки"],
    "Здоровье": ["Аптека"],
    "Развлечения": ["Кино", "Книга"],
    "Покупки": ["Футболка", "Зарядка для телефона"],
}


def daily_expenses(rng: random.Random, day: date, target: float, overspend: bool = True) -> list[tuple[str, int]]:
    """Обычные траты одного дня, примерно в пределах дневного лимита: продукты
    почти каждый день, транспорт по будням, кафе пару раз в неделю, остальное
    изредка.  Иногда день без трат, иногда небольшой перерасход."""
    roll = rng.random()
    if roll < 0.08 or target <= 0:
        return []
    budget = target * (rng.uniform(1.05, 1.25) if overspend and roll > 0.95 else rng.uniform(0.6, 0.98))
    items: list[tuple[str, int]] = []
    if day.weekday() < 5 and rng.random() < 0.75:
        items.append(("Транспорт", rng.choice([60, 60, 120, 350, 420])))
    if rng.random() < 0.3:
        items.append(("Кафе", rng.randrange(25, 90) * 10))
    if rng.random() < 0.1:
        items.append((rng.choice(["Дом", "Здоровье", "Развлечения", "Покупки"]), rng.randrange(30, 150) * 10))
    if not overspend:
        # Аккуратный режим: мелкие траты тоже не выходят за лимит.
        kept, total = [], 0
        for title, value in items:
            if total + value <= budget:
                kept.append((title, value))
                total += value
        items = kept
    rest = budget - sum(value for _, value in items)
    if rest >= 200:
        items.insert(0, ("Продукты", int(rest // 10) * 10))
    return items


def add_expense(app: App, cats: dict, title: str, amount: float, day: date, note: str = "", **extra) -> dict:
    return app.post("/api/transactions", {
        "type": "expense", "amount": amount, "tx_date": day.isoformat(),
        "category_id": cats[title], "note": note, **extra,
    })


def pending_payday(app: App) -> dict | None:
    return next((p for p in app.get("/api/buffer").get("periods", []) if p.get("income_pending")), None)


def confirm_payday(app: App, amount: float | None = None) -> None:
    period = pending_payday(app)
    if period:
        app.put(f"/api/cashflow/income-overrides/{period['override_key']}",
                {"amount": period["payday_amount"] if amount is None else amount, "confirmed": True})


def add_bills(app: App, cats: dict) -> dict:
    bills = {}
    for title, amount, day, category in (("Аренда", 25000, 10, "Дом"), ("Интернет", 700, 15, "Дом"),
                                         ("Телефон", 500, 25, "Другое")):
        app.post("/api/bill-rules", {"title": title, "amount": amount, "day_of_month": day,
                                     "category_id": cats[category], "active": True,
                                     "effective_date": app.today.isoformat()})
    for rule in app.get("/api/bootstrap")["bill_rules"]:
        bills[rule["title"]] = rule
    return bills


PAYROLL_START = date(2026, 8, 1)
PAYROLL_TODAY = date(2026, 10, 14)   # среда, середина периода 8–22 октября
PAYDAY = date(2026, 10, 22)          # аванс
IRREGULAR_START = date(2026, 9, 1)

# Доходы режима «Подработки»: дата, сумма, откуда.
SIDE_JOBS = [
    (date(2026, 9, 3), 22000, "Ремонт"), (date(2026, 9, 9), 6500, "Доставка"),
    (date(2026, 9, 16), 4000, "Репетиторство"), (date(2026, 9, 19), 28000, "Ремонт"),
    (date(2026, 9, 26), 6200, "Доставка"), (date(2026, 9, 30), 4000, "Репетиторство"),
    (date(2026, 10, 5), 26000, "Ремонт"), (date(2026, 10, 8), 7400, "Доставка"),
    (date(2026, 10, 12), 4000, "Репетиторство"),
]


def simulate(app: App, start: date, end: date, rng: random.Random, *, irregular: bool = False,
             confirm: bool = True, overspend: bool = True, extra: dict | None = None) -> None:
    """Проматывает дни: в свой день отмечаются оплаченными обязательные
    платежи, подтверждаются выплаты («Получил ✓»), записываются доходы
    подработок и обычные траты.  Последний день остаётся без трат."""
    cats = category_ids(app)
    bills = {rule["title"]: rule for rule in app.get("/api/bootstrap")["bill_rules"]}
    numbers = "/api/irregular" if irregular else "/api/cashflow"
    for day in days(start, end):
        app.set_date(day)
        if confirm and not irregular:
            confirm_payday(app)
        for when, value, source in SIDE_JOBS if irregular else ():
            if when == day:
                app.post("/api/irregular/incomes", {"amount": value, "tx_date": day.isoformat(),
                                                    "income_source": source})
        for rule in bills.values():
            if day.day == rule["day_of_month"]:
                app.post(f"/api/bills/{rule['id']}/pay", {"due_date": day.isoformat()})
        if day == end:
            break
        target = app.get(numbers)["available_today"]
        for title, value in daily_expenses(rng, day, target, overspend):
            add_expense(app, cats, title, value, day, rng.choice(NOTES.get(title, [""])), **(extra or {}))
        if not irregular and day.day in (12, 27):
            left = app.get(numbers)["available_today"]
            if left >= 300:
                app.post("/api/piggy-bank/deposit", {"amount": int(left // 100) * 100, "source": "daily_budget",
                                                     "movement_date": day.isoformat(), "note": ""})


def seed_payroll(app: App) -> None:
    """Режим «Зарплата и аванс»: оклад 80 000 и премия 10 000 до НДФЛ, три
    обязательных платежа, старт 1 августа с капиталом 40 000 и два с половиной
    месяца повседневных трат.  Снимок «payroll» — утро 14 октября."""
    app.reset(PAYROLL_START)
    cats = category_ids(app)
    app.put("/api/payroll-settings", {"payroll_enabled": True, "salary_gross": 80000, "bonus_gross": 10000,
                                      "tax_rate": 13, "salary_day": 7, "advance_day": 22})
    add_bills(app, cats)
    app.put("/api/cashflow-settings", {"cashflow_enabled": True, "start_date": PAYROLL_START.isoformat(),
                                       "start_capital": 40000})
    app.post("/api/piggy-bank/deposit", {"amount": 5000, "movement_date": "2026-08-01",
                                         "note": "Отложено раньше"})
    simulate(app, PAYROLL_START, PAYROLL_TODAY, random.Random(7))
    app.save("payroll")
    # Тот же человек неделю спустя: 22 октября пришёл аванс, но ещё не подтверждён.
    simulate(app, PAYROLL_TODAY, PAYDAY, random.Random(22), confirm=False)
    app.save("payday")


def seed_irregular(app: App) -> None:
    """Режим «Подработки»: старт 1 сентября (45 000, из них 5 000 резерв),
    доходы от ремонта, доставки и репетиторства, те же платежи."""
    app.reset(IRREGULAR_START)
    cats = category_ids(app)
    add_bills(app, cats)
    start = {"start_date": IRREGULAR_START.isoformat(), "start_total": 45000, "start_reserve": 5000}
    app.put("/api/irregular/settings", {"reserve_percent": 10, "stretch_days": 14, "bills_scope": "month",
                                        "lookahead_days": 30, **start})
    app.put("/api/budget-mode", {"budget_mode": "irregular", **start})
    rng = random.Random(11)
    simulate(app, IRREGULAR_START, date(2026, 9, 23), rng, irregular=True)
    app.post("/api/irregular/reserve/withdraw", {"amount": 3500, "movement_date": "2026-09-23",
                                                 "purpose": "pay_expense", "reason": "Стоматолог",
                                                 "category_id": cats["Здоровье"]})
    simulate(app, date(2026, 9, 23), PAYROLL_TODAY, rng, irregular=True)
    app.save("irregular")


ACCOUNT_ROWS = [
    {"title": "Карта 1", "emoji": "💳", "kind": "debit"},
    {"title": "Карта 2", "emoji": "💳", "kind": "debit"},
    {"title": "Накопительный", "emoji": "🏦", "kind": "savings"},
    {"title": "Кредитка", "emoji": "💳", "kind": "credit", "credit_limit": 50000, "payment_day": 20},
]
CARD2_BALANCE = 3000


def seed_accounts(app: App) -> None:
    """Зарплата + карты и счета: на «Карте 1» деньги на траты, на
    «Накопительном» буфер и копилка, «Карта 2» вне бюджета, «Кредитка» без долга."""
    app.load("payroll")
    payroll_today(app)
    amounts = {e["key"]: e["amount"] for e in app.get("/api/accounts?with_amounts=true")["envelopes"]}
    balances = [amounts["spending"], CARD2_BALANCE, round(amounts["buffer"] + amounts["piggy"], 2), 0]
    app.post("/api/accounts/setup", {
        "accounts": [{**row, "balance": balances[i], "is_main": i == 0} for i, row in enumerate(ACCOUNT_ROWS)],
        "envelopes": {"spending": 0, "buffer": 2, "piggy": 2},
    })
    app.save("accounts")
    # Куртка за 3 200 ₽ куплена кредиткой.
    ids = account_ids(app)
    add_expense(app, category_ids(app), "Покупки", 3200, PAYROLL_TODAY, "Куртка", account_id=ids["Кредитка"])
    app.save("accounts-credit")


def account_ids(app: App) -> dict:
    return {a["title"]: a["id"] for a in app.get("/api/accounts")["accounts"]}


# ---------------------------------------------------------------------------
# Сцены по главам
# ---------------------------------------------------------------------------

def payroll_today(app: App) -> None:
    """Утро 14 октября: уже записаны две траты."""
    cats = category_ids(app)
    add_expense(app, cats, "Продукты", 640, PAYROLL_TODAY, "Молоко и хлеб")
    add_expense(app, cats, "Транспорт", 60, PAYROLL_TODAY, "Метро")


def card(sh: Shooter, title: str):
    """Карточка раздела по её заголовку."""
    page = sh.page
    return page.locator(".page.active .card").filter(has=page.locator(f'h3:text-is("{title}")')).first


LABEL_JS = """(labels, text) => labels.findIndex(label => {
    const own = [...label.childNodes].filter(n => n.nodeType === 3).map(n => n.nodeValue).join(' ').trim();
    const caption = own || label.querySelector('span')?.textContent.trim() || '';
    return caption === text || caption.startsWith(text + ':');
})"""


def label(scope, text: str):
    """Поле формы по его подписи: точное совпадение («Куда» ≠ «Откуда»)."""
    labels = scope.locator("label")
    index = labels.evaluate_all(LABEL_JS, text)
    if index < 0:
        raise RuntimeError(f"Нет поля с подписью «{text}»")
    return labels.nth(index)


def dlg(sh: Shooter, dialog_id: str):
    return sh.page.locator(f"#{dialog_id}")


# --- 1. Что это за приложение -------------------------------------------------

def ch01(app: App, sh: Shooter, texts: dict) -> None:
    app.load("payroll")
    payroll_today(app)
    sh.open()
    sh.shot("01-overview-home.png", [
        M(1, "#todayHero"), M(2, ".quick-actions"), M(3, ".card.period"),
    ], facts=HOME_FACTS)
    app.load("irregular")
    sh.open()
    sh.shot("01-overview-irregular.png", [M(1, "#todayHero"), M(2, "#irrHome")])


# --- 2. Как открыть и начать -----------------------------------------------------

def ch02(app: App, sh: Shooter, texts: dict) -> None:
    app.reset(PAYROLL_TODAY)
    sh.open()
    sh.shot("02-start-first-launch.png", facts=HOME_FACTS)
    sh.shot("02-start-open-settings.png", [M(1, ".topbar .settings-btn", pad=3)])
    sh.nav("settings")
    mode = card(sh, "Как приходят деньги")
    sh.scroll(mode, "start")
    sh.shot("02-start-choose-mode.png", [
        M(1, '[data-budget-mode="payroll"]', pad=2), M(2, '[data-budget-mode="irregular"]', pad=2),
    ])


# --- 3. Главный экран ----------------------------------------------------------

def ch03(app: App, sh: Shooter, texts: dict) -> None:
    app.load("payroll")
    payroll_today(app)
    sh.open()
    sh.shot("03-home-numbers.png", [
        M(1, "#leftToday"), M(2, "#todayHero .hero-row span:first-child", pad=3),
        M(3, "#todayHero .hero-row span:last-child", pad=3),
        M(4, "#cardCash >> xpath=.."), M(5, "#cardBuffer >> xpath=.."),
        M(6, "#cardPiggy >> xpath=.."), M(7, "#cardReservesTotal >> xpath=.."),
    ], facts=HOME_FACTS)
    sh.scroll(".card.period", "start")
    sh.shot("03-home-period.png", [
        M(1, "#nextPayTitle", pad=3), M(2, "#periodStrip .day.today", pad=2),
        M(3, "#periodStrip .day.bill >> nth=-1", pad=2), M(4, "#periodStrip .day.payday", pad=2),
        M(5, ".card.home-bills"), M(6, ".card.home-recent"),
    ], facts={"Заголовок": "#nextPayTitle", "Скоро платежи": "#homeBills"})
    sh.scroll(".card.home-buffer", "start")
    sh.shot("03-home-buffer.png", [M(1, ".card.home-buffer"), M(2, ".card.home-forecast")],
            facts={"В буфере сейчас": "#reserveBalance", "Пояснение": "#reserveReason"})
    # Пример: обед в кафе за 450 ₽.
    sh.scroll_top()
    sh.page.click("#quickExpenseBtn")
    sh.page.fill("#expenseAmount", "450")
    dlg(sh, "expenseDialog").locator(".cat-tile", has_text="Кафе").click()
    sh.page.fill("#expenseNote", "Обед")
    sh.page.click("#saveExpenseBtn")
    sh.settle()
    sh.scroll_top()
    sh.shot("03-home-after-expense.png", [
        M(1, "#leftToday"), M(2, "#todayHero .hero-row span:last-child", pad=3),
        M(3, "#cardCash >> xpath=.."),
    ], facts=HOME_FACTS)
    app.save("payroll-cafe")
    app.set_date(PAYROLL_TODAY + timedelta(days=1))
    sh.open()
    sh.shot("03-home-next-day.png", [
        M(1, "#leftToday"), M(2, "#todayHero .hero-row span:first-child", pad=3),
        M(3, "#cardCash >> xpath=.."),
    ], facts=HOME_FACTS)


# --- 4. Схемы ---------------------------------------------------------------------

SCHEME_CSS = """
<link href="https://fonts.googleapis.com/css2?family=Onest:wght@400;500;600;700&display=swap" rel="stylesheet">
<style>
body{margin:0;background:#f4f6f8;font-family:Onest,system-ui,sans-serif;color:#0e1116}
.wrap{width:390px;box-sizing:border-box;padding:18px 16px 22px}
h1{font-size:19px;margin:0 0 4px}
.sub{color:#6b7280;font-size:13px;margin:0 0 14px}
.box{background:#fff;border-radius:16px;padding:12px 14px;border:2px solid var(--c,#e6e8ee)}
.box b{display:block;font-size:15px}
.box span{display:block;color:#475066;font-size:12.5px;line-height:1.35;margin-top:2px}
.dot{display:inline-block;width:10px;height:10px;border-radius:3px;background:var(--c);margin-right:6px}
.arrow{text-align:center;color:#6b7280;font-size:12.5px;line-height:1.3;padding:5px 0}
.arrow i{display:block;font-style:normal;font-size:18px;color:#9aa1ad;line-height:1}
.row{display:grid;grid-template-columns:1fr 1fr;gap:10px}
.row3{display:grid;grid-template-columns:1fr 1fr 1fr;gap:8px}
.row3 .box{padding:10px}
.note{margin-top:14px;font-size:12.5px;color:#475066;background:#fff;border-radius:14px;padding:10px 12px}
</style>"""

SCHEME_PAYROLL = SCHEME_CSS + """
<div class="wrap" id="scheme">
<h1>Как устроен бюджет: зарплата и аванс</h1>
<p class="sub">Каждый прямоугольник — «конверт»: сколько денег и на что они предназначены.</p>
<div class="box" style="--c:#1f8f57"><b>Зарплата и аванс</b><span>приходят на карту в дни выплаты</span></div>
<div class="arrow"><i>↓</i>вся выплата сначала попадает в буфер</div>
<div class="box" style="--c:#8a6fd6"><b><i class="dot"></i>Буфер</b><span>держит деньги на обязательные платежи и выравнивает большие и маленькие выплаты</span></div>
<div class="row">
 <div><div class="arrow"><i>↓</i>в день платежа</div><div class="box"><b>Обязательные платежи</b><span>аренда, интернет, телефон</span></div></div>
 <div><div class="arrow"><i>↓</i>в день выплаты, на весь период</div><div class="box" style="--c:#2481cc"><b><i class="dot"></i>На карте</b><span>деньги на обычные траты до следующей выплаты</span></div></div>
</div>
<div class="row"><div></div><div><div class="arrow"><i>↓</i>делятся на дни</div><div class="box" style="--c:#2481cc"><b>Лимит на день</b><span>сколько можно потратить сегодня</span></div></div></div>
<div class="note"><b><i class="dot" style="--c:#1f8f57"></i>Копилка</b> — отдельные накопления. Деньги попадают туда и уходят оттуда только когда вы сами переносите их кнопками.</div>
</div>"""

SCHEME_IRREGULAR = SCHEME_CSS + """
<div class="wrap" id="scheme">
<h1>Как устроен бюджет: подработки</h1>
<p class="sub">Каждый доход сразу делится по конвертам, по порядку сверху вниз.</p>
<div class="box" style="--c:#1f8f57"><b>Доход</b><span>любая сумма в любой день</span></div>
<div class="arrow"><i>↓</i>делится сразу</div>
<div class="row3">
 <div class="box" style="--c:#eb6834"><b>1. Резерв</b><span>заданный процент: на болезни, поломки, дни без заработка</span></div>
 <div class="box" style="--c:#8a6fd6"><b>2. На обязательные</b><span>сколько не хватает на ближайшие платежи</span></div>
 <div class="box" style="--c:#2481cc"><b>3. Свободные</b><span>всё остальное</span></div>
</div>
<div class="row"><div></div><div><div class="arrow"><i>↓</i>делятся на дни, по умолчанию на 14</div><div class="box" style="--c:#2481cc"><b>Лимит на день</b><span>сколько можно потратить сегодня</span></div></div></div>
<div class="note"><b><i class="dot" style="--c:#1f8f57"></i>Копилка</b> — отдельные накопления, только по вашим переводам. Деньги, которые ещё не пришли, не учитываются никогда.</div>
</div>"""


def ch04(app: App, sh: Shooter, texts: dict) -> None:
    for name, html in (("04-scheme-payroll.png", SCHEME_PAYROLL), ("04-scheme-irregular.png", SCHEME_IRREGULAR)):
        sh.close()
        sh.context = sh.browser.new_context(viewport={"width": 390, "height": 900}, device_scale_factor=2)
        sh.page = sh.context.new_page()
        sh.page.set_content(html)
        sh.page.wait_for_load_state("networkidle")
        sh.page.evaluate("document.fonts.ready")
        sh.element_shot(name, "#scheme")


# --- 5. Зарплата и аванс ------------------------------------------------------------

def ch05(app: App, sh: Shooter, texts: dict) -> None:
    app.load("payroll")
    payroll_today(app)
    sh.open()
    sh.nav("settings")
    pay = card(sh, "Зарплата и аванс")
    sh.scroll(pay, "start")
    sh.shot("05-payroll-settings.png", [
        M(1, label(pay, "Считать автоматически"), pad=3), M(2, label(pay, "Оклад до НДФЛ"), pad=3),
        M(3, label(pay, "Премия до НДФЛ"), pad=3), M(4, label(pay, "День зарплаты"), pad=3),
        M(5, label(pay, "День аванса"), pad=3), M(6, label(pay, "НДФЛ, %"), pad=3),
    ])
    sh.scroll("#payrollPreview", "center")
    lines = sh.page.locator("#payrollPreview .preview-line")
    sh.shot("05-payroll-shift.png", [
        M(1, "#savePayrollBtn", pad=3), M(2, lines.filter(has_text="Аванс"), pad=3),
        M(3, lines.filter(has_text="Зарплата + премия"), pad=3),
    ], facts={"Расчёт": "#payrollPreview"})
    start = card(sh, "Стартовый капитал")
    sh.scroll(start, "center")
    sh.shot("05-payroll-start-capital.png", [
        M(1, label(start, "Считать безопасный дневной лимит"), pad=3),
        M(2, label(start, "Капитал на начало дня"), pad=3), M(3, label(start, "Дата старта"), pad=3),
        M(4, label(start, "Прогноз, месяцев"), pad=3), M(5, "#saveCashflowBtn", pad=3),
        M(6, "#saveSettingsBtn", pad=3),
    ])
    # Раздел «Резервы» → «Буфер»
    sh.nav("buffer")
    sh.shot("05-payroll-buffer-page.png", [
        M(1, "#bufferPageBalance >> xpath=.."), M(2, "#bufferHorizon >> xpath=.."),
        M(3, "#bufferPageDaily >> xpath=..", pad=2), M(4, "#bufChart"),
    ], facts={"В буфере сейчас": "#bufferPageBalance", "Прогноз до": "#bufferHorizon",
              "Можно сегодня": "#bufferPageDaily", "В день": "#planDaily", "Конец прогноза": "#bufChartEnd"})
    cur = sh.page.locator("#planCards article.pc.cur")
    sh.scroll(cur, "center")
    sh.shot("05-payroll-period-card.png", [
        M(1, cur.locator(".pc-move"), pad=3), M(2, cur.locator(".pc-pay"), pad=3),
        M(3, cur.locator(".pc-flow"), pad=3), M(4, cur.locator(".pc-foot"), pad=3),
    ], facts={"Текущий период": "#planCards article.pc.cur",
              "Следующий период": "#planCards article.pc >> nth=1"})
    nxt = sh.page.locator("#planCards article.pc").nth(1)
    sh.scroll(nxt, "center")
    sh.shot("05-payroll-next-period.png", [M(1, nxt.locator(".pc-move"), pad=3),
                                          M(2, nxt.locator(".pc-foot"), pad=3)],
            facts={"Следующий период": "#planCards article.pc >> nth=1"})

    # Отпуск 26–30 октября, отпускные 15 000 ₽.
    sh.nav("settings")
    vac = card(sh, "Отпускные")
    sh.scroll(vac, "center")
    sh.page.click("#addVacationBtn")
    sh.page.fill("#vacationStart", "2026-10-26")
    sh.page.locator("#vacationStart").dispatch_event("change")
    sh.page.fill("#vacationEnd", "2026-10-30")
    sh.page.fill("#vacationAmount", "15000")
    d = dlg(sh, "vacationDialog")
    sh.shot("05-payroll-vacation-dialog.png", [
        M(1, label(d, "Начало"), pad=3), M(2, label(d, "Окончание"), pad=3),
        M(3, label(d, "Отпускные к получению"), pad=3), M(4, label(d, "Дата выплаты отпускных"), pad=3),
        M(5, d.locator("button.primary"), pad=3),
    ])
    d.locator("button.primary").click()
    sh.settle()
    sh.scroll("#payrollPreview", "center")
    lines = sh.page.locator("#payrollPreview .preview-line")
    sh.shot("05-payroll-vacation-preview.png", [
        M(1, lines.filter(has_text="Рабочих дней отпуска"), pad=3),
        M(2, lines.filter(has_text="Зарплата + премия"), pad=3),
    ], facts={"Расчёт": "#payrollPreview"})
    sh.scroll(vac, "center")
    sh.shot("05-payroll-vacation-list.png", [M(1, "#vacationsList .list-row", pad=3)])

    # Повышение зарплаты
    raise_card = card(sh, "Повышение зарплаты или изменение премии")
    sh.scroll(raise_card, "center")
    sh.page.click("#addPayrollChangeBtn")
    sh.page.fill("#payrollChangeMonth", "2027-01")
    sh.page.fill("#payrollChangeSalary", "90000")
    d = dlg(sh, "payrollChangeDialog")
    sh.shot("05-payroll-raise.png", [
        M(1, label(d, "Начать с месяца начисления"), pad=3), M(2, label(d, "Оклад до НДФЛ"), pad=3),
        M(3, label(d, "Премия до НДФЛ"), pad=3), M(4, "#savePayrollChangeBtn", pad=3),
    ])
    d.locator("[data-close]").click()

    # День выплаты: 22 октября, аванс пришёл, но не подтверждён.
    app.load("payday")
    texts["payday"] = app.bot_texts()
    sh.open()
    sh.shot("05-payroll-payday-home.png", [
        M(1, "#leftToday"), M(2, "#cardCash >> xpath=.."), M(3, "#cardBuffer >> xpath=.."),
        M(4, "#periodStrip .day.payday", pad=2),
    ], facts={**HOME_FACTS, "Заголовок": "#nextPayTitle"})
    sh.nav("buffer")
    cur = sh.page.locator("#planCards article.pc.cur")
    sh.scroll(cur, "center")
    sh.shot("05-payroll-confirm-salary.png", [
        M(1, cur.locator(".pc-status"), pad=3), M(2, cur.locator(".income-btn.accent"), pad=3),
        M(3, cur.locator(".income-btn:not(.accent)"), pad=3),
    ], facts={"Текущий период": "#planCards article.pc.cur"})
    cur.locator(".income-btn.accent").click()
    sh.page.fill("#cashflowIncomeAmount", "33800")
    d = dlg(sh, "cashflowIncomeDialog")
    sh.shot("05-payroll-confirm-dialog.png", [
        M(1, "#cashflowIncomeCaption", pad=3), M(2, label(d, "Сколько пришло"), pad=3),
        M(3, "#saveCashflowIncomeBtn", pad=3),
    ])
    sh.page.click("#saveCashflowIncomeBtn")
    sh.settle()
    cur = sh.page.locator("#planCards article.pc.cur")
    sh.scroll(cur, "center")
    sh.shot("05-payroll-after-confirm.png", [
        M(1, cur.locator(".pc-pay"), pad=3), M(2, cur.locator(".pc-flow > div").nth(0), pad=2),
        M(3, cur.locator(".pc-foot"), pad=3),
    ], facts={"Текущий период": "#planCards article.pc.cur"})
    sh.nav("home")
    sh.shot("05-payroll-payday-confirmed-home.png", [
        M(1, "#leftToday"), M(2, "#cardCash >> xpath=..", pad=4), M(3, "#cardBuffer >> xpath=..", pad=4),
    ], facts=HOME_FACTS)
    texts["payday-confirmed"] = app.bot_texts()
    app.save("payday-confirmed")
    app.set_date(PAYDAY + timedelta(days=1))
    sh.open()
    sh.shot("05-payroll-next-day.png", [
        M(1, "#leftToday"), M(2, "#todayHero .hero-row span:first-child", pad=3),
        M(3, "#cardCash >> xpath=.."), M(4, "#cardBuffer >> xpath=.."),
    ], facts=HOME_FACTS)
    # Отпускные пришли 23 октября: «Доход» → «В буфер».
    sh.page.click("#quickIncomeBtn")
    sh.page.fill("#incomeTxAmount", "15000")
    sh.page.select_option("#incomeTxDestination", "buffer")
    sh.page.fill("#incomeTxNote", "Отпускные")
    d = dlg(sh, "incomeTxDialog")
    sh.shot("05-payroll-extra-income.png", [
        M(1, label(d, "Сумма"), pad=3), M(2, label(d, "Дата"), pad=3),
        M(3, label(d, "Направить доход"), pad=3), M(4, "#saveIncomeTxBtn", pad=3),
    ])
    d.locator("[data-close]").click()
    # Без «Получил ✓»: тот же следующий день, но выплата так и не подтверждена.
    app.load("payday", PAYDAY + timedelta(days=1))
    sh.open()
    sh.shot("05-payroll-not-confirmed.png", [
        M(1, "#leftToday"), M(2, "#cardCash >> xpath=.."), M(3, "#cardBuffer >> xpath=.."),
    ], facts=HOME_FACTS)


# --- 6. Подработки --------------------------------------------------------------------

def ch06(app: App, sh: Shooter, texts: dict) -> None:
    app.load("irregular")
    texts["irregular"] = app.bot_texts()
    sh.open()
    sh.shot("06-irregular-home.png", [
        M(1, "#todayHero"), M(2, "#irrFree >> xpath=.."), M(3, "#irrReserve >> xpath=.."),
        M(4, "#irrBills >> xpath=.."), M(5, "#irrNoIncome >> xpath=.."),
    ], facts={**HOME_FACTS, "Свободных денег": "#irrFree", "Растягиваем": "#irrStretch",
              "Резерв": "#irrReserve", "Резерв хватит": "#irrReserveDays", "Отложено на обязательные": "#irrBills",
              "Ближайший": "#irrNextBill", "Дней без дохода": "#irrNoIncome", "Последний доход": "#irrLastIncome"})
    sh.nav("settings")
    irr = card(sh, "Подработки")
    sh.scroll(irr, "start")
    sh.shot("06-irregular-settings.png", [
        M(1, label(irr, "Процент в резерв"), pad=3), M(2, label(irr, "Цель резерва"), pad=3),
        M(3, label(irr, "Растягивать деньги на, дней"), pad=3),
        M(4, label(irr, "Откладывать на обязательные"), pad=3),
    ])
    sh.scroll("#saveIrregularBtn", "center")
    sh.shot("06-irregular-start.png", [
        M(1, label(irr, "Дата старта"), pad=3), M(2, label(irr, "Денег сейчас, без копилки"), pad=3),
        M(3, label(irr, "Из них резерв"), pad=3), M(4, "#saveIrregularBtn", pad=3),
    ])
    # Новый доход 12 000 ₽ за ремонт
    sh.nav("home")
    before = {k: sh.text(v) for k, v in {"Свободных денег": "#irrFree", "Резерв": "#irrReserve",
                                         "Лимит на день": "#dailyAvailable"}.items()}
    sh.facts["06-before-income"] = before
    sh.page.click("#quickIncomeBtn")
    sh.page.fill("#irrIncomeAmount", "12000")
    sh.page.fill("#irrIncomeSource", "Ремонт")
    sh.page.wait_for_timeout(900)
    d = dlg(sh, "irrIncomeDialog")
    sh.shot("06-irregular-income-dialog.png", [
        M(1, label(d, "Сумма"), pad=3), M(2, label(d, "Откуда"), pad=3),
        M(3, label(d, "Процент в резерв"), pad=3), M(4, label(d, "Направление"), pad=3),
        M(5, "#irrIncomePreview", pad=3), M(6, "#saveIrrIncomeBtn", pad=3),
    ], facts={"Разбивка": "#irrIncomePreview"})
    sh.page.click("#saveIrrIncomeBtn")
    sh.settle()
    sh.shot("06-irregular-after-income.png", [
        M(1, "#dailyAvailable >> xpath=..", pad=3), M(2, "#irrFree >> xpath=.."),
        M(3, "#irrReserve >> xpath=.."), M(4, "#irrNoIncome >> xpath=.."),
    ], facts={**HOME_FACTS, "Свободных денег": "#irrFree", "Растягиваем": "#irrStretch", "Резерв": "#irrReserve",
              "Дней без дохода": "#irrNoIncome"})
    app.load("irregular")
    sh.open()
    sh.nav("buffer")
    sh.shot("06-irregular-reserve.png", [
        M(1, "#irrReserveBalance", pad=3), M(2, "#irrWithdrawBtn", pad=3),
        M(3, "#irrDepositBtn", pad=3), M(4, "#irrReturnBtn", pad=3),
        M(5, "#irrReserveMovements", pad=3),
    ], facts={"Резерв": "#irrReserveBalance", "Пояснение": "#irrReserveNote"})
    sh.page.click("#irrWithdrawBtn")
    sh.page.fill("#irrReserveAmount", "2000")
    sh.page.fill("#irrReserveReason", "Сломался телефон")
    sh.page.select_option("#irrReserveCategory", label="🛍️ Покупки")
    d = dlg(sh, "irrReserveDialog")
    sh.shot("06-irregular-reserve-withdraw.png", [
        M(1, label(d, "Как использовать"), pad=3), M(2, label(d, "Сумма"), pad=3),
        M(3, label(d, "Категория расхода"), pad=3), M(4, label(d, "Причина"), pad=3),
        M(5, "#saveIrrReserveBtn", pad=3),
    ])
    sh.page.click("#saveIrrReserveBtn")
    sh.settle()
    sh.shot("06-irregular-reserve-after.png", [
        M(1, "#irrReserveBalance", pad=3), M(2, "#irrReserveMovements .list-row >> nth=0", pad=3),
    ], facts={"Резерв": "#irrReserveBalance", "История": "#irrReserveMovements .list-row >> nth=0"})
    sh.nav("home")
    sh.facts["06-irregular-reserve-after-home"] = {
        k: sh.text(v) for k, v in {"Лимит на день": "#dailyAvailable", "Свободных денег": "#irrFree",
                                   "Резерв": "#irrReserve"}.items()}
    sh.nav("buffer")
    sh.tab("obligations")
    sh.shot("06-irregular-obligations.png", [
        M(1, "#irrObligationsTotal", pad=3), M(2, "#irrObligationsList", pad=3),
    ], facts={"Отложено": "#irrObligationsTotal", "Платежи": "#irrObligationsList"})
    # Две недели без дохода
    app.load("irregular")
    simulate(app, PAYROLL_TODAY, date(2026, 10, 28), random.Random(28), irregular=True, overspend=False)
    sh.open()
    sh.shot("06-irregular-no-income.png", [
        M(1, "#irrNoIncome >> xpath=.."), M(2, "#irrNoIncomeWarning", pad=3),
        M(3, "#irrTakeReserveBtn", pad=3),
    ], facts={"Предупреждение": "#irrNoIncomeText", "Дней без дохода": "#irrNoIncome"})


# --- 7. Траты ------------------------------------------------------------------------

def ch07(app: App, sh: Shooter, texts: dict) -> None:
    app.load("payroll")
    payroll_today(app)
    sh.open()
    sh.shot("07-expense-open.png", [M(1, "#quickExpenseBtn", pad=3), M(2, "#navAddExpense", pad=2)],
            facts=HOME_FACTS)
    sh.page.click("#quickExpenseBtn")
    sh.page.fill("#expenseAmount", "450")
    d = dlg(sh, "expenseDialog")
    d.locator(".cat-tile", has_text="Кафе").click()
    sh.page.fill("#expenseNote", "Обед")
    sh.shot("07-expense-dialog.png", [
        M(1, label(d, "Сумма"), pad=3), M(2, d.locator(".cat-tile", has_text="Кафе"), pad=2),
        M(3, label(d, "Дата"), pad=3), M(4, label(d, "Комментарий"), pad=3),
        M(5, "#saveExpenseBtn", pad=3),
    ])
    sh.page.click("#saveExpenseBtn")
    sh.settle()
    sh.scroll_top()
    sh.shot("07-expense-result.png", [
        M(1, "#leftToday"), M(2, "#todayHero .hero-row span:last-child", pad=3),
    ], facts=HOME_FACTS)
    sh.nav("operations")
    first = sh.page.locator("#transactionsList .transaction-week").first
    sh.shot("07-operations-list.png", [
        M(1, ".page.active .segmented", pad=3), M(2, "#txSearch", pad=3), M(3, "#txChips", pad=3),
        M(4, first.locator(".transaction-week-toggle"), pad=2),
        M(5, first.locator(".actions .tiny:not(.danger)").first, pad=2),
        M(6, first.locator(".actions .tiny.danger").first, pad=2),
    ])
    sh.page.fill("#txSearch", "кафе")
    sh.page.wait_for_timeout(300)
    sh.shot("07-operations-search.png", [M(1, "#txSearch", pad=3)])
    sh.page.fill("#txSearch", "")
    sh.page.locator("#txSearch").dispatch_event("input")
    sh.page.wait_for_timeout(200)
    sh.page.locator("#transactionsList .list-row", has_text="Обед").first.locator(".actions .tiny:not(.danger)").click()
    sh.page.fill("#expenseAmount", "520")
    sh.shot("07-expense-edit.png", [M(1, label(d, "Сумма"), pad=3), M(2, "#saveExpenseBtn", pad=3)])
    d.locator("[data-close]").click()
    # Перерасход: кроссовки за 2 000 ₽ при остатке на сегодня 1 028,94 ₽.
    app.load("payroll-cafe")
    cats = category_ids(app)
    add_expense(app, cats, "Покупки", 2000, PAYROLL_TODAY, "Кроссовки")
    app.save("payroll-overspend")
    sh.open()
    sh.shot("07-overspend.png", [
        M(1, "#leftToday"), M(2, "#overspendAmount", pad=3),
        M(3, "#spreadOverspendBtn", pad=3), M(4, "#coverFromPiggyBtn", pad=3),
    ], facts={**HOME_FACTS, "Перерасход": "#overspendAmount"})
    texts["overspend"] = app.bot_texts()
    sh.page.click("#coverFromPiggyBtn")
    sh.page.fill("#piggyAmount", "971.06")
    d = dlg(sh, "piggyDialog")
    sh.shot("07-overspend-piggy.png", [M(1, label(d, "Сумма"), pad=3), M(2, "#savePiggyBtn", pad=3)])
    sh.page.click("#savePiggyBtn")
    sh.settle()
    sh.scroll_top()
    sh.shot("07-overspend-covered.png", [
        M(1, "#leftToday"), M(2, "#cardPiggy >> xpath=.."),
    ], facts=HOME_FACTS)
    app.load("payroll-overspend", PAYROLL_TODAY + timedelta(days=1))
    sh.open()
    sh.shot("07-overspend-next-day.png", [
        M(1, "#todayHero .hero-row span:first-child", pad=3), M(2, "#cardCash >> xpath=.."),
    ], facts=HOME_FACTS)


# --- 8. Обязательные платежи ----------------------------------------------------------

def ch08(app: App, sh: Shooter, texts: dict) -> None:
    # Новый платёж — на пустой базе, как при первой настройке.
    app.reset(PAYROLL_TODAY)
    sh.open()
    sh.nav("settings")
    bills = card(sh, "Обязательные платежи")
    sh.scroll(bills, "center")
    sh.page.click("#addBillRuleBtn")
    sh.page.fill("#ruleTitle", "Аренда")
    sh.page.fill("#ruleAmount", "25000")
    sh.page.fill("#ruleDay", "10")
    sh.page.select_option("#ruleBillCategory", label="🏠 Дом")
    d = dlg(sh, "ruleDialog")
    sh.shot("08-bills-dialog.png", [
        M(1, label(d, "Название"), pad=3), M(2, label(d, "Сумма"), pad=3),
        M(3, label(d, "День месяца"), pad=3), M(4, "#billCategoryWrap", pad=3),
        M(5, label(d, "Активно"), pad=3), M(6, label(d, "Применять условия с даты"), pad=3),
        M(7, "#saveRuleBtn", pad=3),
    ])
    d.locator("[data-close]").click()
    app.load("payroll")
    payroll_today(app)
    sh.open()
    sh.nav("settings")
    bills = card(sh, "Обязательные платежи")
    sh.scroll(bills, "center")
    row = sh.page.locator("#billRulesList .list-row").first
    sh.shot("08-bills-settings.png", [
        M(1, "#addBillRuleBtn", pad=3), M(2, row.locator(".row-text"), pad=3),
        M(3, row.locator(".actions .tiny:not(.danger)"), pad=2), M(4, row.locator(".actions .tiny.danger"), pad=2),
    ])
    sh.nav("buffer")
    sh.tab("plan")
    unpaid = sh.page.locator("#planList .bill-row:not(.paid)").first
    paid = sh.page.locator("#planList .bill-row.paid").first
    sh.shot("08-bills-plan.png", [
        M(1, paid.locator(".bill-paid-caption"), pad=3), M(2, unpaid.locator(".bill-payment-action .tiny"), pad=3),
    ], facts={"Платежи": "#planList"})
    unpaid.locator(".bill-payment-action .tiny").click()
    sh.settle()
    paid_now = sh.page.locator("#planList .bill-row", has_text="Интернет")
    sh.shot("08-bills-paid.png", [M(1, paid_now, pad=3)], facts={"Платежи": "#planList"})
    sh.nav("home")
    sh.shot("08-bills-paid-home.png", [M(1, "#leftToday"), M(2, "#cardBuffer >> xpath=..")],
            facts=HOME_FACTS)
    sh.nav("buffer")
    sh.tab("plan")
    sh.page.locator("#planList .bill-row", has_text="Интернет").locator(".tiny").click()
    sh.page.fill("#billPaymentAmount", "650")
    sh.page.select_option("#billRemainderDestination", "piggy")
    d = dlg(sh, "billPaymentDialog")
    sh.shot("08-bills-partial-dialog.png", [
        M(1, "#billPaymentPlanned", pad=3), M(2, label(d, "Оплачено фактически"), pad=3),
        M(3, label(d, "Куда направить остаток"), pad=3), M(4, d.locator("button.primary"), pad=3),
    ])
    d.locator("button.primary").click()
    sh.settle()
    sh.shot("08-bills-partial-result.png", [
        M(1, sh.page.locator("#planList .bill-row", has_text="Интернет").locator(".bill-paid-caption"), pad=3),
    ], facts={"Платежи": "#planList"})
    # Просроченный платёж: 16 октября, интернет так и не оплачен.
    app.load("payroll", date(2026, 10, 16))
    sh.open()
    sh.scroll(".card.home-bills", "center")
    sh.shot("08-bills-overdue.png", [M(1, "#homeBills .list-row", pad=3)], facts={"Скоро платежи": "#homeBills"})


# --- 9. Копилка ------------------------------------------------------------------------

def ch09(app: App, sh: Shooter, texts: dict) -> None:
    app.load("payroll")
    payroll_today(app)
    sh.open()
    sh.nav("buffer")
    sh.tab("piggy")
    sh.shot("09-piggy-page.png", [
        M(1, "#piggyBalance", pad=3), M(2, "#piggyDepositBtn", pad=3), M(3, "#piggyTransferBtn", pad=3),
        M(4, "#piggyToCardBtn", pad=3), M(5, "#piggyWithdrawBtn", pad=3),
        M(6, "#piggyMovements .list-row >> nth=0", pad=3),
    ], facts={"В копилке": "#piggyBalance"})
    sh.nav("home")
    sh.page.click("#quickPiggyBtn")
    sh.page.fill("#piggyAmount", "400")
    d = dlg(sh, "piggyDialog")
    sh.shot("09-piggy-transfer-day.png", [M(1, label(d, "Сумма"), pad=3), M(2, "#savePiggyBtn", pad=3)])
    sh.page.click("#savePiggyBtn")
    sh.settle()
    sh.scroll_top()
    sh.shot("09-piggy-result.png", [
        M(1, "#leftToday"), M(2, "#cardPiggy >> xpath=.."),
    ], facts=HOME_FACTS)
    sh.nav("buffer")
    sh.tab("piggy")
    sh.page.click("#piggyDepositBtn")
    sh.page.fill("#piggyAmount", "3000")
    sh.page.fill("#piggyNote", "Друг вернул долг")
    d = dlg(sh, "piggyDialog")
    sh.shot("09-piggy-deposit.png", [
        M(1, label(d, "Сумма"), pad=3), M(2, label(d, "Дата"), pad=3),
        M(3, label(d, "Комментарий"), pad=3), M(4, "#savePiggyBtn", pad=3),
    ])
    d.locator("[data-close]").click()
    sh.page.click("#piggyToCardBtn")
    sh.page.fill("#piggyAmount", "1000")
    sh.shot("09-piggy-to-card.png", [
        M(1, label(d, "Как учесть"), pad=3), M(2, label(d, "Сумма"), pad=3), M(3, "#savePiggyBtn", pad=3),
    ])
    sh.page.click("#savePiggyBtn")
    sh.settle()
    sh.nav("home")
    sh.shot("09-piggy-to-card-result.png", [
        M(1, "#leftToday"), M(2, "#todayHero .hero-row span:first-child", pad=3),
        M(3, "#cardPiggy >> xpath=.."),
    ], facts=HOME_FACTS)
    sh.nav("buffer")
    sh.tab("piggy")
    sh.page.click("#piggyWithdrawBtn")
    sh.page.fill("#piggyAmount", "2500")
    sh.page.fill("#piggyNote", "Подарок маме")
    sh.shot("09-piggy-withdraw.png", [M(1, label(d, "Сумма"), pad=3), M(2, "#savePiggyBtn", pad=3)])
    d.locator("[data-close]").click()


# --- 10. Категории -----------------------------------------------------------------

def ch10(app: App, sh: Shooter, texts: dict) -> None:
    app.load("payroll")
    sh.open()
    sh.nav("settings")
    cats = card(sh, "Категории")
    sh.scroll(cats, "center")
    item = sh.page.locator("#categoriesList .category-order-item", has_text="Кафе")
    sh.shot("10-categories-list.png", [
        M(1, "#addCategoryBtn", pad=3), M(2, item.locator(".category-order-name"), pad=2),
        M(3, item.locator(".category-drag-handle"), pad=2),
    ])
    item.locator(".category-order-name").click()
    d = dlg(sh, "categoryDialog")
    sh.shot("10-categories-edit.png", [
        M(1, label(d, "Эмодзи"), pad=3), M(2, label(d, "Название"), pad=3),
        M(3, "#saveCategoryBtn", pad=3), M(4, label(d, "Траты и платежи этой категории"), pad=3),
        M(5, "#deleteCategoryBtn", pad=3),
    ])
    d.locator("[data-close]").click()
    sh.page.click("#addCategoryBtn")
    sh.page.fill("#categoryEmoji", "🐾")
    sh.page.fill("#categoryTitle", "Питомец")
    sh.shot("10-categories-new.png", [
        M(1, label(d, "Эмодзи"), pad=3), M(2, label(d, "Название"), pad=3), M(3, "#saveCategoryBtn", pad=3),
    ])
    d.locator("[data-close]").click()
    # Перетаскивание «Кафе» на первое место
    sh.scroll(cats, "center")
    handle = item.locator(".category-drag-handle").bounding_box()
    first = sh.page.locator("#categoriesList .category-order-item").first.bounding_box()
    x, y = handle["x"] + handle["width"] / 2, handle["y"] + handle["height"] / 2
    sh.page.mouse.move(x, y)
    sh.page.mouse.down()
    for step in range(1, 9):
        sh.page.mouse.move(x, y + (first["y"] + 10 - y) * step / 10)
        sh.page.wait_for_timeout(30)
    sh.page.wait_for_timeout(250)
    sh.shot("10-categories-drag.png", [M(1, "#categoriesList .category-order-item.dragging", pad=2)])
    sh.page.mouse.up()
    sh.settle()
    # Удаление «Кафе» с переносом трат в «Продукты»: до и после в аналитике.
    app.load("payroll")
    sh.open()
    sh.nav("analytics")
    sh.page.wait_for_timeout(500)
    sh.facts["10-before-delete"] = {"Категории": sh.text("#analyticsCategories")}
    sh.nav("settings")
    sh.scroll(card(sh, "Категории"), "center")
    sh.page.locator("#categoriesList .category-order-item", has_text="Кафе").locator(".category-order-name").click()
    sh.page.select_option("#categoryMoveTo", label="Перенести в 🛒 Продукты")
    d = dlg(sh, "categoryDialog")
    sh.shot("10-categories-delete.png", [M(1, label(d, "Траты и платежи этой категории"), pad=3),
                                        M(2, "#deleteCategoryBtn", pad=3)])
    sh.page.click("#deleteCategoryBtn")
    sh.settle()
    sh.nav("analytics")
    sh.page.wait_for_timeout(500)
    sh.scroll(".analytics-cats-card", "start")
    sh.shot("10-categories-after-delete.png", [M(1, "#analyticsCategories .analytics-cat >> nth=0", pad=3)],
            facts={"Категории": "#analyticsCategories"})


# --- 11. Аналитика ------------------------------------------------------------------

def ch11(app: App, sh: Shooter, texts: dict) -> None:
    app.load("payroll")
    payroll_today(app)
    sh.open()
    sh.nav("analytics")
    sh.page.wait_for_timeout(500)
    sh.shot("11-analytics-total.png", [
        M(1, ".page.active .analytics-modes", pad=3), M(2, ".analytics-nav", pad=3),
        M(3, ".analytics-total-row", pad=3), M(4, "#analyticsKinds", pad=3),
    ], facts={"Подпись": "#analyticsTotalCaption", "Всего": "#analyticsTotal", "Сравнение": "#analyticsDelta",
              "Виды": "#analyticsKinds", "Месяц": "#analyticsLabel"})
    sh.page.locator("#analyticsCategories .analytics-cat").first.click()
    sh.page.wait_for_timeout(200)
    # The list of operations is next to the category's button, not inside it.
    first = sh.page.locator("#analyticsCategories .analytics-cat-item").first
    sh.scroll(".analytics-cats-card", "start")
    sh.shot("11-analytics-categories.png", [
        M(1, first.locator(".top"), pad=3), M(2, first.locator(".meta"), pad=3),
        M(3, first.locator(".analytics-drill > span").first, pad=3), M(4, first.locator(".analytics-sort"), pad=3),
    ], facts={"Категории": "#analyticsCategories"})
    sh.scroll("#analyticsDaysCard", "center")
    sh.shot("11-analytics-days.png", [M(1, "#analyticsDays", pad=3), M(2, "#analyticsDaysLegend", pad=3)],
            facts={"Легенда": "#analyticsDaysLegend"})
    sh.scroll(".analytics-bills-card", "center")
    sh.shot("11-analytics-bills.png", [M(1, "#analyticsBillsTotal", pad=3), M(2, "#analyticsBills", pad=3)],
            facts={"Платежи": ".analytics-bills-card"})
    sh.scroll_top()
    sh.page.click('.page.active [data-amode="period"]')
    sh.settle()
    sh.shot("11-analytics-period.png", [M(1, '.page.active [data-amode="period"]', pad=2),
                                        M(2, "#analyticsLabel", pad=3)],
            facts={"Период": "#analyticsLabel", "Всего": "#analyticsTotal"})
    sh.scroll("#analyticsDaysCard", "center")
    sh.shot("11-analytics-days-limit.png", [M(1, "#analyticsDays", pad=3), M(2, "#analyticsDaysLegend", pad=3)],
            facts={"Легенда": "#analyticsDaysLegend"})
    sh.scroll_top()
    sh.page.click('.page.active [data-amode="year"]')
    sh.settle()
    sh.scroll("#analyticsMonthsCard", "center")
    sh.shot("11-analytics-year.png", [M(1, "#analyticsMonths", pad=3), M(2, "#analyticsMonthsNote", pad=3)],
            facts={"По месяцам": "#analyticsMonthsNote"})
    app.load("irregular")
    sh.open()
    sh.nav("analytics")
    sh.page.wait_for_timeout(500)
    sh.scroll("#irrAnalytics", "start")
    sh.shot("11-analytics-income.png", [
        M(1, "#irrAnalyticsTotal", pad=3), M(2, "#irrAnalyticsStats", pad=3),
        M(3, "#irrAnalyticsWeeks", pad=3),
    ], facts={"Доходы": "#irrAnalytics"})


# --- 12. Карты и счета -------------------------------------------------------------

def ch12(app: App, sh: Shooter, texts: dict) -> None:
    # Мастер включения — заполняем через интерфейс те же карты, что и в демо-данных.
    app.load("payroll")
    payroll_today(app)
    sh.open()
    sh.nav("settings")
    acc = card(sh, "Карты и счета")
    sh.scroll(acc, "start")
    sh.shot("12-accounts-enable.png", [M(1, label(acc, "Учёт по картам и счетам"), pad=3)])
    sh.page.click("#accountsEnabled")
    sh.page.wait_for_selector("#accountsWizard[open]")
    sh.page.wait_for_timeout(300)
    for i, row in enumerate(ACCOUNT_ROWS):
        if i:
            sh.page.click("#wizAddRow")
        sh.page.fill(f"#wizEmoji-{i}", row["emoji"])
        sh.page.fill(f"#wizTitle-{i}", row["title"])
        if row["kind"] != "debit":
            sh.page.select_option(f"#wizKind-{i}", row["kind"])
        if row["kind"] == "credit":
            sh.page.fill(f"#wizLimit-{i}", "50000")
            sh.page.fill(f"#wizDay-{i}", "20")
    w = dlg(sh, "accountsWizard")
    sh.scroll("#accountsWizardNext", "end")
    sh.shot("12-accounts-wizard-1.png", [
        M(1, w.locator(".wizard-account").nth(3).locator("label").filter(has_text="Тип"), pad=3),
        M(2, w.locator(".wizard-account").nth(3).locator(".grid2"), pad=3),
        M(3, "#wizAddRow", pad=3), M(4, "#accountsWizardNext", pad=3),
    ])
    sh.page.click("#accountsWizardNext")
    sh.page.select_option("#wizEnvelope-buffer", "2")
    sh.page.select_option("#wizEnvelope-piggy", "2")
    sh.shot("12-accounts-wizard-2.png", [
        M(1, label(w, "Основной счёт"), pad=3), M(2, w.locator("#accountsWizardBody label").nth(1), pad=3),
        M(3, w.locator("#accountsWizardBody label").nth(2), pad=3),
        M(4, w.locator("#accountsWizardBody label").nth(3), pad=3), M(5, "#accountsWizardNext", pad=3),
    ])
    sh.page.click("#accountsWizardNext")
    sh.page.fill("#wizBalance-1", str(CARD2_BALANCE))
    sh.shot("12-accounts-wizard-3.png", [
        M(1, w.locator("#accountsWizardBody label").nth(0), pad=3),
        M(2, w.locator("#accountsWizardBody label").nth(1), pad=3),
        M(3, w.locator("#accountsWizardBody label").nth(2), pad=3), M(4, "#accountsWizardNext", pad=3),
    ], facts={"Остатки": "#accountsWizardBody"})
    sh.page.click("#accountsWizardNext")
    sh.settle()

    # Дальше — снимок с теми же картами, созданный через API.
    app.load("accounts")
    texts["accounts"] = app.bot_texts()
    sh.open()
    sh.scroll("#homeAccounts", "start")
    rows = sh.page.locator("#accList .acc-item")
    sh.shot("12-accounts-home.png", [
        M(1, "#homeAccounts .acc-total", pad=3), M(2, rows.nth(0), pad=3), M(3, rows.nth(1), pad=3),
        M(4, rows.nth(2), pad=3), M(5, "#accCredit .acc-item", pad=3),
    ], facts={"Карты и счета": "#homeAccounts", **HOME_FACTS})
    sh.nav("settings")
    acc = card(sh, "Карты и счета")
    sh.scroll(acc, "start")
    first = sh.page.locator("#accountsList .account-item").first
    sh.shot("12-accounts-settings.png", [
        M(1, "#addAccountBtn", pad=3), M(2, first.locator(".category-order-name"), pad=2),
        M(3, first.locator(".category-drag-handle"), pad=2),
    ])
    sh.scroll("#envelopePlacesBlock", "center")
    sh.shot("12-accounts-envelopes.png", [
        M(1, "#envelopePlaces label >> nth=0", pad=3), M(2, "#envelopePlaces label >> nth=1", pad=3),
        M(3, "#envelopePlaces label >> nth=2", pad=3),
    ])
    # Трата с выбором карты
    sh.nav("home")
    sh.page.click("#quickExpenseBtn")
    sh.page.fill("#expenseAmount", "1200")
    d = dlg(sh, "expenseDialog")
    d.locator(".cat-tile", has_text="Продукты").click()
    sh.scroll("#expenseAccountField", "center")
    sh.shot("12-accounts-expense.png", [
        M(1, "#expenseAccountField", pad=3), M(2, "#expenseFundingField", pad=3),
    ])
    d.locator("[data-close]").click()
    # Перевод 1 000 ₽ с «Карты 2» на «Карту 1»
    sh.scroll("#homeAccounts", "center")
    ids = account_ids(app)
    sh.page.locator("#accList .acc-row", has_text="Карта 2").click()
    sh.page.wait_for_selector("#accountHistoryDialog[open]")
    sh.page.click('#accountHistoryActions [data-history-action="transfer"]')
    sh.page.select_option("#transferTo", str(ids["Карта 1"]))
    sh.page.fill("#transferAmount", "1000")
    t = dlg(sh, "transferDialog")
    sh.shot("12-accounts-transfer.png", [
        M(1, label(t, "Откуда"), pad=3), M(2, label(t, "Куда"), pad=3),
        M(3, label(t, "Сумма"), pad=3), M(4, "#saveTransferBtn", pad=3),
    ])
    sh.page.click("#saveTransferBtn")
    sh.settle()
    dlg(sh, "accountHistoryDialog").locator("[data-close]").click()
    sh.scroll("#homeAccounts", "start")
    sh.shot("12-accounts-after-transfer.png", [
        M(1, "#homeAccounts .acc-total", pad=3), M(2, "#accList .acc-item >> nth=0", pad=3),
        M(3, "#accList .acc-item >> nth=1", pad=3),
    ], facts={"Карты и счета": "#homeAccounts"})
    # Окно карты: остаток, кнопки и история
    sh.page.locator("#accList .acc-row").first.click()
    sh.page.wait_for_selector("#accountHistoryDialog[open]")
    sh.page.wait_for_timeout(300)
    sh.shot("12-accounts-history.png", [
        M(1, "#accountHistoryDialog .acc-total", pad=3),
        M(2, '#accountHistoryActions [data-history-action="transfer"]', pad=2),
        M(3, '#accountHistoryActions [data-history-action="expense"]', pad=2),
        M(4, '#accountHistoryActions [data-history-action="adjust"]', pad=2),
        M(5, "#accountHistoryList .list-row >> nth=0", pad=3),
    ], facts={"История": "#accountHistoryList", "Остаток": "#accountHistoryBalance"})
    dlg(sh, "accountHistoryDialog").locator("[data-close]").click()

    # Проценты по накопительному счёту: «Исправить остаток» → вне бюджета → в копилку.
    app.load("accounts")
    sh.open()
    sh.scroll("#homeAccounts", "center")
    sh.page.locator("#accList .acc-row", has_text="Накопительный").click()
    sh.page.wait_for_selector("#accountHistoryDialog[open]")
    sh.page.click('#accountHistoryActions [data-history-action="adjust"]')
    target = sh.page.input_value("#adjustTarget")
    savings = next(a for a in app.get("/api/accounts")["accounts"] if a["title"] == "Накопительный")
    sh.page.fill("#adjustTarget", f"{savings['balance'] + 250:.2f}".replace(".", ","))
    sh.page.fill("#adjustNote", "Проценты по счёту")
    a = dlg(sh, "adjustDialog")
    sh.shot("12-accounts-adjust.png", [
        M(1, "#adjustCaption", pad=3), M(2, label(a, "Сколько на счёте на самом деле"), pad=3),
        M(3, label(a, "Комментарий"), pad=3), M(4, "#saveAdjustBtn", pad=3),
    ])
    sh.facts["12-adjust-before"] = {"Было": target}
    sh.page.click("#saveAdjustBtn")
    sh.settle()
    dlg(sh, "accountHistoryDialog").locator("[data-close]").click()
    sav = sh.page.locator("#accList .acc-item", has_text="Накопительный")
    sh.scroll(sav, "center")
    sh.shot("12-accounts-outside.png", [M(1, sav.locator(".acc-diff"), pad=3),
                                        M(2, sav.locator("[data-add-to]"), pad=2)],
            facts={"Накопительный": sav})
    sav.locator("[data-add-to]").click()
    sh.settle()
    sh.scroll_top()
    sh.shot("15-interest-result.png", [M(1, "#cardPiggy >> xpath=..")], facts=HOME_FACTS)
    sav = sh.page.locator("#accList .acc-item", has_text="Накопительный")
    sh.facts["15-interest-savings"] = {"Накопительный": sh.text(sav)}

    # «Не хватает»: в банке на «Карте 1» на 300 ₽ меньше, чем в приложении.
    app.load("accounts")
    card1 = next(a for a in app.get("/api/accounts")["accounts"] if a["title"] == "Карта 1")
    app.post(f"/api/accounts/{card1['id']}/adjust", {"target": round(card1["balance"] - 300, 2),
                                                     "note": "Сверка с банком"})
    sh.open()
    one = sh.page.locator("#accList .acc-item", has_text="Карта 1")
    sh.scroll(one, "center")
    sh.shot("12-accounts-missing.png", [
        M(1, one.locator(".acc-diff span"), pad=3), M(2, one.locator("[data-transfer-to]"), pad=2),
        M(3, one.locator("[data-adjust-account]"), pad=2),
    ], facts={"Карта 1": one})

    # Кредитка: покупка, плашка долга, погашение, план погашения, проценты.
    app.load("accounts-credit")
    texts["accounts-credit"] = app.bot_texts()
    sh.open()
    sh.scroll("#homeAccounts", "start")
    sh.shot("12-credit-banner.png", [
        M(1, "#accDebtText", pad=3), M(2, "#accDebtBtn", pad=2), M(3, "#accCredit .acc-item", pad=3),
    ], facts={"Карты и счета": "#homeAccounts"})
    sh.scroll_top()
    sh.shot("12-credit-home-numbers.png", [M(1, "#leftToday"), M(2, "#cardCash >> xpath=..")],
            facts=HOME_FACTS)
    sh.scroll("#homeAccounts", "center")
    sh.page.click("#accDebtBtn")
    r = dlg(sh, "repayDialog")
    sh.shot("12-credit-repay.png", [
        M(1, "#repayCaption", pad=3), M(2, label(r, "С какой карты"), pad=3), M(3, label(r, "Сумма"), pad=3),
        M(4, label(r, "Из чего"), pad=3), M(5, label(r, "Из них проценты или комиссия банка"), pad=3),
        M(6, "#saveRepayBtn", pad=3),
    ])
    sh.page.click("#saveRepayBtn")
    sh.settle()
    sh.scroll_top()
    sh.shot("12-credit-repay-result.png", [
        M(1, "#leftToday"), M(2, "#overspendPanel", pad=3),
    ], facts={**HOME_FACTS, "Перерасход": "#overspendAmount"})
    sh.scroll("#homeAccounts", "start")
    sh.shot("12-credit-repay-accounts.png", [
        M(1, "#accList .acc-item >> nth=0", pad=3), M(2, "#accCredit .acc-item", pad=3),
    ], facts={"Карты и счета": "#homeAccounts"})
    app.load("accounts-credit")
    sh.open()
    sh.scroll("#homeAccounts", "center")
    sh.page.locator("#accCredit .acc-row").click()
    sh.page.wait_for_selector("#accountHistoryDialog[open]")
    sh.page.wait_for_timeout(300)
    sh.shot("12-credit-history.png", [
        M(1, '#accountHistoryActions [data-history-action="repay"]', pad=2),
        M(2, '#accountHistoryActions [data-history-action="schedule"]', pad=2),
        M(3, '#accountHistoryActions [data-history-action="adjust"]', pad=2),
    ])
    sh.page.click('#accountHistoryActions [data-history-action="schedule"]')
    s = dlg(sh, "scheduleDialog")
    sh.shot("12-credit-schedule.png", [
        M(1, "#scheduleCaption", pad=3), M(2, label(s, "Дата платежа"), pad=3),
        M(3, label(s, "Сумма"), pad=3), M(4, "#saveScheduleBtn", pad=3),
    ])
    sh.page.click("#saveScheduleBtn")
    sh.settle()
    dlg(sh, "accountHistoryDialog").locator("[data-close]").click()
    sh.scroll("#homeAccounts", "center")
    sh.shot("12-credit-scheduled.png", [M(1, "#accCredit .acc-split", pad=3)],
            facts={"Кредитка": "#accCredit"})
    sh.nav("buffer")
    sh.tab("plan")
    sh.shot("12-credit-scheduled-bill.png", [M(1, sh.page.locator("#planList .bill-row", has_text="Кредитка"), pad=3)],
            facts={"Платежи": "#planList"})
    # Запланированное погашение оплачено 20 октября, а по выписке банка долг ещё 150 ₽.
    app.save("accounts-scheduled")
    ids = account_ids(app)
    simulate(app, PAYROLL_TODAY, date(2026, 10, 20), random.Random(20), overspend=False,
             extra={"account_id": ids["Карта 1"], "funding_envelope": "spending"})
    sh.open()
    sh.scroll("#homeAccounts", "start")
    sh.shot("12-credit-paid.png", [M(1, "#accCredit .acc-item", pad=3)], facts={"Карты и счета": "#homeAccounts"})
    sh.scroll("#homeAccounts", "center")
    sh.page.locator("#accCredit .acc-row").click()
    sh.page.wait_for_selector("#accountHistoryDialog[open]")
    sh.page.click('#accountHistoryActions [data-history-action="adjust"]')
    sh.page.fill("#adjustTarget", "150")
    sh.page.fill("#adjustNote", "Проценты по выписке")
    a = dlg(sh, "adjustDialog")
    sh.shot("12-credit-adjust.png", [
        M(1, "#adjustCaption", pad=3), M(2, label(a, "Долг по выписке банка"), pad=3),
        M(3, "#adjustHint", pad=3), M(4, "#saveAdjustBtn", pad=3),
    ])
    sh.page.click("#saveAdjustBtn")
    sh.settle()
    dlg(sh, "accountHistoryDialog").locator("[data-close]").click()
    sh.scroll("#homeAccounts", "start")
    sh.shot("15-credit-interest.png", [M(1, "#accDebtText", pad=3), M(2, "#accCredit .acc-item", pad=3)],
            facts={"Карты и счета": "#homeAccounts"})

    # День выплаты с картами: «Получил ✓» спрашивает карту, потом «Нужно перевести».
    app.load("accounts")
    ids = account_ids(app)
    simulate(app, PAYROLL_TODAY, PAYDAY, random.Random(22), confirm=False, overspend=False,
             extra={"account_id": ids["Карта 1"], "funding_envelope": "spending"})
    sh.open()
    sh.nav("buffer")
    cur = sh.page.locator("#planCards article.pc.cur")
    cur.locator(".income-btn.accent").click()
    d = dlg(sh, "cashflowIncomeDialog")
    sh.shot("12-accounts-payday-confirm.png", [
        M(1, label(d, "Сколько пришло"), pad=3), M(2, "#cashflowIncomeAccountField", pad=3),
        M(3, "#saveCashflowIncomeBtn", pad=3),
    ])
    sh.page.click("#saveCashflowIncomeBtn")
    sh.settle()
    sh.nav("home")
    sh.scroll("#homeAccounts", "start")
    sh.page.locator("#accHints summary").click()
    sh.page.wait_for_timeout(200)
    sh.shot("12-accounts-hints.png", [
        M(1, "#accHints summary", pad=3), M(2, "#accHints .acc-hint >> nth=0 >> span", pad=3),
        M(3, "#accHints .acc-hint >> nth=0 >> button", pad=2),
    ], facts={"Карты и счета": "#homeAccounts"})
    sh.page.locator("#accHints .acc-hint button").first.click()
    sh.settle()
    sh.scroll("#homeAccounts", "start")
    sh.shot("12-accounts-hints-done.png", [M(1, "#homeAccounts .acc-list", pad=3)],
            facts={"Карты и счета": "#homeAccounts"})

    # Архив: «Карта 2» закрыта в банке.  Сначала деньги переводятся, потом карта уходит в архив.
    app.load("accounts")
    ids = account_ids(app)
    app.post("/api/accounts/transfers", {"from_account_id": ids["Карта 2"], "to_account_id": ids["Карта 1"],
                                         "amount": CARD2_BALANCE, "note": "Закрываю карту"})
    sh.open()
    sh.nav("settings")
    acc = card(sh, "Карты и счета")
    sh.scroll(acc, "center")
    sh.page.locator("#accountsList .account-item", has_text="Карта 2").locator(".category-order-name").click()
    a = dlg(sh, "accountDialog")
    sh.shot("12-accounts-archive-button.png", [
        M(1, "#adjustAccountBtn", pad=3), M(2, "#archiveAccountBtn", pad=3), M(3, "#deleteAccountBtn", pad=3),
    ])
    sh.page.click("#archiveAccountBtn")
    sh.settle()
    sh.page.locator("#accountsArchive summary").click()
    sh.scroll("#accountsArchive", "center")
    sh.shot("12-accounts-archive.png", [
        M(1, "#accountsArchive summary", pad=3), M(2, "[data-restore-account]", pad=2),
    ])


# --- 13. Бот --------------------------------------------------------------------------

def ch13(app: App, sh: Shooter, texts: dict) -> None:
    app.load("payroll")
    texts["payroll-morning"] = app.bot_texts()   # утро: сегодня ещё ничего не потрачено
    app.load("payroll")
    payroll_today(app)
    texts["payroll"] = app.bot_texts()
    sh.open()
    sh.nav("settings")
    note = card(sh, "Уведомления в Telegram")
    sh.scroll(note, "center")
    sh.shot("13-notifications-settings.png", [
        M(1, label(note, "Утренний отчёт"), pad=3), M(2, "#saveReminderSettingsBtn", pad=3),
    ])
    # Напоминание о кредитке: за три дня до дня платежа (20 октября).
    app.load("accounts-credit", date(2026, 10, 17))
    texts["credit-reminder"] = app.bot_texts()


# --- 14. Резервная копия и Excel ------------------------------------------------------

def ch14(app: App, sh: Shooter, texts: dict) -> None:
    app.load("accounts")
    sh.open()
    sh.nav("settings")
    backup = card(sh, "Резервная копия")
    sh.scroll(backup, "center")
    sh.page.evaluate("window.scrollBy(0, 120)")
    sh.shot("14-backup-card.png", [
        M(1, "#exportBackupBtn", pad=3), M(2, "#importBackupBtn", pad=3), M(3, "#exportExcelBtn", pad=3),
    ])
    # Что лежит в файле Excel: листы и заголовки колонок.
    from openpyxl import load_workbook
    import io

    raw = urllib.request.urlopen(f"{app.url}api/export/xlsx", timeout=60).read()
    book = load_workbook(io.BytesIO(raw), read_only=True)
    sheets = {}
    for ws in book.worksheets:
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        sheets[ws.title] = [str(v) for v in first if v is not None]
    sh.facts["14-excel-sheets"] = sheets
    backup_json = app.get("/api/backup")
    sh.facts["14-backup"] = {"backup_version": backup_json.get("backup_version"),
                             "tables": sorted((backup_json.get("data") or {}).keys())}


# --- 15. Типичные ситуации --------------------------------------------------------------

def ch15(app: App, sh: Shooter, texts: dict) -> None:
    # Забыли записать вчерашнюю трату.
    app.load("payroll")
    payroll_today(app)
    sh.open()
    sh.facts["15-forgot-before"] = {k: sh.text(v) for k, v in HOME_FACTS.items()}
    sh.page.click("#quickExpenseBtn")
    sh.page.fill("#expenseAmount", "350")
    d = dlg(sh, "expenseDialog")
    d.locator(".cat-tile", has_text="Транспорт").click()
    sh.page.fill("#expenseDate", "2026-10-13")
    sh.page.fill("#expenseNote", "Такси")
    sh.shot("15-forgot-yesterday.png", [M(1, label(d, "Дата"), pad=3), M(2, "#saveExpenseBtn", pad=3)])
    sh.page.click("#saveExpenseBtn")
    sh.settle()
    sh.scroll_top()
    sh.shot("15-forgot-result.png", [
        M(1, "#leftToday"), M(2, "#todayHero .hero-row span:first-child", pad=3),
        M(3, "#cardCash >> xpath=.."),
    ], facts=HOME_FACTS)
    # Заплатили не с той карты: в трате меняется «С какой карты».
    app.load("accounts")
    ids = account_ids(app)
    cats = category_ids(app)
    add_expense(app, cats, "Кафе", 380, PAYROLL_TODAY, "Кофе с собой", account_id=ids["Карта 1"],
                funding_envelope="spending")
    sh.open()
    sh.nav("operations")
    sh.page.locator("#transactionsList .list-row", has_text="Кофе с собой").first.locator(".actions .tiny:not(.danger)").click()
    sh.page.select_option("#expenseAccount", str(ids["Карта 2"]))
    sh.page.locator("#expenseAccount").dispatch_event("change")
    # У «Карты 2» нет конвертов, и «Из чего» само становится «Вне бюджета».
    # Обычная трата должна остаться в бюджете на траты.
    sh.page.select_option("#expenseFunding", "spending")
    d = dlg(sh, "expenseDialog")
    sh.scroll("#expenseAccountField", "center")
    sh.shot("15-wrong-card.png", [M(1, "#expenseAccountField", pad=3), M(2, "#expenseFundingField", pad=3),
                                  M(3, "#saveExpenseBtn", pad=3)])
    sh.page.click("#saveExpenseBtn")
    sh.settle()
    sh.nav("home")
    sh.scroll("#homeAccounts", "start")
    sh.page.locator("#accHints summary").click()
    sh.page.wait_for_timeout(200)
    sh.shot("15-wrong-card-result.png", [M(1, "#accList .acc-item >> nth=0", pad=3),
                                         M(2, "#accList .acc-item >> nth=1", pad=3),
                                         M(3, "#accHints .acc-hint", pad=3)],
            facts={"Карты и счета": "#homeAccounts"})
    # Перешли с зарплаты на подработки.
    app.load("payroll")
    sh.open()
    sh.nav("settings")
    sh.page.click('[data-budget-mode="irregular"]')
    sh.page.wait_for_timeout(300)
    irr = card(sh, "Подработки")
    sh.page.fill("#irrStartTotal", "20000")
    sh.page.fill("#irrStartReserve", "3000")
    sh.scroll("#saveIrregularBtn", "center")
    sh.shot("15-switch-mode.png", [
        M(1, label(irr, "Дата старта"), pad=3), M(2, label(irr, "Денег сейчас, без копилки"), pad=3),
        M(3, label(irr, "Из них резерв"), pad=3), M(4, "#saveIrregularBtn", pad=3),
    ])
    sh.page.click("#saveIrregularBtn")
    sh.settle()
    sh.nav("home")
    sh.shot("15-switch-mode-result.png", [M(1, "#irrHome")],
            facts={**HOME_FACTS, "Свободных денег": "#irrFree", "Растягиваем": "#irrStretch", "Резерв": "#irrReserve",
                   "Отложено на обязательные": "#irrBills"})


# ---------------------------------------------------------------------------
# Запуск
# ---------------------------------------------------------------------------

# Слова в «ёлочках», которых нет на экранах дословно: названия в другом падеже
# и кнопка самого Telegram.
INFLECTED = {
    "Настройках", "Настроек", "Операциях", "Буфере", "Лимита на день", "Премию до НДФЛ", "Продуктах",
    "Карте 1", "Карту 1", "Карты 1", "Картой 2", "Карты 2", "Карте 2", "Накопительном", "Накопительного",
    "Кредитке", "Запустить",
}


def normalize(text: str) -> str:
    import re

    return re.sub(r"\s+", " ", text.replace("\u00a0", " ").replace("\u202f", " ")).strip()


def check_guide() -> int:
    """Ссылки, картинки, подписи и названия в «ёлочках» в docs/guide.  Названия
    ищутся в тексте экранов, снятых скриптом, в сообщениях бота и в исходниках."""
    import re

    guide = IMAGES.parent
    sources = [path.read_text(encoding="utf-8") for pattern in
               ("app/static/*.html", "app/static/*.js", "app/*.py", "bot.py") for path in ROOT.glob(pattern)]
    # Сообщения из кода с подставленными пустыми значениями: «… копии.${warning} Продолжить?».
    sources += [re.sub(r"\$\{[^}]*\}", "", text) for text in sources]
    for name in ("screen-texts.txt", "bot-messages.md"):
        if (GENERATED / name).exists():
            sources.append((GENERATED / name).read_text(encoding="utf-8"))
    if (GENERATED / "numbers.json").exists():
        sources.append(json.dumps(json.loads((GENERATED / "numbers.json").read_text(encoding="utf-8")),
                                  ensure_ascii=False))
    titles = [re.sub(r"^#\s*\d+\.\s*", "", page.read_text(encoding="utf-8").splitlines()[0])
              for page in guide.glob("[0-9][0-9]-*.md")]
    corpus = normalize("\n".join(sources + titles))
    problems: list[str] = []
    used: set[str] = set()
    for page in sorted(guide.glob("*.md")):
        text = page.read_text(encoding="utf-8")
        lines = text.splitlines()
        for target in re.findall(r"\]\(([^)\s]+)\)", text):
            if target.startswith(("http://", "https://", "#")):
                continue
            if not (guide / target.split("#")[0]).exists():
                problems.append(f"{page.name}: нет файла {target}")
            if target.startswith("images/"):
                used.add(target[len("images/"):])
        for number, line in enumerate(lines):
            image = re.match(r"!\[(.*)\]\((images/[^)]+)\)", line)
            if not image:
                continue
            if not image.group(1).strip():
                problems.append(f"{page.name}:{number + 1}: у картинки нет alt-текста")
            following = next((l for l in lines[number + 1:] if l.strip()), "")
            if not following.startswith("*"):
                problems.append(f"{page.name}:{number + 1}: под картинкой нет подписи")
        if page.name == "HOWTO-UPDATE.md":
            continue
        for quoted in re.findall(r"«([^«»]+)»", text):
            if quoted in INFLECTED:
                continue
            # «…» в цитате — место для числа или названия.
            parts = [normalize(part).strip(" .,:") for part in quoted.split("…")]
            if not all(part in corpus for part in parts if part):
                problems.append(f"{page.name}: «{quoted}» нет ни на одном экране")
    for unused in sorted({p.name for p in IMAGES.glob("*.png")} - used):
        problems.append(f"картинка не используется: images/{unused}")
    for problem in problems:
        print(problem)
    print("Проверка пройдена" if not problems else f"Замечаний: {len(problems)}")
    return 1 if problems else 0


def main() -> None:
    if len(sys.argv) > 1 and sys.argv[1] == "--serve":
        return serve(int(sys.argv[2]), Path(sys.argv[3]))
    if len(sys.argv) > 1 and sys.argv[1] == "--bot-texts":
        return bot_texts(Path(sys.argv[2]))
    parser = argparse.ArgumentParser(description="Пересоздать скриншоты инструкции docs/guide/images")
    parser.add_argument("--only", help="снять только главы с этими номерами через запятую, например 3,5")
    parser.add_argument("--keep", action="store_true", help="не удалять временную папку (для отладки)")
    parser.add_argument("--check", action="store_true", help="только проверить ссылки, картинки и названия в тексте")
    args = parser.parse_args()
    if args.check:
        sys.exit(check_guide())
    run(args)


CHAPTERS = {1: ch01, 2: ch02, 3: ch03, 4: ch04, 5: ch05, 6: ch06, 7: ch07, 8: ch08, 9: ch09,
            10: ch10, 11: ch11, 12: ch12, 13: ch13, 14: ch14, 15: ch15}


def run(args) -> None:
    from playwright.sync_api import sync_playwright

    only = {int(x) for x in args.only.split(",")} if args.only else set(CHAPTERS)
    IMAGES.mkdir(parents=True, exist_ok=True)
    GENERATED.mkdir(parents=True, exist_ok=True)
    if not args.only:
        # Полное пересоздание: старые картинки удаляются, чтобы не осталось лишних.
        for old in IMAGES.glob("*.png"):
            old.unlink()
    workdir = Path(tempfile.mkdtemp(prefix="guide-shots-"))
    app = App(workdir)
    texts: dict = {}
    try:
        print("Демо-данные: зарплата и аванс, подработки, карты и счета")
        seed_payroll(app)
        seed_irregular(app)
        seed_accounts(app)
        with sync_playwright() as p:
            # --lang: поля даты и времени в русском формате, как на телефоне.
            launch = {"args": ["--lang=ru-RU"], "env": {**os.environ, "LANG": "ru_RU.UTF-8"}}
            if os.environ.get("GUIDE_CHROMIUM"):
                launch["executable_path"] = os.environ["GUIDE_CHROMIUM"]
            browser = p.chromium.launch(**launch)
            sh = Shooter(browser, app)
            for number in sorted(only):
                print(f"Глава {number}")
                CHAPTERS[number](app, sh, texts)
            sh.close()
            browser.close()
        write_generated(sh.facts, texts, sh.screen_texts, merge=bool(args.only))
        print(f"Готово: {len(sh.names)} картинок в {IMAGES.relative_to(ROOT)}")
    finally:
        app.stop()
        if args.keep:
            print("Временная папка:", workdir)
        else:
            shutil.rmtree(workdir, ignore_errors=True)


def write_generated(facts: dict, texts: dict, screen_texts: set, merge: bool) -> None:
    numbers_file = GENERATED / "numbers.json"
    bot_file = GENERATED / "bot-texts.json"
    screen_file = GENERATED / "screen-texts.txt"
    if merge and screen_file.exists():
        screen_texts |= set(screen_file.read_text(encoding="utf-8").splitlines())
    screen_file.write_text("\n".join(sorted(screen_texts)) + "\n", encoding="utf-8")
    if merge:
        for path, fresh in ((numbers_file, facts), (bot_file, texts)):
            if path.exists():
                old = json.loads(path.read_text())
                old.update(fresh)
                fresh.clear()
                fresh.update(old)
    numbers_file.write_text(json.dumps(facts, ensure_ascii=False, indent=2) + "\n")
    bot_file.write_text(json.dumps(texts, ensure_ascii=False, indent=2) + "\n")
    lines = ["# Сообщения бота на демо-данных", "",
             "Файл создаёт scripts/make_guide_screenshots.py. Тексты собраны функциями bot.py",
             "на тех же демо-данных, что и скриншоты: из них взяты цитаты в главах.", ""]
    for key, value in texts.items():
        for kind, text in value.items():
            lines += [f"## {key} · {kind}", "", "```text", text, "```", ""]
    (GENERATED / "bot-messages.md").write_text("\n".join(lines))


if __name__ == "__main__":
    main()
