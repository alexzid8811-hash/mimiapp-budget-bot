import hashlib
import hmac
import io
import json
import sqlite3
import time
from urllib.parse import urlencode

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import clock
from app.backup import export_user_data, restore_user_data
from app.db import (
    all_user_ids, connect, db_path, ensure_user, init_db, set_current_user, user_db_path, user_scope,
)
from app.excel_export import build_workbook

SECRET_A = "СЕКРЕТ-ПОЛЬЗОВАТЕЛЯ-А"
SECRET_B = "СЕКРЕТ-ПОЛЬЗОВАТЕЛЯ-Б"
TOKEN = "123456:TEST"


def _seed(user_id: int, secret: str) -> None:
    ensure_user(user_id, f"user{user_id}")
    with connect(user_id) as con:
        category_id = con.execute(
            "INSERT INTO categories(user_id,title,emoji) VALUES(?,?,'🧰')", (user_id, f"Кат {secret}")
        ).lastrowid
        con.execute(
            "INSERT INTO bill_rules(user_id,title,amount,day_of_month,category_id) VALUES(?,?,500,10,?)",
            (user_id, f"Платёж {secret}", category_id),
        )
        con.execute(
            "INSERT INTO income_rules(user_id,title,amount,day_of_month,kind) VALUES(?,?,900,5,'other')",
            (user_id, f"Доход {secret}"),
        )
        con.execute(
            "INSERT INTO transactions(user_id,type,amount,tx_date,category_id,note) "
            "VALUES(?,'expense',100,'2026-01-05',?,?)",
            (user_id, category_id, f"Заметка {secret}"),
        )
        con.execute(
            "INSERT INTO piggy_bank_movements(user_id,direction,amount,movement_date,note) "
            "VALUES(?,'deposit',50,'2026-01-06',?)",
            (user_id, f"Копилка {secret}"),
        )
        con.execute(
            "INSERT INTO vacations(user_id,start_date,end_date,amount,payment_date,note) "
            "VALUES(?,'2026-02-01','2026-02-10',1000,'2026-01-28',?)",
            (user_id, f"Отпуск {secret}"),
        )


def _setup(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_PATH", str(tmp_path / "budget.sqlite3"))
    init_db()
    _seed(1, SECRET_A)
    _seed(2, SECRET_B)


def _init_data(user_id: int) -> str:
    fields = {
        "auth_date": str(int(time.time())),
        "user": json.dumps({"id": user_id, "first_name": f"u{user_id}"}, separators=(",", ":")),
    }
    check = "\n".join(f"{k}={fields[k]}" for k in sorted(fields))
    secret = hmac.new(b"WebAppData", TOKEN.encode(), hashlib.sha256).digest()
    fields["hash"] = hmac.new(secret, check.encode(), hashlib.sha256).hexdigest()
    return urlencode(fields)


def test_every_user_has_a_separate_database_file(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch)
    assert user_db_path(1) != user_db_path(2)
    assert all_user_ids() == [1, 2]
    # Raw bytes of one user's file never contain the other user's data.
    raw_a = user_db_path(1).read_bytes() + (user_db_path(1).with_name("1.sqlite3-wal").read_bytes()
                                           if user_db_path(1).with_name("1.sqlite3-wal").exists() else b"")
    assert SECRET_A.encode() in raw_a and SECRET_B.encode() not in raw_a


def test_connect_without_user_fails_closed(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_PATH", str(tmp_path / "budget.sqlite3"))
    init_db()
    token = set_current_user(None)
    try:
        with pytest.raises(RuntimeError):
            connect()
    finally:
        from app.db import reset_current_user
        reset_current_user(token)


def test_backup_contains_only_own_data(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch)
    dump_a = json.dumps(export_user_data(1), ensure_ascii=False)
    dump_b = json.dumps(export_user_data(2), ensure_ascii=False)
    assert SECRET_A in dump_a and SECRET_B not in dump_a
    assert SECRET_B in dump_b and SECRET_A not in dump_b


def test_excel_contains_only_own_data(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch)
    for user_id, own, foreign in ((1, SECRET_A, SECRET_B), (2, SECRET_B, SECRET_A)):
        wb = load_workbook(io.BytesIO(build_workbook(user_id)))
        text = " ".join(
            str(cell.value) for ws in wb.worksheets for row in ws.iter_rows() for cell in row if cell.value
        )
        assert own in text
        assert foreign not in text


def test_restore_does_not_touch_other_users(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch)
    before_b = json.dumps(export_user_data(2)["data"], ensure_ascii=False, sort_keys=True)
    restore_user_data(1, export_user_data(1))
    restore_user_data(1, export_user_data(2))
    after_b = json.dumps(export_user_data(2)["data"], ensure_ascii=False, sort_keys=True)
    assert before_b == after_b


def test_http_requests_only_see_the_authenticated_users_data(tmp_path, monkeypatch):
    from app.cashflow_app import app

    monkeypatch.setenv("DATABASE_PATH", str(tmp_path / "budget.sqlite3"))
    monkeypatch.setenv("BOT_TOKEN", TOKEN)
    monkeypatch.setenv("DEV_MODE", "false")
    monkeypatch.setattr(clock, "today", lambda: __import__("datetime").date(2026, 9, 15))
    init_db()
    # No default user at all for the requests below.
    token = set_current_user(None)
    try:
        client = TestClient(app)
        for user_id, secret in ((11, SECRET_A), (22, SECRET_B)):
            headers = {"X-Telegram-Init-Data": _init_data(user_id)}
            assert client.get("/api/bootstrap", headers=headers).status_code == 200
            created = client.post(
                "/api/transactions", headers=headers,
                json={"type": "expense", "amount": 10, "note": f"Заметка {secret}"},
            )
            assert created.status_code == 200, created.text
        for user_id, own, foreign in ((11, SECRET_A, SECRET_B), (22, SECRET_B, SECRET_A)):
            headers = {"X-Telegram-Init-Data": _init_data(user_id)}
            for path in ("/api/transactions", "/api/backup", "/api/export/xlsx"):
                response = client.get(path, headers=headers)
                assert response.status_code == 200, (path, response.text)
                body = response.content if path.endswith("xlsx") else response.text
                if path.endswith("xlsx"):
                    wb = load_workbook(io.BytesIO(body))
                    body = " ".join(str(c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r if c.value)
                assert own in body, path
                assert foreign not in body, path
        assert all_user_ids() == [11, 22]
    finally:
        from app.db import reset_current_user
        reset_current_user(token)


def test_legacy_shared_database_is_split_without_leaving_other_users_data(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_PATH", str(tmp_path / "budget.sqlite3"))
    # Build an old-style shared file with the real schema, then add two users.
    from app.db import SCHEMA
    legacy = db_path()
    with sqlite3.connect(legacy) as con:
        con.executescript(SCHEMA)
        for uid, secret in ((1, SECRET_A), (2, SECRET_B)):
            con.execute("INSERT INTO users(id,first_name) VALUES(?,?)", (uid, f"u{uid}"))
            con.execute("INSERT INTO settings(user_id) VALUES(?)", (uid,))
            con.execute("INSERT INTO categories(user_id,title) VALUES(?,?)", (uid, f"Кат {secret}"))
            con.execute(
                "INSERT INTO transactions(user_id,type,amount,tx_date,note) VALUES(?,'expense',5,'2026-01-01',?)",
                (uid, f"Заметка {secret}"),
            )
    con.close()

    init_db()
    init_db()  # idempotent

    assert all_user_ids() == [1, 2]
    assert not legacy.exists() and legacy.with_name(legacy.name + ".migrated").exists()
    for uid, own, foreign in ((1, SECRET_A, SECRET_B), (2, SECRET_B, SECRET_A)):
        raw = user_db_path(uid).read_bytes()
        assert own.encode() in raw
        assert foreign.encode() not in raw  # also not in free pages (VACUUM)
        with connect(uid) as c:
            assert [r[0] for r in c.execute("SELECT DISTINCT user_id FROM transactions")] == [uid]
            assert [r[0] for r in c.execute("SELECT id FROM users")] == [uid]


def test_morning_reports_are_built_per_user(tmp_path, monkeypatch):
    from datetime import datetime

    from app.morning_reports import pending_morning_reports, record_morning_report

    _setup(tmp_path, monkeypatch)
    monkeypatch.setattr(clock, "today", lambda: datetime(2026, 9, 15).date())
    reports = pending_morning_reports(datetime(2026, 9, 16, 23, 0))
    assert sorted(r["user_id"] for r in reports) == [1, 2]
    record_morning_report(reports[0])
    assert [r["user_id"] for r in pending_morning_reports(datetime(2026, 9, 16, 23, 0))] == [reports[1]["user_id"]]
