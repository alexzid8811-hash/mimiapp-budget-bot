"""API of cards and accounts: switch, setup wizard, CRUD, order, archive."""
from datetime import date

import pytest
from fastapi.testclient import TestClient

from app import clock
from app.cashflow_app import app
from app.db import connect, ensure_user, init_db, user_scope

TODAY = date(2026, 9, 13)


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_PATH", str(tmp_path / "budget.sqlite3"))
    monkeypatch.setenv("DEV_MODE", "true")
    monkeypatch.setattr(clock, "today", lambda: TODAY)
    monkeypatch.setattr("app.russian_calendar._remote_year", lambda _: None)
    init_db()
    ensure_user(1)
    with connect() as con:
        con.execute("UPDATE income_rules SET amount=100000 WHERE user_id=1")
        con.execute(
            "UPDATE settings SET cashflow_enabled=1,cashflow_start_date=?,"
            "cashflow_start_capital=1000,forecast_months=1 WHERE user_id=1",
            (TODAY.isoformat(),),
        )
    with TestClient(app) as value:
        yield value


def ok(response):
    assert response.status_code == 200, response.text
    return response.json()


def setup(client, accounts=None, envelopes=None):
    accounts = accounts or [
        {"title": "Карта 1", "kind": "debit", "balance": 40000, "is_main": True},
        {"title": "Карта 2", "kind": "debit", "balance": 15000},
        {"title": "Накопительный", "kind": "savings", "balance": 100000},
        {"title": "Кредитка", "kind": "credit", "balance": 0, "credit_limit": 50000, "payment_day": 25},
    ]
    return ok(client.post("/api/accounts/setup", json={"accounts": accounts, "envelopes": envelopes or {}}))


def by_title(snapshot, title):
    return next(a for a in snapshot["accounts"] if a["title"] == title)


def test_feature_is_off_until_accounts_with_balances_are_added(client):
    snap = ok(client.get("/api/accounts"))
    assert snap["enabled"] is False and snap["accounts"] == []
    # The switch alone cannot turn it on: first the cards and their money.
    response = client.put("/api/accounts/settings", json={"enabled": True})
    assert response.status_code == 422
    assert "Сначала добавьте карты" in response.json()["detail"]


def test_wizard_prefill_reads_the_envelopes_without_enabling(client):
    snap = ok(client.get("/api/accounts", params={"with_amounts": True}))
    assert snap["enabled"] is False and snap["budget_configured"] is True
    amounts = {e["key"]: e["amount"] for e in snap["envelopes"]}
    assert set(amounts) == {"spending", "buffer", "piggy"}
    assert amounts["spending"] + amounts["buffer"] == 1000


def test_setup_creates_accounts_main_and_envelopes(client):
    snap = setup(client, envelopes={"piggy": 1})
    assert snap["enabled"] is True
    assert snap["total"] == 155000 and snap["credit_debt"] == 0
    main = by_title(snap, "Карта 1")
    assert main["is_main"] and snap["main_account_id"] == main["id"]
    assert main["opening_date"] == TODAY.isoformat()
    assert [a["title"] for a in snap["accounts"]] == ["Карта 1", "Карта 2", "Накопительный", "Кредитка"]
    credit = by_title(snap, "Кредитка")
    assert credit["debt"] == 0 and credit["available"] == 50000 and credit["payment_day"] == 25
    placed = {e["key"]: e["account_id"] for e in snap["envelopes"]}
    assert placed == {"spending": main["id"], "buffer": main["id"], "piggy": by_title(snap, "Карта 2")["id"]}
    assert main["envelopes"] == ["spending", "buffer"]
    # The wizard runs once.
    assert client.post("/api/accounts/setup", json={"accounts": [{"title": "X", "is_main": True}]}).status_code == 409


@pytest.mark.parametrize("payload, message", [
    ({"accounts": [{"title": "A"}]}, "основной"),
    ({"accounts": [{"title": "К", "kind": "credit", "is_main": True}]}, "Кредитка"),
    ({"accounts": [{"title": "A", "is_main": True}, {"title": "К", "kind": "credit"}],
      "envelopes": {"piggy": 1}}, "кредитку"),
    ({"accounts": [{"title": "A", "is_main": True}], "envelopes": {"reserve": 0}}, "режиме"),
    ({"accounts": [{"title": "К", "kind": "credit", "balance": -5, "is_main": False}]}, "отрицательным"),
])
def test_setup_validation(client, payload, message):
    response = client.post("/api/accounts/setup", json=payload)
    assert response.status_code == 422
    assert message in str(response.json()["detail"])
    assert ok(client.get("/api/accounts"))["accounts"] == []


def test_crud_main_and_credit_rules(client):
    snap = setup(client)
    card2 = by_title(snap, "Карта 2")
    credit = by_title(snap, "Кредитка")
    snap = ok(client.put(f"/api/accounts/{card2['id']}", json={"title": "Зарплатная", "emoji": "💼", "is_main": True}))
    assert by_title(snap, "Зарплатная")["is_main"] and snap["main_account_id"] == card2["id"]
    response = client.put(f"/api/accounts/{credit['id']}", json={"title": "Кредитка", "is_main": True})
    assert response.status_code == 422
    snap = ok(client.put(f"/api/accounts/{credit['id']}",
                         json={"title": "Кредитка", "credit_limit": 70000, "payment_day": 5}))
    assert by_title(snap, "Кредитка")["available"] == 70000
    snap = ok(client.post("/api/accounts", json={"title": "Кошелёк", "emoji": "👛", "kind": "cash", "balance": 500}))
    assert by_title(snap, "Кошелёк")["kind"] == "cash" and snap["total"] == 155500
    # A new account stands at the end of the active list.
    assert snap["accounts"][-1]["title"] == "Кошелёк"
    assert client.put("/api/accounts/999", json={"title": "x"}).status_code == 404


def test_order_follows_drag_and_drop(client):
    snap = setup(client)
    ids = [a["id"] for a in snap["accounts"]]
    ok(client.put("/api/accounts/order", json={"account_ids": list(reversed(ids))}))
    assert [a["id"] for a in ok(client.get("/api/accounts"))["accounts"]] == list(reversed(ids))
    assert client.put("/api/accounts/order", json={"account_ids": ids[:2]}).status_code == 422
    assert client.put("/api/accounts/order", json={"account_ids": [ids[0]] * 4}).status_code == 422


def test_archive_rules_and_restore(client):
    snap = setup(client, envelopes={"piggy": 1})
    main, card2, savings = (by_title(snap, t) for t in ("Карта 1", "Карта 2", "Накопительный"))
    # The main account and an account with envelopes stay active.
    response = client.post(f"/api/accounts/{main['id']}/archive")
    assert response.status_code == 422 and "основной" in response.json()["detail"]
    response = client.post(f"/api/accounts/{card2['id']}/archive")
    assert response.status_code == 422 and "копилка" in response.json()["detail"]
    # Money on the account: move it first.
    response = client.post(f"/api/accounts/{savings['id']}/archive")
    assert response.status_code == 422 and "остаток" in response.json()["detail"]

    empty = ok(client.post("/api/accounts", json={"title": "Пустая", "balance": 0}))
    empty_id = by_title(empty, "Пустая")["id"]
    snap = ok(client.post(f"/api/accounts/{empty_id}/archive"))
    archived = by_title(snap, "Пустая")
    assert archived["archived"] is True
    assert snap["accounts"][-1]["id"] == empty_id
    ok(client.put("/api/accounts/order", json={"account_ids": [a["id"] for a in snap["accounts"] if not a["archived"]]}))
    snap = ok(client.post(f"/api/accounts/{empty_id}/restore"))
    restored = by_title(snap, "Пустая")
    assert restored["archived"] is False
    assert [a["id"] for a in snap["accounts"] if not a["archived"]][-1] == empty_id


def test_delete_only_without_operations_otherwise_archive(client):
    snap = setup(client)
    card2 = by_title(snap, "Карта 2")
    ok(client.post("/api/transactions", json={"type": "expense", "amount": 15000, "account_id": card2["id"]}))
    snap = ok(client.delete(f"/api/accounts/{card2['id']}"))
    assert snap["outcome"] == "archived" and by_title(snap, "Карта 2")["archived"] is True
    # From the archive it cannot be deleted for good either.
    assert client.delete(f"/api/accounts/{card2['id']}").status_code == 422

    fresh = by_title(ok(client.post("/api/accounts", json={"title": "Новая", "balance": 100})), "Новая")
    snap = ok(client.delete(f"/api/accounts/{fresh['id']}"))
    assert snap["outcome"] == "deleted"
    assert all(a["title"] != "Новая" for a in snap["accounts"])


def test_expense_from_a_card_changes_its_balance_and_the_budget_as_before(client):
    snap = setup(client)
    card1 = by_title(snap, "Карта 1")
    before = ok(client.get("/api/cashflow"))["available_today"]
    tx = ok(client.post("/api/transactions", json={"type": "expense", "amount": 2000, "account_id": card1["id"]}))
    assert tx["account_id"] == card1["id"]
    snap = ok(client.get("/api/accounts"))
    assert by_title(snap, "Карта 1")["balance"] == 38000 and snap["total"] == 153000
    assert ok(client.get("/api/cashflow"))["available_today"] == before - 2000
    # An edit without the account (old app) keeps it; an explicit null removes it.
    ok(client.put(f"/api/transactions/{tx['id']}", json={"type": "expense", "amount": 1000}))
    assert by_title(ok(client.get("/api/accounts")), "Карта 1")["balance"] == 39000
    ok(client.put(f"/api/transactions/{tx['id']}", json={"type": "expense", "amount": 1000, "account_id": None}))
    assert by_title(ok(client.get("/api/accounts")), "Карта 1")["balance"] == 40000
    listed = ok(client.get("/api/transactions"))[0]
    assert listed["account_id"] is None


def test_operation_account_validation(client):
    snap = setup(client)
    credit = by_title(snap, "Кредитка")
    response = client.post("/api/transactions", json={"type": "income", "amount": 10, "account_id": credit["id"]})
    assert response.status_code == 422
    assert client.post("/api/transactions", json={"type": "expense", "amount": 10, "account_id": 999}).status_code == 422
    empty = by_title(ok(client.post("/api/accounts", json={"title": "Пустая"})), "Пустая")
    ok(client.post(f"/api/accounts/{empty['id']}/archive"))
    response = client.post("/api/transactions", json={"type": "expense", "amount": 10, "account_id": empty["id"]})
    assert response.status_code == 422 and "архиве" in response.json()["detail"]


def test_history_names_the_account_and_marks_the_archive(client):
    snap = setup(client)
    card2 = by_title(snap, "Карта 2")
    ok(client.post("/api/transactions", json={"type": "expense", "amount": 15000, "account_id": card2["id"]}))
    ok(client.post(f"/api/accounts/{card2['id']}/archive"))
    row = ok(client.get("/api/transactions"))[0]
    assert row["account_title"] == "Карта 2" and row["account_archived"] == 1


def test_operation_before_the_opening_date_does_not_touch_the_balance(client):
    snap = setup(client)
    card1 = by_title(snap, "Карта 1")
    ok(client.post("/api/transactions", json={"type": "expense", "amount": 500, "tx_date": "2026-09-10",
                                              "account_id": card1["id"]}))
    assert by_title(ok(client.get("/api/accounts")), "Карта 1")["balance"] == 40000


def test_switch_off_keeps_everything_and_counts_operations_on_return(client):
    setup(client)
    snap = ok(client.put("/api/accounts/settings", json={"enabled": False}))
    assert snap["enabled"] is False and len(snap["accounts"]) == 4
    with connect() as con:
        # created_at has a one-second resolution: make the operations "later".
        con.execute("UPDATE settings SET accounts_disabled_at='2000-01-01 00:00:00'")
    ok(client.post("/api/transactions", json={"type": "expense", "amount": 100}))
    ok(client.post("/api/piggy-bank/deposit", json={"amount": 50}))
    snap = ok(client.put("/api/accounts/settings", json={"enabled": True}))
    assert snap["enabled"] is True and snap["operations_without_account"] == 2
    assert ok(client.get("/api/accounts"))["operations_without_account"] == 0


def test_last_account_leaving_switches_the_feature_off(client):
    snap = setup(client, accounts=[{"title": "Единственная", "balance": 0, "is_main": True}])
    only = snap["accounts"][0]
    snap = ok(client.delete(f"/api/accounts/{only['id']}"))
    assert snap["outcome"] == "deleted" and snap["enabled"] is False and snap["main_account_id"] is None


def test_users_do_not_see_each_others_accounts(client):
    setup(client)
    ensure_user(2)
    with user_scope(2), connect() as con:
        assert con.execute("SELECT COUNT(*) FROM accounts").fetchone()[0] == 0
    with user_scope(1), connect() as con:
        assert con.execute("SELECT COUNT(*) FROM accounts").fetchone()[0] == 4


def test_budget_and_old_responses_do_not_change_while_the_feature_is_off(client):
    before = {path: ok(client.get(path)) for path in ("/api/cashflow", "/api/dashboard")}
    ok(client.post("/api/transactions", json={"type": "expense", "amount": 300}))
    after_expense = ok(client.get("/api/cashflow"))
    assert after_expense["available_today"] == before["/api/cashflow"]["available_today"] - 300
    assert ok(client.get("/api/accounts"))["enabled"] is False
    tx = ok(client.get("/api/transactions"))[0]
    assert tx["account_id"] is None and tx["funding_envelope"] is None


# ---------- Stage 2: transfers, corrections, history, payroll income ----------


def balances(client):
    snap = ok(client.get("/api/accounts"))
    return {a["title"]: a["balance"] for a in snap["accounts"]}, snap


def test_transfer_moves_money_between_two_accounts_only(client):
    snap = setup(client)
    c1, c2 = by_title(snap, "Карта 1")["id"], by_title(snap, "Карта 2")["id"]
    budget = ok(client.get("/api/cashflow"))["available_today"]
    ok(client.post("/api/accounts/transfers", json={"from_account_id": c1, "to_account_id": c2, "amount": 10000}))
    values, snap = balances(client)
    assert (values["Карта 1"], values["Карта 2"], snap["total"]) == (30000, 25000, 155000)
    assert ok(client.get("/api/cashflow"))["available_today"] == budget
    # Cash withdrawal is a transfer to a cash account.
    cash = by_title(ok(client.post("/api/accounts", json={"title": "Наличные", "kind": "cash"})), "Наличные")["id"]
    ok(client.post("/api/accounts/transfers", json={"from_account_id": c1, "to_account_id": cash, "amount": 3000}))
    values, snap = balances(client)
    assert (values["Карта 1"], values["Наличные"], snap["total"]) == (27000, 3000, 155000)


@pytest.mark.parametrize("change, status", [
    ({"to_account_id": "same"}, 422),
    ({"to_account_id": "credit"}, 422),
    ({"from_account_id": "credit"}, 422),
    ({"date": "2026-09-20"}, 422),
    ({"to_account_id": 999}, 404),
])
def test_transfer_validation(client, change, status):
    snap = setup(client)
    ids = {"same": by_title(snap, "Карта 1")["id"], "credit": by_title(snap, "Кредитка")["id"]}
    payload = {"from_account_id": ids["same"], "to_account_id": by_title(snap, "Карта 2")["id"], "amount": 10}
    payload.update({k: ids.get(v, v) for k, v in change.items()})
    assert client.post("/api/accounts/transfers", json=payload).status_code == status


def test_transfer_edit_and_delete_recalculate(client):
    snap = setup(client)
    c1, c2 = by_title(snap, "Карта 1")["id"], by_title(snap, "Карта 2")["id"]
    ok(client.post("/api/accounts/transfers", json={"from_account_id": c1, "to_account_id": c2, "amount": 100}))
    history = ok(client.get(f"/api/accounts/{c1}/history"))
    item = history["items"][0]
    assert item["type"] == "transfer" and item["amount"] == -100 and "Карта 2" in item["title"]
    ok(client.put(f"/api/accounts/transfers/{item['id']}",
                  json={"from_account_id": c1, "to_account_id": c2, "amount": 250}))
    assert balances(client)[0]["Карта 1"] == 39750
    ok(client.delete(f"/api/accounts/transfers/{item['id']}"))
    assert balances(client)[0]["Карта 1"] == 40000


def test_correct_balance_writes_the_difference_and_keeps_the_budget(client):
    snap = setup(client)
    c2 = by_title(snap, "Карта 2")["id"]
    budget = ok(client.get("/api/cashflow"))["available_today"]
    ok(client.post(f"/api/accounts/{c2}/adjust", json={"target": 15500, "note": "проценты"}))
    values, snap = balances(client)
    assert values["Карта 2"] == 15500 and snap["total"] == 155500
    assert ok(client.get("/api/cashflow"))["available_today"] == budget
    item = ok(client.get(f"/api/accounts/{c2}/history"))["items"][0]
    assert (item["type"], item["amount"], item["title"]) == ("adjustment", 500, "Исправление остатка")
    assert client.post(f"/api/accounts/{c2}/adjust", json={"target": 15500}).status_code == 422
    assert client.post(f"/api/accounts/{c2}/adjust", json={"target": 1, "date": "2026-09-01"}).status_code == 422
    ok(client.delete(f"/api/accounts/adjustments/{item['id']}"))
    assert balances(client)[0]["Карта 2"] == 15000


def test_correct_credit_debt_is_interest_or_cashback(client):
    snap = setup(client)
    credit = by_title(snap, "Кредитка")["id"]
    ok(client.post(f"/api/accounts/{credit}/adjust", json={"target": 700}))
    item = ok(client.get(f"/api/accounts/{credit}/history"))["items"][0]
    assert (item["reason"], item["amount"], item["title"]) == ("interest", -700, "Проценты банка")
    state = by_title(ok(client.get("/api/accounts")), "Кредитка")
    assert (state["debt"], state["available"]) == (700, 49300)
    ok(client.post(f"/api/accounts/{credit}/adjust", json={"target": 200}))
    item = ok(client.get(f"/api/accounts/{credit}/history"))["items"][0]
    assert (item["reason"], item["amount"]) == ("cashback", 500)
    assert ok(client.get("/api/accounts"))["total"] == 155000


def test_history_lists_operations_before_opening_as_not_counted(client):
    snap = setup(client)
    c1 = by_title(snap, "Карта 1")["id"]
    ok(client.post("/api/transactions", json={"type": "expense", "amount": 50, "tx_date": "2026-09-01", "account_id": c1}))
    ok(client.post("/api/transactions", json={"type": "income", "amount": 70, "account_id": c1}))
    items = ok(client.get(f"/api/accounts/{c1}/history"))["items"]
    assert [(i["type"], i["amount"], i["counted"]) for i in items] == [("income", 70, True), ("expense", -50, False)]


def test_income_arrives_on_the_chosen_card(client):
    snap = setup(client)
    c2 = by_title(snap, "Карта 2")["id"]
    ok(client.post("/api/transactions", json={"type": "income", "amount": 5000, "income_destination": "daily",
                                              "account_id": c2}))
    assert balances(client)[0]["Карта 2"] == 20000


def advance_key(client):
    rows = ok(client.get("/api/cashflow"))["periods"]
    return next(r for r in rows if r["payday"] == "2026-09-22")["override_key"]


def test_confirmed_payout_arrives_on_the_main_account(client, monkeypatch):
    setup(client)
    monkeypatch.setattr(clock, "today", lambda: date(2026, 9, 22))
    key = advance_key(client)
    # Not confirmed: the balances stay as they were.
    assert balances(client)[0]["Карта 1"] == 40000
    ok(client.put(f"/api/cashflow/income-overrides/{key}", json={"amount": 60000, "confirmed": True}))
    values, snap = balances(client)
    assert values["Карта 1"] == 100000 and snap["total"] == 215000
    # A later correction of the amount follows; the reset removes the income.
    ok(client.put(f"/api/cashflow/income-overrides/{key}", json={"amount": 58000}))
    assert balances(client)[0]["Карта 1"] == 98000
    ok(client.delete(f"/api/cashflow/income-overrides/{key}"))
    assert balances(client)[0]["Карта 1"] == 40000


def test_confirmed_payout_can_go_to_another_card(client, monkeypatch):
    snap = setup(client)
    c2 = by_title(snap, "Карта 2")["id"]
    monkeypatch.setattr(clock, "today", lambda: date(2026, 9, 22))
    ok(client.put(f"/api/cashflow/income-overrides/{advance_key(client)}",
                  json={"amount": 1000, "confirmed": True, "account_id": c2}))
    values, _ = balances(client)
    assert (values["Карта 1"], values["Карта 2"]) == (40000, 16000)


def test_payout_counted_without_confirmation_offers_to_record_the_income(client, monkeypatch):
    setup(client)
    monkeypatch.setattr(clock, "today", lambda: date(2026, 10, 7))
    snap = ok(client.get("/api/accounts"))
    hints = [h for h in snap["hints"] if h["type"] == "payroll_income"]
    # 22.09 is counted by the budget on its own; 07.10 still waits for «Получил ✓».
    assert [h["payday"] for h in hints] == ["2026-09-22"]
    assert hints[0]["confirmed"] is False and hints[0]["amount"] > 0
    snap = ok(client.post("/api/accounts/payroll-income", json={"period_start": hints[0]["period_start"]}))
    assert by_title(snap, "Карта 1")["balance"] == 40000 + hints[0]["amount"]
    assert snap["hints"] == []
    assert client.post("/api/accounts/payroll-income", json={"period_start": hints[0]["period_start"]}).status_code == 422


def test_irregular_income_keeps_its_card(client):
    snap = setup(client)
    c2 = by_title(snap, "Карта 2")["id"]
    ok(client.put("/api/budget-mode", json={"budget_mode": "irregular", "start_date": TODAY.isoformat(),
                                             "start_total": 10000, "start_reserve": 0}))
    income = ok(client.post("/api/irregular/incomes", json={"amount": 1000, "account_id": c2}))
    assert balances(client)[0]["Карта 2"] == 16000
    ok(client.put(f"/api/irregular/incomes/{income['id']}", json={"amount": 2000}))
    assert balances(client)[0]["Карта 2"] == 17000
    listed = ok(client.get("/api/irregular/incomes"))[0]
    assert listed["account_id"] == c2


# ---------- Stage 3: envelopes on cards, linked transfers, hints ----------


def amounts(client):
    return {e["key"]: e["amount"] for e in ok(client.get("/api/accounts", params={"with_amounts": True}))["envelopes"]}


def two_cards(client, envelopes):
    """Card 1 (main) and card 2, their balances equal to their envelopes."""
    ok(client.post("/api/piggy-bank/deposit", json={"amount": 8000}))
    env = amounts(client)
    cards = {0: 0, 1: 0}
    for key, value in env.items():
        cards[envelopes.get(key, 0)] += value
    snap = setup(client, accounts=[
        {"title": "Карта 1", "balance": cards[0], "is_main": True},
        {"title": "Карта 2", "balance": cards[1]},
        {"title": "Кредитка", "kind": "credit", "credit_limit": 50000},
    ], envelopes=envelopes)
    return by_title(snap, "Карта 1")["id"], by_title(snap, "Карта 2")["id"]


def state(client):
    snap = ok(client.get("/api/accounts"))
    cards = {a["title"]: a for a in snap["accounts"]}
    return snap, cards


def differences(client):
    _, cards = state(client)
    return {t: a["breakdown"]["difference"] for t, a in cards.items() if a["breakdown"] is not None}


def linked(client):
    with connect() as con:
        return [dict(r) for r in con.execute(
            "SELECT from_account_id,to_account_id,amount,kind,piggy_movement_id,reserve_movement_id "
            "FROM account_transfers WHERE kind='linked' ORDER BY id")]


def test_breakdown_after_setup_is_balanced_and_several_envelopes_share_a_card(client):
    two_cards(client, {"piggy": 1})
    snap, cards = state(client)
    assert cards["Карта 1"]["envelopes"] == ["spending", "buffer"]
    labels = [e["label"] for e in cards["Карта 1"]["breakdown"]["envelopes"]]
    assert labels == ["на траты", "буфер"]
    assert differences(client) == {"Карта 1": 0, "Карта 2": 0}
    assert cards["Кредитка"]["breakdown"] is None


def test_piggy_from_the_day_creates_a_transfer_only_between_different_cards(client):
    c1, c2 = two_cards(client, {"piggy": 1})
    movement = ok(client.post("/api/piggy-bank/deposit", json={"amount": 50, "source": "daily_budget"}))["movement"]
    assert linked(client) == [{"from_account_id": c1, "to_account_id": c2, "amount": 50, "kind": "linked",
                               "piggy_movement_id": movement["id"], "reserve_movement_id": None}]
    assert differences(client) == {"Карта 1": 0, "Карта 2": 0}
    # «Копилка → карта» goes back.
    ok(client.post("/api/piggy-bank/to-card", json={"amount": 20, "purpose": "today"}))
    assert linked(client)[-1]["from_account_id"] == c2 and linked(client)[-1]["to_account_id"] == c1
    assert differences(client) == {"Карта 1": 0, "Карта 2": 0}
    # Deleting the operation deletes its transfer.
    ok(client.delete(f"/api/piggy-bank/movements/{movement['id']}"))
    assert [t["amount"] for t in linked(client)] == [20]
    assert differences(client) == {"Карта 1": 0, "Карта 2": 0}


def test_piggy_on_the_same_card_needs_no_transfer(client):
    two_cards(client, {})
    ok(client.post("/api/piggy-bank/deposit", json={"amount": 50, "source": "daily_budget"}))
    assert linked(client) == []
    assert differences(client) == {"Карта 1": 0, "Карта 2": 0}


def test_piggy_outside_money_from_outside_from_a_card_or_already_there(client):
    c1, c2 = two_cards(client, {"piggy": 1})
    before = state(client)[1]
    ok(client.post("/api/piggy-bank/deposit", json={"amount": 100}))                              # from outside
    ok(client.post("/api/piggy-bank/deposit", json={"amount": 30, "account_id": c1}))            # from card 1
    _, cards = state(client)
    assert cards["Карта 2"]["balance"] == before["Карта 2"]["balance"] + 130
    assert cards["Карта 1"]["balance"] == before["Карта 1"]["balance"] - 30
    # Outside money taken from card 1 was not outside the budget there.
    assert differences(client) == {"Карта 1": -30, "Карта 2": 0}
    # «Исправить остаток» +500 (interest), then «Добавить в копилку»: no money moves.
    ok(client.post(f"/api/accounts/{c2}/adjust", json={"target": cards["Карта 2"]["balance"] + 500}))
    assert differences(client)["Карта 2"] == 500
    ok(client.post("/api/piggy-bank/deposit", json={"amount": 500, "account_move": False}))
    assert differences(client) == {"Карта 1": -30, "Карта 2": 0}
    assert state(client)[1]["Карта 2"]["balance"] == cards["Карта 2"]["balance"] + 500


def test_piggy_withdrawal_to_a_card_or_outside(client):
    c1, c2 = two_cards(client, {"piggy": 1})
    before = state(client)[1]
    ok(client.post("/api/piggy-bank/withdraw", json={"amount": 100, "account_id": c1}))
    ok(client.post("/api/piggy-bank/withdraw", json={"amount": 40}))
    _, cards = state(client)
    assert cards["Карта 2"]["balance"] == before["Карта 2"]["balance"] - 140
    assert cards["Карта 1"]["balance"] == before["Карта 1"]["balance"] + 100
    assert linked(client)[-1]["to_account_id"] is None
    # Money withdrawn to card 1 is outside the budget there.
    assert differences(client) == {"Карта 1": 100, "Карта 2": 0}


def test_expense_from_the_piggy_bank_keeps_the_daily_budget(client):
    c1, c2 = two_cards(client, {"piggy": 1})
    flow = ok(client.get("/api/cashflow"))
    piggy = ok(client.get("/api/piggy-bank"))["balance"]
    tx = ok(client.post("/api/transactions", json={"type": "expense", "amount": 2000, "account_id": c2,
                                                   "funding_envelope": "piggy", "category_id": 1}))
    assert tx["funding_envelope"] == "piggy"
    assert ok(client.get("/api/cashflow"))["available_today"] == flow["available_today"]
    assert ok(client.get("/api/piggy-bank"))["balance"] == piggy - 2000
    assert differences(client) == {"Карта 1": 0, "Карта 2": 0}
    assert ok(client.get("/api/transactions"))[0]["category_id"] == 1
    movement = next(m for m in ok(client.get("/api/piggy-bank"))["movements"] if m["expense_transaction_id"])
    assert client.delete(f"/api/piggy-bank/movements/{movement['id']}").status_code == 422
    # Editing follows; switching to «из бюджета» moves it back to the day.
    ok(client.put(f"/api/transactions/{tx['id']}", json={"type": "expense", "amount": 1500}))
    assert ok(client.get("/api/piggy-bank"))["balance"] == piggy - 1500
    ok(client.put(f"/api/transactions/{tx['id']}", json={"type": "expense", "amount": 1500,
                                                         "funding_envelope": "spending"}))
    assert ok(client.get("/api/piggy-bank"))["balance"] == piggy
    assert ok(client.get("/api/cashflow"))["available_today"] == flow["available_today"] - 1500
    # More than the piggy bank holds is refused.
    response = client.post("/api/transactions", json={"type": "expense", "amount": piggy + 1, "account_id": c2,
                                                      "funding_envelope": "piggy"})
    assert response.status_code == 422


def test_expense_outside_the_budget_changes_only_the_card(client):
    c1, _ = two_cards(client, {})
    flow = ok(client.get("/api/cashflow"))
    ok(client.post("/api/transactions", json={"type": "expense", "amount": 700, "account_id": c1,
                                              "funding_envelope": "none"}))
    assert ok(client.get("/api/cashflow"))["available_today"] == flow["available_today"]
    assert differences(client)["Карта 1"] == -700
    assert client.post("/api/transactions", json={"type": "expense", "amount": 1, "funding_envelope": "reserve"}).status_code == 422


def test_expense_from_spending_on_another_card_asks_for_a_transfer(client):
    c1, c2 = two_cards(client, {"piggy": 1})
    ok(client.post("/api/transactions", json={"type": "expense", "amount": 300, "account_id": c2,
                                              "funding_envelope": "spending"}))
    snap, _ = state(client)
    hints = [h for h in snap["hints"] if h["type"] == "transfer"]
    assert hints == [{"type": "transfer", "from_account_id": c1, "to_account_id": c2, "amount": 300}]
    ok(client.post("/api/accounts/transfers", json={"from_account_id": c1, "to_account_id": c2,
                                                    "amount": 300, "kind": "hint"}))
    assert [h for h in state(client)[0]["hints"] if h["type"] == "transfer"] == []
    assert differences(client) == {"Карта 1": 0, "Карта 2": 0}


def test_bill_remainder_to_the_piggy_bank_moves_between_cards(client):
    with connect() as con:
        con.execute("INSERT INTO bill_rules(user_id,title,amount,day_of_month) VALUES(1,'Интернет',1000,14)")
    c1, c2 = two_cards(client, {"piggy": 1})
    bill = client.get("/api/plan").json()[0]
    payment = ok(client.post(f"/api/bills/{bill['id']}/pay", json={"due_date": bill["due_date"]}))
    assert payment["account_id"] == c1                      # the buffer's card by default
    ok(client.put(f"/api/bill-payments/{payment['id']}", json={"amount": 800, "remainder_destination": "piggy"}))
    assert [(t["from_account_id"], t["to_account_id"], t["amount"]) for t in linked(client)] == [(c1, c2, 200)]
    assert differences(client) == {"Карта 1": 0, "Карта 2": 0}
    ok(client.put(f"/api/bill-payments/{payment['id']}", json={"amount": 700, "remainder_destination": "piggy"}))
    assert [t["amount"] for t in linked(client)] == [300]
    ok(client.put(f"/api/bill-payments/{payment['id']}", json={"amount": 700, "remainder_destination": "budget"}))
    assert linked(client) == []
    assert differences(client) == {"Карта 1": 0, "Карта 2": 0}


def test_moving_an_envelope_with_its_money(client):
    c1, c2 = two_cards(client, {})
    snap = ok(client.put("/api/accounts/envelopes/piggy", json={"account_id": c2, "move_money": True}))
    assert by_title(snap, "Карта 2")["envelopes"] == ["piggy"]
    assert differences(client) == {"Карта 1": 0, "Карта 2": 0}
    with connect() as con:
        assert con.execute("SELECT kind,amount FROM account_transfers").fetchall()[0]["kind"] == "reassign"
    # Without the money: the old card has it outside the budget, the new one misses it.
    ok(client.put("/api/accounts/envelopes/piggy", json={"account_id": c1}))
    diff = differences(client)
    assert diff["Карта 2"] == -diff["Карта 1"] > 0
    credit = by_title(snap, "Кредитка")["id"]
    assert client.put("/api/accounts/envelopes/piggy", json={"account_id": credit}).status_code == 422
    assert client.put("/api/accounts/envelopes/reserve", json={"account_id": c1}).status_code == 422


def test_payday_hint_when_buffer_and_spending_lie_apart(client, monkeypatch):
    c1, c2 = two_cards(client, {"spending": 1})
    monkeypatch.setattr(clock, "today", lambda: date(2026, 9, 22))
    rows = ok(client.get("/api/cashflow"))["periods"]
    key = next(r for r in rows if r["payday"] == "2026-09-22")["override_key"]
    ok(client.put(f"/api/cashflow/income-overrides/{key}", json={"amount": 100000, "confirmed": True}))
    snap, cards = state(client)
    hint = next(h for h in snap["hints"] if h["type"] == "transfer")
    assert (hint["from_account_id"], hint["to_account_id"]) == (c1, c2)
    assert differences(client) == {"Карта 1": 0, "Карта 2": -hint["amount"]}
    ok(client.post("/api/accounts/transfers", json={"from_account_id": c1, "to_account_id": c2,
                                                    "amount": hint["amount"], "kind": "hint"}))
    assert [h for h in state(client)[0]["hints"] if h["type"] == "transfer"] == []
    assert differences(client) == {"Карта 1": 0, "Карта 2": 0}


def irregular(client, reserve_total=20000):
    ok(client.put("/api/budget-mode", json={"budget_mode": "irregular", "start_date": TODAY.isoformat(),
                                             "start_total": 30000, "start_reserve": reserve_total}))


def test_irregular_income_reserve_share_is_a_hint_and_reserve_moves_are_linked(client):
    irregular(client)
    c1, c2 = two_cards(client, {"piggy": 1, "reserve": 1})
    ok(client.post("/api/irregular/incomes", json={"amount": 20000, "account_id": c1}))
    snap, _ = state(client)
    hint = next(h for h in snap["hints"] if h["type"] == "transfer")
    assert (hint["from_account_id"], hint["to_account_id"], hint["amount"]) == (c1, c2, 2000)
    assert differences(client) == {"Карта 1": 0, "Карта 2": -2000}
    ok(client.post("/api/accounts/transfers", json={"from_account_id": c1, "to_account_id": c2,
                                                    "amount": 2000, "kind": "hint"}))
    assert differences(client) == {"Карта 1": 0, "Карта 2": 0}
    # Reserve -> free money: a linked transfer back to card 1.
    ok(client.post("/api/irregular/reserve/withdraw", json={"amount": 4000, "purpose": "to_free", "reason": "нужно"}))
    assert linked(client)[-1]["from_account_id"] == c2 and linked(client)[-1]["reserve_movement_id"]
    assert differences(client) == {"Карта 1": 0, "Карта 2": 0}
    # An expense paid from the reserve is paid from its card; the day does not change.
    flow = ok(client.get("/api/irregular"))
    ok(client.post("/api/irregular/reserve/withdraw", json={"amount": 5000, "purpose": "pay_expense",
                                                           "reason": "телефон", "category_id": 1}))
    tx = ok(client.get("/api/transactions"))[0]
    assert (tx["account_id"], tx["funding_envelope"]) == (c2, "reserve")
    assert ok(client.get("/api/irregular"))["available_today"] == flow["available_today"]
    assert differences(client) == {"Карта 1": 0, "Карта 2": 0}
    # Free money back to the reserve, outside money into the reserve from card 1.
    ok(client.post("/api/irregular/reserve/deposit", json={"amount": 300, "kind": "from_free"}))
    ok(client.post("/api/irregular/reserve/deposit", json={"amount": 100, "account_id": c1}))
    assert differences(client) == {"Карта 1": -100, "Карта 2": 0}
    movement = next(m for m in ok(client.get("/api/irregular/reserve"))["movements"]
                    if m["kind"] == "deposit" and m["amount"] == 100)
    ok(client.delete(f"/api/irregular/reserve/movements/{movement['id']}"))
    assert differences(client) == {"Карта 1": 0, "Карта 2": 0}


def test_expense_from_the_reserve_with_the_expense_form(client):
    irregular(client)
    _, c2 = two_cards(client, {"reserve": 1})
    reserve = ok(client.get("/api/irregular"))["reserve_balance"]
    flow = ok(client.get("/api/irregular"))["available_today"]
    tx = ok(client.post("/api/transactions", json={"type": "expense", "amount": 500, "account_id": c2,
                                                   "funding_envelope": "reserve", "category_id": 2}))
    snap = ok(client.get("/api/irregular"))
    assert (snap["reserve_balance"], snap["available_today"]) == (reserve - 500, flow)
    assert differences(client) == {"Карта 1": 0, "Карта 2": 0}
    ok(client.delete(f"/api/transactions/{tx['id']}"))
    assert ok(client.get("/api/irregular"))["reserve_balance"] == reserve


def test_linked_transfers_are_not_created_while_the_feature_is_off(client):
    two_cards(client, {"piggy": 1})
    ok(client.put("/api/accounts/settings", json={"enabled": False}))
    ok(client.post("/api/piggy-bank/deposit", json={"amount": 50, "source": "daily_budget"}))
    assert linked(client) == []


def test_envelope_amounts_do_not_depend_on_cards(client):
    before = amounts(client)
    two_cards(client, {"piggy": 1})
    after = amounts(client)
    assert after["spending"] == before["spending"] and after["buffer"] == before["buffer"]
    ok(client.put("/api/accounts/envelopes/buffer", json={"account_id": by_title(ok(client.get("/api/accounts")), "Карта 2")["id"],
                                                          "move_money": True}))
    assert amounts(client) == after


# ---------- Stage 4: credit cards ----------


def categories_spent(client):
    data = ok(client.get("/api/analytics", params={"mode": "month", "anchor": TODAY.isoformat()}))
    return {c["title"]: c["amount"] for c in data["categories"]}, data


def test_example_1_balances_total_and_credit_card_through_the_api(client, monkeypatch):
    snap = setup(client)
    c1, c2, credit = (by_title(snap, t)["id"] for t in ("Карта 1", "Карта 2", "Кредитка"))
    budget = lambda: ok(client.get("/api/cashflow"))["available_today"]  # noqa: E731
    start = budget()
    # 1. 2 000 from card 1 out of the daily budget.
    ok(client.post("/api/transactions", json={"type": "expense", "amount": 2000, "account_id": c1,
                                              "funding_envelope": "spending"}))
    snap, cards = state(client)
    assert (cards["Карта 1"]["balance"], snap["total"], budget()) == (38000, 153000, start - 2000)
    # 2. 3 000 with the credit card: only the debt grows.
    tx = ok(client.post("/api/transactions", json={"type": "expense", "amount": 3000, "account_id": credit,
                                                   "category_id": 1}))
    assert tx["funding_envelope"] == "none"
    snap, cards = state(client)
    assert (snap["total"], cards["Кредитка"]["debt"], cards["Кредитка"]["available"]) == (153000, 3000, 47000)
    assert budget() == start - 2000 and snap["credit_debt"] == 3000
    assert categories_spent(client)[0]["Продукты"] == 3000
    # 3. Card 1 -> card 2.
    ok(client.post("/api/accounts/transfers", json={"from_account_id": c1, "to_account_id": c2, "amount": 10000}))
    snap, cards = state(client)
    assert (cards["Карта 1"]["balance"], cards["Карта 2"]["balance"], snap["total"]) == (28000, 25000, 153000)
    # 4. Repay 3 000 from card 1 out of the daily budget.
    ok(client.post(f"/api/accounts/{credit}/repay", json={"from_account_id": c1, "amount": 3000}))
    snap, cards = state(client)
    assert (cards["Карта 1"]["balance"], snap["total"], cards["Кредитка"]["debt"]) == (25000, 150000, 0)
    assert (cards["Кредитка"]["available"], snap["credit_debt"]) == (50000, 0)
    assert budget() == start - 5000
    # The repayment is no expense in the analytics: the purchase already is.
    assert categories_spent(client)[0] == {"Продукты": 3000, "Без категории": 2000}
    # 5. Salary 60 000 confirmed on the main card.
    monkeypatch.setattr(clock, "today", lambda: date(2026, 9, 22))
    ok(client.put(f"/api/cashflow/income-overrides/{advance_key(client)}", json={"amount": 60000, "confirmed": True}))
    snap, cards = state(client)
    assert (cards["Карта 1"]["balance"], snap["total"]) == (85000, 210000)


def test_example_5_bank_interest_in_a_repayment(client):
    snap = setup(client)
    c1, credit = by_title(snap, "Карта 1")["id"], by_title(snap, "Кредитка")["id"]
    budget = ok(client.get("/api/cashflow"))["available_today"]
    ok(client.post("/api/transactions", json={"type": "expense", "amount": 3000, "account_id": credit}))
    assert ok(client.get("/api/cashflow"))["available_today"] == budget
    ok(client.post(f"/api/accounts/{credit}/repay", json={"from_account_id": c1, "amount": 3000, "interest": 500}))
    snap, cards = state(client)
    assert (cards["Кредитка"]["debt"], cards["Кредитка"]["available"]) == (500, 49500)
    assert cards["Карта 1"]["balance"] == 37000 and snap["credit_debt"] == 500
    assert ok(client.get("/api/cashflow"))["available_today"] == budget - 3000
    spent, _ = categories_spent(client)
    assert spent["Проценты банка"] == 500
    # The interest expense belongs to the repayment.
    interest = next(t for t in ok(client.get("/api/transactions")) if t["category_title"] == "Проценты банка")
    assert client.delete(f"/api/transactions/{interest['id']}").status_code == 422
    assert client.put(f"/api/transactions/{interest['id']}", json={"type": "expense", "amount": 1}).status_code == 422
    ok(client.post(f"/api/accounts/{credit}/repay", json={"from_account_id": c1, "amount": 500}))
    snap, cards = state(client)
    assert cards["Кредитка"]["debt"] == 0 and snap["credit_debt"] == 0


def test_the_same_through_correct_debt_with_interest(client):
    snap = setup(client)
    c1, credit = by_title(snap, "Карта 1")["id"], by_title(snap, "Кредитка")["id"]
    ok(client.post("/api/transactions", json={"type": "expense", "amount": 3000, "account_id": credit}))
    ok(client.post(f"/api/accounts/{credit}/repay", json={"from_account_id": c1, "amount": 3000}))
    budget = ok(client.get("/api/cashflow"))["available_today"]
    ok(client.post(f"/api/accounts/{credit}/adjust", json={"target": 500}))
    snap, cards = state(client)
    assert cards["Кредитка"]["debt"] == 500 and categories_spent(client)[0]["Проценты банка"] == 500
    assert ok(client.get("/api/cashflow"))["available_today"] == budget
    item = ok(client.get(f"/api/accounts/{credit}/history"))["items"][0]
    ok(client.delete(f"/api/accounts/adjustments/{item['id']}"))
    assert "Проценты банка" not in categories_spent(client)[0]


def test_repayment_from_the_piggy_bank_or_outside_the_budget(client):
    c1, c2 = two_cards(client, {"piggy": 1})
    credit = by_title(ok(client.get("/api/accounts")), "Кредитка")["id"]
    ok(client.post("/api/transactions", json={"type": "expense", "amount": 900, "account_id": credit}))
    budget = ok(client.get("/api/cashflow"))["available_today"]
    piggy = ok(client.get("/api/piggy-bank"))["balance"]
    ok(client.post(f"/api/accounts/{credit}/repay", json={"from_account_id": c2, "amount": 600, "funding": "piggy"}))
    assert ok(client.get("/api/piggy-bank"))["balance"] == piggy - 600
    assert ok(client.get("/api/cashflow"))["available_today"] == budget
    assert differences(client) == {"Карта 1": 0, "Карта 2": 0}
    movement = next(m for m in ok(client.get("/api/piggy-bank"))["movements"] if m["note"].startswith("Погашение"))
    assert client.delete(f"/api/piggy-bank/movements/{movement['id']}").status_code == 422
    ok(client.post(f"/api/accounts/{credit}/repay", json={"from_account_id": c1, "amount": 300, "funding": "none"}))
    assert differences(client)["Карта 1"] == -300
    snap, cards = state(client)
    assert cards["Кредитка"]["debt"] == 0
    # Deleting a repayment brings the debt and the piggy bank back.
    item = next(i for i in ok(client.get(f"/api/accounts/{credit}/history"))["items"] if i["amount"] == 600)
    ok(client.delete(f"/api/accounts/transfers/{item['id']}"))
    assert ok(client.get("/api/piggy-bank"))["balance"] == piggy
    assert state(client)[1]["Кредитка"]["debt"] == 600


def test_repayment_validation(client):
    snap = setup(client)
    c1, credit = by_title(snap, "Карта 1")["id"], by_title(snap, "Кредитка")["id"]
    assert client.post(f"/api/accounts/{credit}/repay", json={"from_account_id": c1, "amount": 100, "interest": 101}).status_code == 422
    assert client.post(f"/api/accounts/{c1}/repay", json={"from_account_id": credit, "amount": 100}).status_code == 422
    assert client.post(f"/api/accounts/{credit}/repay", json={"from_account_id": c1, "amount": 100, "funding": "reserve"}).status_code == 422
    # A plain transfer to a credit card is not a repayment.
    assert client.post("/api/accounts/transfers", json={"from_account_id": c1, "to_account_id": credit, "amount": 1}).status_code == 422


def test_planned_repayment_is_a_bill_the_buffer_puts_aside(client, monkeypatch):
    snap = setup(client)
    credit = by_title(snap, "Кредитка")["id"]
    ok(client.post("/api/transactions", json={"type": "expense", "amount": 3000, "account_id": credit}))
    snap = ok(client.post(f"/api/accounts/{credit}/schedule", json={}))
    card = by_title(snap, "Кредитка")
    assert card["scheduled"] == {"bill_id": card["scheduled"]["bill_id"], "due_date": "2026-09-25",
                                 "amount": 3000.0, "outdated": False}
    assert card["due_date"] == "2026-09-25" and snap["credit_due_date"] == "2026-09-25"
    bills = client.get("/api/bootstrap").json()["bill_rules"]
    assert [(b["title"], b["once_date"]) for b in bills] == [("Кредитка Кредитка", "2026-09-25")]
    # A one-off payment: only in September.
    with user_scope(1):
        from app import planning
        events = planning.bill_events(1, date(2026, 9, 1), date(2026, 12, 31))
    assert [e["due_date"] for e in events] == ["2026-09-25"]
    # The debt grows: the plan is out of date and can be updated in place.
    ok(client.post("/api/transactions", json={"type": "expense", "amount": 500, "account_id": credit}))
    assert by_title(ok(client.get("/api/accounts")), "Кредитка")["scheduled"]["outdated"] is True
    snap = ok(client.post(f"/api/accounts/{credit}/schedule", json={}))
    assert by_title(snap, "Кредитка")["scheduled"]["amount"] == 3500
    assert len(client.get("/api/bootstrap").json()["bill_rules"]) == 1
    # Paid on the day: the buffer pays, the debt goes, no expense in analytics.
    monkeypatch.setattr(clock, "today", lambda: date(2026, 9, 25))
    budget = ok(client.get("/api/cashflow"))["available_today"]
    bill = next(b for b in client.get("/api/plan").json() if b["title"].startswith("Кредитка"))
    payment = ok(client.post(f"/api/bills/{bill['id']}/pay", json={"due_date": bill["due_date"]}))
    assert payment["account_id"] is None
    snap, cards = state(client)
    assert cards["Кредитка"]["debt"] == 0 and cards["Кредитка"]["scheduled"] is None
    assert ok(client.get("/api/cashflow"))["available_today"] == budget
    _, data = categories_spent(client)
    assert data["totals"]["bills"] == 0
    assert all(not b["title"].startswith("Кредитка") for b in data["bills"])
    item = next(i for i in ok(client.get(f"/api/accounts/{credit}/history"))["items"] if i["type"] == "transfer")
    assert item["deletable"] is False
    assert client.delete(f"/api/accounts/transfers/{item['id']}").status_code == 422
    # Unmarking the payment brings the debt back.
    ok(client.delete(f"/api/transactions/{payment['id']}"))
    assert state(client)[1]["Кредитка"]["debt"] == 3500


def test_credit_card_with_a_debt_stays_until_repaid(client):
    snap = setup(client)
    credit = by_title(snap, "Кредитка")["id"]
    ok(client.post("/api/transactions", json={"type": "expense", "amount": 10, "account_id": credit}))
    response = client.post(f"/api/accounts/{credit}/archive")
    assert response.status_code == 422 and "долг" in response.json()["detail"]
    # A credit card holds no envelope.
    assert client.put("/api/accounts/envelopes/spending", json={"account_id": credit}).status_code == 422


def test_moving_a_purchase_off_the_credit_card_returns_it_to_the_budget(client):
    snap = setup(client)
    c1, credit = by_title(snap, "Карта 1")["id"], by_title(snap, "Кредитка")["id"]
    budget = ok(client.get("/api/cashflow"))["available_today"]
    tx = ok(client.post("/api/transactions", json={"type": "expense", "amount": 400, "account_id": credit}))
    ok(client.put(f"/api/transactions/{tx['id']}", json={"type": "expense", "amount": 400, "account_id": c1}))
    assert ok(client.get("/api/cashflow"))["available_today"] == budget - 400
    ok(client.put(f"/api/transactions/{tx['id']}", json={"type": "expense", "amount": 400, "account_id": credit}))
    assert ok(client.get("/api/cashflow"))["available_today"] == budget


def test_irregular_repayment_from_free_money_or_the_reserve(client):
    irregular(client)
    c1, c2 = two_cards(client, {"reserve": 1})
    credit = by_title(ok(client.get("/api/accounts")), "Кредитка")["id"]
    ok(client.post("/api/transactions", json={"type": "expense", "amount": 1000, "account_id": credit}))
    flow = ok(client.get("/api/irregular"))
    ok(client.post(f"/api/accounts/{credit}/repay", json={"from_account_id": c1, "amount": 400}))
    assert ok(client.get("/api/irregular"))["free_balance"] == flow["free_balance"] - 400
    ok(client.post(f"/api/accounts/{credit}/repay", json={"from_account_id": c2, "amount": 600, "funding": "reserve"}))
    after = ok(client.get("/api/irregular"))
    assert after["reserve_balance"] == flow["reserve_balance"] - 600
    assert differences(client) == {"Карта 1": 0, "Карта 2": 0}
    assert state(client)[1]["Кредитка"]["debt"] == 0


# ---------- Stage 5: backup v11 ----------

from app.backup import export_user_data, restore_user_data  # noqa: E402

ACCOUNT_TABLES = ("accounts", "envelope_accounts", "account_transfers", "account_adjustments", "account_incomes")


def rich_history(client):
    """Every kind of accounts record: transfers, linked ones, corrections,
    incomes, a credit card repayment with interest, a planned repayment."""
    c1, c2 = two_cards(client, {"piggy": 1})
    credit = by_title(ok(client.get("/api/accounts")), "Кредитка")["id"]
    ok(client.post("/api/piggy-bank/deposit", json={"amount": 50, "source": "daily_budget"}))
    ok(client.post("/api/transactions", json={"type": "expense", "amount": 300, "account_id": c1, "category_id": 1}))
    ok(client.post("/api/transactions", json={"type": "expense", "amount": 200, "account_id": c2,
                                              "funding_envelope": "piggy"}))
    ok(client.post("/api/transactions", json={"type": "expense", "amount": 3000, "account_id": credit}))
    ok(client.post("/api/piggy-bank/deposit", json={"amount": 70}))
    ok(client.post("/api/accounts/transfers", json={"from_account_id": c1, "to_account_id": c2, "amount": 10}))
    ok(client.post(f"/api/accounts/{c2}/adjust", json={"target": 1}))
    ok(client.post(f"/api/accounts/{credit}/repay", json={"from_account_id": c1, "amount": 1000, "interest": 100}))
    ok(client.post(f"/api/accounts/{credit}/adjust", json={"target": 2500}))
    ok(client.post(f"/api/accounts/{credit}/schedule", json={}))
    return c1, c2, credit


def comparable(snap):
    names = {a["id"]: a["title"] for a in snap["accounts"]}
    accounts = [{k: v for k, v in a.items() if k not in ("id", "created_at", "scheduled")} for a in snap["accounts"]]
    hints = [{**h, "from_account_id": names.get(h.get("from_account_id")), "to_account_id": names.get(h.get("to_account_id"))}
             for h in snap["hints"]]
    envelopes = [{**e, "account_id": names.get(e["account_id"])} for e in snap["envelopes"]]
    return {"accounts": accounts, "hints": hints, "envelopes": envelopes, "total": snap["total"],
            "credit_debt": snap["credit_debt"], "enabled": snap["enabled"],
            "main": names.get(snap["main_account_id"])}


def test_backup_v11_round_trip_keeps_cards_and_their_history(client):
    rich_history(client)
    before = ok(client.get("/api/accounts"))
    budget = ok(client.get("/api/cashflow"))["available_today"]
    backup = export_user_data(1)
    assert backup["backup_version"] == 11
    assert all(backup["data"][t] for t in ACCOUNT_TABLES)
    ensure_user(2)
    restore_user_data(2, backup)
    with user_scope(2):
        from app import accounts_engine, engine
        after = accounts_engine.snapshot(2)
        assert engine.compute(2)["available_today"] == budget
    assert comparable(after) == comparable(before)
    # And back over user 1 itself: the same picture.
    restore_user_data(1, export_user_data(1))
    assert comparable(ok(client.get("/api/accounts"))) == comparable(before)
    # Linked records keep working after the ids were remapped.
    interest = next(t for t in ok(client.get("/api/transactions")) if t["category_title"] == "Проценты банка")
    assert client.delete(f"/api/transactions/{interest['id']}").status_code == 422


@pytest.mark.parametrize("version", range(1, 11))
def test_older_backups_restore_with_the_feature_off(client, version):
    rich_history(client)
    backup = export_user_data(1)
    old = {**backup, "backup_version": version, "data": {k: v for k, v in backup["data"].items() if k not in ACCOUNT_TABLES}}
    old["data"]["settings"] = {**old["data"]["settings"], "accounts_enabled": 1}
    for row in old["data"]["transactions"]:
        row.pop("account_id", None)
        row.pop("funding_envelope", None)
    for row in old["data"]["bill_rules"]:
        row.pop("credit_account_id", None)
    old["data"]["emergency_reserve_movements"] = []
    restore_user_data(1, old)
    snap = ok(client.get("/api/accounts"))
    assert snap["enabled"] is False and snap["accounts"] == [] and snap["main_account_id"] is None
    ok(client.get("/api/cashflow"))


def test_backup_with_broken_account_links_is_refused(client):
    rich_history(client)
    backup = export_user_data(1)
    broken = json_copy(backup)
    broken["data"]["transactions"][0]["account_id"] = 999
    with pytest.raises(ValueError):
        restore_user_data(1, broken)
    broken = json_copy(backup)
    credit = next(a["id"] for a in broken["data"]["accounts"] if a["kind"] == "credit")
    broken["data"]["envelope_accounts"][0]["account_id"] = credit
    with pytest.raises(ValueError):
        restore_user_data(1, broken)
    # Nothing was touched.
    assert len(ok(client.get("/api/accounts"))["accounts"]) == 3


def json_copy(value):
    import json
    return json.loads(json.dumps(value))


def test_backup_keeps_a_repayment_from_the_reserve(client):
    irregular(client)
    c1, c2 = two_cards(client, {"reserve": 1})
    credit = by_title(ok(client.get("/api/accounts")), "Кредитка")["id"]
    ok(client.post("/api/transactions", json={"type": "expense", "amount": 600, "account_id": credit}))
    ok(client.post(f"/api/accounts/{credit}/repay", json={"from_account_id": c2, "amount": 600, "funding": "reserve"}))
    before = ok(client.get("/api/accounts"))
    reserve = ok(client.get("/api/irregular"))["reserve_balance"]
    restore_user_data(1, export_user_data(1))
    assert comparable(ok(client.get("/api/accounts"))) == comparable(before)
    assert ok(client.get("/api/irregular"))["reserve_balance"] == reserve


def test_excel_has_cards_and_accounts(client):
    import io

    from openpyxl import load_workbook

    from app.excel_export import build_workbook

    rich_history(client)
    wb = load_workbook(io.BytesIO(build_workbook(1)))
    ops = list(wb["Операции"].iter_rows(values_only=True))
    assert ops[0][5:7] == ("Счёт", "Из чего")
    rows = {(r[2], r[3]): r for r in ops[1:]}
    assert rows[("Без категории", -200)][5:7] == ("Карта 2", "Из копилки")
    assert rows[("Без категории", -3000)][5:7] == ("Кредитка", "Вне бюджета")
    sheet = list(wb["Карты и счета"].iter_rows(values_only=True))
    flat = [v for row in sheet for v in row if v is not None]
    assert "Всего (без кредиток)" in flat and "Кредитки" in flat and "Движения" in flat
    credit = next(row for row in sheet if row[0] == "Кредитка" and isinstance(row[1], (int, float)))
    assert credit[1] == 2500 and credit[4] == 1000 and credit[5] == 500     # interest 100 + correction 400
    card1 = next(row for row in sheet if row[0] == "Карта 1" and row[1] == "Дебетовая карта")
    assert card1[2] == "Да" and "на траты" in card1[5]
    # Credit card purchases count in the months; the planned repayment is no expense.
    pivot = list(wb["Расходы по месяцам"].iter_rows(values_only=True))
    assert all(row[0] != "Обязательные платежи" for row in pivot)


# ---------- Stage 5: the bot ----------


def report_for(day):
    from datetime import datetime

    from app.morning_reports import pending_morning_reports

    return pending_morning_reports(datetime(day.year, day.month, day.day, 12, 0))[0]


def test_morning_report_has_the_total_and_the_credit_debt(client, monkeypatch):
    import bot

    assert report_for(TODAY)["accounts"] is None
    assert "Всего денег" not in bot.morning_report_text(report_for(TODAY))
    snap = setup(client)
    credit = by_title(snap, "Кредитка")["id"]
    ok(client.post("/api/transactions", json={"type": "expense", "amount": 3000, "account_id": credit}))
    report = report_for(TODAY)
    assert report["accounts"] == {"total": 155000.0, "credit_debt": 3000.0, "credit_due_date": "2026-09-25",
                                  "credit_reminders": []}
    text = bot.morning_report_text(report)
    assert "💰 Всего денег: 155 000,00 ₽" in text
    assert "💳 Долг по кредитке: 3 000,00 ₽ · до 25.09" in text
    assert "⏰" not in text
    # Three days before the payment day and on the day: a reminder.
    monkeypatch.setattr(clock, "today", lambda: date(2026, 9, 22))
    text = bot.morning_report_text(report_for(date(2026, 9, 22)))
    assert "⏰ Платёж по кредитке «Кредитка» через 3 дн., 25.09: долг 3 000,00 ₽" in text
    monkeypatch.setattr(clock, "today", lambda: date(2026, 9, 24))
    assert "⏰" not in bot.morning_report_text(report_for(date(2026, 9, 24)))
    monkeypatch.setattr(clock, "today", lambda: date(2026, 9, 25))
    assert "⏰ Платёж по кредитке «Кредитка» сегодня: долг 3 000,00 ₽" in bot.morning_report_text(report_for(date(2026, 9, 25)))


def test_irregular_morning_report_has_the_total_too(client):
    import bot

    irregular(client)
    two_cards(client, {})
    report = report_for(TODAY)
    assert report["mode"] == "irregular" and report["accounts"]["credit_debt"] == 0
    assert "💰 Всего денег:" in bot.morning_report_text(report)


def test_spending_today_lists_credit_piggy_and_outside_apart(client):
    import bot
    from app.quick_stats import quick_numbers

    c1, c2 = two_cards(client, {"piggy": 1})
    credit = by_title(ok(client.get("/api/accounts")), "Кредитка")["id"]
    ok(client.post("/api/transactions", json={"type": "expense", "amount": 30, "account_id": c1, "note": "Кофе"}))
    ok(client.post("/api/transactions", json={"type": "expense", "amount": 3000, "account_id": credit, "note": "Пылесос"}))
    ok(client.post("/api/transactions", json={"type": "expense", "amount": 200, "account_id": c2,
                                              "funding_envelope": "piggy", "note": "Подарок"}))
    ok(client.post("/api/transactions", json={"type": "expense", "amount": 70, "account_id": c1,
                                              "funding_envelope": "none", "note": "Чужие деньги"}))
    numbers = quick_numbers(1)
    assert numbers["spent_today"] == 30
    assert [i["title"] for i in numbers["expenses"]] == ["Кофе"]
    text = bot.spent_text(numbers)
    assert "Кредиткой (растёт долг, не из суммы на день):\n• Пылесос — 3 000,00 ₽" in text
    assert "Из копилки:\n• Подарок — 200,00 ₽" in text
    assert "Вне бюджета:\n• Чужие деньги — 70,00 ₽" in text
