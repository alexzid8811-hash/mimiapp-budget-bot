"""Irregular-income mode through the HTTP API and the database."""
import sqlite3
from datetime import date

import pytest
from fastapi.testclient import TestClient

from app import clock, engine
from app.cashflow_app import app
from app.db import connect, ensure_user, init_db, user_db_path
from app.irregular_engine import snapshot

NOW = {"today": date(2026, 10, 1)}


@pytest.fixture
def client(tmp_path, monkeypatch):
    """Start 1 October with no money; internet 1 000 ₽ on the 10th and
    utilities 6 000 ₽ on the 25th; reserve 10 %, N = 14, K = 30."""
    NOW["today"] = date(2026, 10, 1)
    monkeypatch.setenv("DATABASE_PATH", str(tmp_path / "budget.sqlite3"))
    monkeypatch.setenv("DEV_MODE", "true")
    monkeypatch.setattr(clock, "today", lambda: NOW["today"])
    monkeypatch.setattr("app.russian_calendar._remote_year", lambda _: None)
    init_db()
    ensure_user(1)
    with TestClient(app) as value:
        for title, amount, day in (("Интернет", 1000, 10), ("Коммуналка", 6000, 25)):
            r = value.post("/api/bill-rules", json={
                "title": title, "amount": amount, "day_of_month": day, "effective_date": "2026-10-01"})
            assert r.status_code == 200, r.text
        r = value.put("/api/budget-mode", json={
            "budget_mode": "irregular", "start_date": "2026-10-01", "start_total": 0, "start_reserve": 0})
        assert r.status_code == 200, r.text
        yield value


def at(day):
    NOW["today"] = day


def income(client, amount, day, **extra):
    r = client.post("/api/irregular/incomes", json={"amount": amount, "tx_date": day, **extra})
    assert r.status_code == 200, r.text
    return r.json()


def spend(client, amount, day):
    r = client.post("/api/transactions", json={"type": "expense", "amount": amount, "tx_date": day,
                                               "category_id": None})
    assert r.status_code == 200, r.text
    return r.json()


def test_mode_requires_start(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_PATH", str(tmp_path / "b.sqlite3"))
    monkeypatch.setenv("DEV_MODE", "true")
    init_db()
    with TestClient(app) as c:
        assert c.put("/api/budget-mode", json={"budget_mode": "irregular"}).status_code == 422
        assert c.get("/api/irregular").json()["configured"] is False
        future = c.put("/api/budget-mode", json={"budget_mode": "irregular", "start_date": "2999-01-01"})
        assert future.status_code == 422


def test_reference_scenario_through_api(client):
    at(date(2026, 10, 3))
    preview = client.post("/api/irregular/preview", json={"amount": 20000, "tx_date": "2026-10-03"}).json()
    assert (preview["reserve"], preview["bills"], preview["free"]) == (2000, 7000, 11000)
    created = income(client, 20000, "2026-10-03", income_source="ремонт")
    assert created["summary"] == "Пришло 20 000 ₽ → 2 000 ₽ в резерв (10%), 7 000 ₽ на обязательные, 11 000 ₽ свободных"
    flow = client.get("/api/irregular").json()
    assert flow["today_target"] == 785.71
    assert flow["stretch_until"] == "2026-10-16" and flow["days_left"] == 14
    assert flow["reserve_balance"] == 2000 and flow["bills_reserved"] == 7000
    assert flow["next_bill"]["title"] == "Интернет" and flow["next_bill"]["reserved"] == 1000
    assert flow["bills_shortfall"] is None

    at(date(2026, 10, 4))
    assert client.get("/api/irregular").json()["today_target"] == 846.15
    at(date(2026, 10, 6))
    income(client, 5000, "2026-10-06")
    flow = client.get("/api/irregular").json()
    assert flow["today_target"] == 1107.14
    assert flow["stretch_until"] == "2026-10-19"
    assert flow["free_balance"] == 15500
    assert flow["reserve_balance"] == 2500
    assert flow["days_without_income"] == 0


def test_generic_income_in_irregular_mode_is_split(client):
    at(date(2026, 10, 3))
    r = client.post("/api/transactions", json={"type": "income", "amount": 20000, "tx_date": "2026-10-03"})
    assert r.status_code == 200
    assert r.json()["income_destination"] == "split" and r.json()["reserve_percent"] == 10
    assert client.get("/api/irregular").json()["reserve_balance"] == 2000


def test_percent_change_does_not_touch_past_incomes(client):
    at(date(2026, 10, 3))
    income(client, 20000, "2026-10-03")
    assert client.put("/api/irregular/settings", json={
        "reserve_percent": 50, "stretch_days": 14, "lookahead_days": 30}).status_code == 200
    flow = client.get("/api/irregular").json()
    assert flow["reserve_balance"] == 2000
    later = income(client, 1000, "2026-10-03")
    assert later["split"]["reserve"] == 500
    # Percent can be overridden for one income.
    once = income(client, 1000, "2026-10-03", reserve_percent=0)
    assert once["split"]["reserve"] == 0


def test_deleting_income_rolls_back_or_is_rejected(client):
    at(date(2026, 10, 3))
    first = income(client, 20000, "2026-10-03")
    second = income(client, 10000, "2026-10-03")
    assert client.get("/api/irregular").json()["reserve_balance"] == 3000
    assert client.delete(f"/api/transactions/{second['id']}").status_code == 200
    flow = client.get("/api/irregular").json()
    assert flow["reserve_balance"] == 2000 and flow["free_balance"] == 11000
    r = client.post("/api/irregular/reserve/withdraw", json={
        "amount": 1500, "purpose": "to_free", "reason": "Мало заказов"})
    assert r.status_code == 200, r.text
    rejected = client.delete(f"/api/transactions/{first['id']}")
    assert rejected.status_code == 422
    assert "Резерв" in rejected.json()["detail"]
    assert client.get("/api/irregular").json()["reserve_balance"] == 500
    # Editing the income down is also checked against the whole history.
    assert client.put(f"/api/irregular/incomes/{first['id']}", json={
        "amount": 1000, "tx_date": "2026-10-03"}).status_code == 422


def test_backdated_income_recalculates_later_days(client):
    at(date(2026, 10, 8))
    income(client, 20000, "2026-10-03")
    income(client, 1000, "2026-10-07")
    before = client.get("/api/irregular").json()
    income(client, 3000, "2026-10-05")
    after = client.get("/api/irregular").json()
    assert after["free_balance"] == before["free_balance"] + 2700
    assert after["reserve_balance"] == before["reserve_balance"] + 300
    assert client.post("/api/irregular/incomes", json={"amount": 1, "tx_date": "2026-09-30"}).status_code == 422
    assert client.post("/api/irregular/incomes", json={"amount": 1, "tx_date": "2026-10-09"}).status_code == 422


def test_three_ways_to_take_from_reserve(client):
    at(date(2026, 10, 3))
    income(client, 20000, "2026-10-03")
    category = client.get("/api/bootstrap").json()["categories"][4]["id"]
    assert client.post("/api/irregular/reserve/withdraw", json={
        "amount": 100, "purpose": "to_free", "reason": ""}).status_code == 422  # reason is required
    paid = client.post("/api/irregular/reserve/withdraw", json={
        "amount": 500, "purpose": "pay_expense", "reason": "Стоматолог", "category_id": category})
    assert paid.status_code == 200, paid.text
    flow = client.get("/api/irregular").json()
    assert flow["reserve_balance"] == 1500
    assert flow["today_target"] == 785.71 and flow["spent_today"] == 0
    expense = next(t for t in client.get("/api/transactions").json() if t["type"] == "expense")
    assert expense["note"] == "Стоматолог" and expense["category_id"] == category
    # Editing the expense moves its reserve withdrawal along.
    assert client.put(f"/api/transactions/{expense['id']}", json={
        "type": "expense", "amount": 700, "tx_date": "2026-10-03", "category_id": category,
        "note": "Стоматолог"}).status_code == 200
    assert client.get("/api/irregular").json()["reserve_balance"] == 1300
    movement = client.get("/api/irregular/reserve").json()["movements"][0]
    assert client.delete(f"/api/irregular/reserve/movements/{movement['id']}").status_code == 422

    spend(client, 1000, "2026-10-03")
    flow = client.get("/api/irregular").json()
    assert flow["overspend"] == 214.29
    assert client.post("/api/irregular/reserve/withdraw", json={
        "amount": 300, "purpose": "cover_overspend", "reason": "Сломался телефон"}).status_code == 422
    assert client.post("/api/irregular/reserve/withdraw", json={
        "amount": 214.29, "purpose": "cover_overspend", "reason": "Сломался телефон"}).status_code == 200
    flow = client.get("/api/irregular").json()
    assert flow["overspend"] == 0 and flow["available_today"] == 0

    before = client.get("/api/irregular").json()["tomorrow_limit"]
    assert client.post("/api/irregular/reserve/withdraw", json={
        "amount": 1000, "purpose": "to_free", "reason": "Нет заказов"}).status_code == 200
    flow = client.get("/api/irregular").json()
    assert flow["tomorrow_limit"] > before
    assert client.post("/api/irregular/reserve/withdraw", json={
        "amount": 1000, "purpose": "to_free", "reason": "Слишком много"}).status_code == 422

    history = client.get("/api/irregular/reserve").json()["movements"]
    reasons = [m["reason"] for m in history]
    assert {"Стоматолог", "Сломался телефон", "Нет заказов"} <= set(reasons)
    assert history[0]["balance"] == flow["reserve_balance"]


def test_reserve_deposits(client):
    at(date(2026, 10, 3))
    income(client, 20000, "2026-10-03")
    assert client.post("/api/irregular/reserve/deposit", json={
        "amount": 1000, "kind": "deposit", "reason": "Подарок", "movement_date": "2026-10-02"}).status_code == 200
    assert client.post("/api/irregular/reserve/deposit", json={
        "amount": 11000.01, "kind": "from_free"}).status_code == 422
    assert client.post("/api/irregular/reserve/deposit", json={"amount": 1000, "kind": "from_free"}).status_code == 200
    flow = client.get("/api/irregular").json()
    assert flow["reserve_balance"] == 4000 and flow["free_balance"] == 10000


def test_reserve_target_through_settings(client):
    at(date(2026, 10, 3))
    assert client.put("/api/irregular/settings", json={
        "reserve_percent": 10, "reserve_target": 2500, "stretch_days": 14, "lookahead_days": 30}).status_code == 200
    income(client, 20000, "2026-10-03")
    assert income(client, 10000, "2026-10-03")["split"]["reserve"] == 500
    assert income(client, 10000, "2026-10-03")["split"]["reserve"] == 0


def test_bill_paid_more_and_less_than_put_aside(client):
    at(date(2026, 10, 3))
    income(client, 20000, "2026-10-03")
    plan = client.get("/api/plan").json()
    internet = next(b for b in plan if b["title"] == "Интернет")
    assert internet["reserved"] == 1000
    paid = client.post(f"/api/bills/{internet['id']}/pay", json={"due_date": "2026-10-10"}).json()
    flow = client.get("/api/irregular").json()
    assert flow["bills_reserved"] == 6000 and flow["free_balance"] == 11000
    # Paid more: the difference comes from free money.
    assert client.put(f"/api/bill-payments/{paid['id']}", json={"amount": 1300}).status_code == 200
    flow = client.get("/api/irregular").json()
    assert flow["free_balance"] == 10700 and flow["spent_today"] == 300
    # Paid less: the remainder returns to free money ...
    assert client.put(f"/api/bill-payments/{paid['id']}", json={"amount": 800}).status_code == 200
    assert client.get("/api/irregular").json()["free_balance"] == 11200
    # ... or goes to the piggy bank.
    assert client.put(f"/api/bill-payments/{paid['id']}", json={
        "amount": 800, "remainder_destination": "piggy"}).status_code == 200
    flow = client.get("/api/irregular").json()
    assert flow["free_balance"] == 11000 and flow["piggy_bank_balance"] == 200


def test_small_income_warns_about_bills(client):
    at(date(2026, 10, 3))
    income(client, 5000, "2026-10-03")
    flow = client.get("/api/irregular").json()
    assert flow["bills_shortfall"] == {"amount": 2500, "date": "2026-10-25"}
    assert flow["free_balance"] == 0 and flow["today_target"] == 0


def test_window_extends_without_income(client):
    at(date(2026, 10, 3))
    income(client, 20000, "2026-10-03")
    at(date(2026, 10, 20))
    flow = client.get("/api/irregular").json()
    assert flow["no_income_warning"] is True
    assert flow["days_without_income"] == 17
    assert flow["stretch_until"] == "2026-10-30"


def test_reserve_days_need_spending_history(client):
    at(date(2026, 10, 3))
    income(client, 20000, "2026-10-03")
    assert client.get("/api/irregular").json()["reserve_days"] is None
    spend(client, 500, "2026-10-03")
    flow = client.get("/api/irregular").json()
    assert flow["avg_daily_spending"] == 500 and flow["reserve_days"] == 4


def test_switching_mode_keeps_data_and_payroll_numbers(client):
    at(date(2026, 10, 3))
    with connect() as con:
        con.execute("UPDATE income_rules SET amount=50000 WHERE user_id=1")
        con.execute("UPDATE settings SET cashflow_enabled=1,cashflow_start_date='2026-10-01',"
                    "cashflow_start_capital=10000 WHERE user_id=1")
    assert client.put("/api/budget-mode", json={"budget_mode": "payroll"}).status_code == 200
    before = engine.compute(1)
    assert client.put("/api/budget-mode", json={"budget_mode": "irregular"}).status_code == 200
    client.post("/api/irregular/reserve/deposit", json={"amount": 300, "kind": "deposit", "reason": "x"})
    assert client.put("/api/irregular/settings", json={
        "reserve_percent": 33, "stretch_days": 20, "lookahead_days": 10}).status_code == 200
    assert client.put("/api/budget-mode", json={"budget_mode": "payroll"}).status_code == 200
    assert engine.compute(1) == before
    # Irregular data survived the switch.
    assert client.put("/api/budget-mode", json={"budget_mode": "irregular"}).status_code == 200
    assert snapshot(1)["reserve_balance"] == 300
    assert client.get("/api/bootstrap").json()["settings"]["budget_mode"] == "irregular"


def test_payroll_mode_rejects_irregular_destinations(client):
    assert client.put("/api/budget-mode", json={"budget_mode": "payroll"}).status_code == 200
    r = client.post("/api/transactions", json={"type": "income", "amount": 10, "income_destination": "split"})
    assert r.status_code == 422


def test_existing_user_file_is_migrated(tmp_path, monkeypatch):
    from app.db import SCHEMA

    monkeypatch.setenv("DATABASE_PATH", str(tmp_path / "budget.sqlite3"))
    path = user_db_path(7)
    with sqlite3.connect(path) as con:
        con.executescript(SCHEMA)
        con.execute("INSERT INTO users(id,first_name) VALUES(7,'Old')")
        con.execute("INSERT INTO settings(user_id,salary_gross) VALUES(7,99000)")
        con.execute("INSERT INTO transactions(user_id,type,amount,tx_date,note) VALUES(7,'income',123,'2026-01-01','старый')")
    con.close()
    init_db()
    with connect(7) as con:
        settings = {r[1] for r in con.execute("PRAGMA table_info(settings)")}
        transactions = {r[1] for r in con.execute("PRAGMA table_info(transactions)")}
        tables = {r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        row = con.execute("SELECT salary_gross,budget_mode,irregular_reserve_percent,irregular_stretch_days,"
                          "irregular_bills_lookahead_days FROM settings WHERE user_id=7").fetchone()
        tx = con.execute("SELECT amount,note,income_source,reserve_percent FROM transactions WHERE user_id=7").fetchone()
    assert {"budget_mode", "irregular_reserve_percent", "irregular_start_date"} <= settings
    assert {"reserve_percent", "income_source"} <= transactions
    assert "emergency_reserve_movements" in tables
    assert tuple(row) == (99000, "payroll", 10, 14, 30)
    assert tuple(tx) == (123, "старый", "", None)


def test_income_analytics(client):
    at(date(2026, 10, 3))
    income(client, 20000, "2026-10-03", income_source="ремонт")
    at(date(2026, 10, 20))
    income(client, 5000, "2026-10-06", income_source="Доставка")
    income(client, 3000, "2026-10-08", income_source="Ремонт")
    income(client, 1000, "2026-10-08")
    client.post("/api/irregular/reserve/withdraw", json={"amount": 700, "purpose": "to_free", "reason": "Дождь"})
    at(date(2026, 12, 15))
    income(client, 10000, "2026-11-20", income_source="ремонт")
    october = client.get("/api/analytics/income?mode=month&anchor=2026-10-01").json()
    assert october["total"] == 29000 and october["count"] == 4
    assert october["sources"][0] == {"title": "ремонт", "amount": 23000, "count": 2, "share": round(23000 / 29000, 4)}
    assert {s["title"] for s in october["sources"]} == {"ремонт", "Доставка", "Без источника"}
    assert sum(d["amount"] for d in october["days"]) == 29000
    assert sum(w["amount"] for w in october["weeks"]) == 29000
    assert october["reserve"]["put_aside"] == 2900
    assert october["reserve"]["taken"] == 700
    assert october["reserve"]["taken_items"][0]["reason"] == "Дождь"
    # Complete months: October and November (start in October).
    assert october["stats"]["3"] == {"months": 2, "average": 19500, "median": 19500}
    assert october["median_interval_days"] == 3   # 3→6, 6→8, 8→20.11
    assert october["longest_gap"] == {"days": 43, "from": "2026-10-08", "to": "2026-11-20"}


def _numbers():
    flow = snapshot(1)
    keys = ("today_target", "available_today", "free_balance", "stretch_until", "reserve_balance",
            "bills_reserved", "bills_shortfall", "days_without_income", "piggy_bank_balance")
    return {k: flow[k] for k in keys}


def _rich_history(client):
    at(date(2026, 10, 3))
    income(client, 20000, "2026-10-03", income_source="ремонт")
    category = client.get("/api/bootstrap").json()["categories"][0]["id"]
    client.post("/api/irregular/reserve/withdraw", json={
        "amount": 300, "purpose": "pay_expense", "reason": "Лекарства", "category_id": category})
    at(date(2026, 10, 6))
    income(client, 5000, "2026-10-06", destination="piggy")
    income(client, 700, "2026-10-06", destination="reserve")
    spend(client, 1200, "2026-10-06")
    client.post("/api/irregular/reserve/deposit", json={"amount": 100, "kind": "from_free"})
    assert client.put("/api/irregular/settings", json={
        "reserve_percent": 15, "reserve_target": 50000, "stretch_days": 10, "lookahead_days": 20}).status_code == 200


def test_backup_round_trip_keeps_numbers(client):
    from app.backup import export_user_data, restore_user_data

    _rich_history(client)
    before = _numbers()
    backup = export_user_data(1)
    assert backup["backup_version"] == 10
    assert len(backup["data"]["emergency_reserve_movements"]) == 2
    with connect() as con:
        con.execute("DELETE FROM emergency_reserve_movements WHERE user_id=1")
        con.execute("DELETE FROM transactions WHERE user_id=1")
        con.execute("UPDATE settings SET budget_mode='payroll',irregular_start_date=NULL WHERE user_id=1")
    restore_user_data(1, backup)
    assert _numbers() == before
    assert snapshot(1)["settings"]["reserve_target"] == 50000
    # The expense paid from the reserve is still linked to it.
    assert client.get("/api/irregular/reserve").json()["movements"][-2]["expense_id"] is not None


def test_version_nine_backup_is_restored(client):
    from app.backup import export_user_data, restore_user_data

    at(date(2026, 10, 3))
    backup = export_user_data(1)
    backup["backup_version"] = 9
    for key in [k for k in backup["data"]["settings"] if k.startswith("irregular_") or k == "budget_mode"]:
        del backup["data"]["settings"][key]
    del backup["data"]["emergency_reserve_movements"]
    for row in backup["data"]["transactions"]:
        row.pop("reserve_percent", None)
        row.pop("income_source", None)
    counts = restore_user_data(1, backup)
    assert counts["emergency_reserve_movements"] == 0
    assert snapshot(1)["settings"]["budget_mode"] == "payroll"


def test_backup_with_broken_reserve_link_is_rejected(client):
    from app.backup import export_user_data, restore_user_data

    _rich_history(client)
    backup = export_user_data(1)
    backup["data"]["emergency_reserve_movements"][0]["expense_transaction_id"] = 999
    with pytest.raises(ValueError):
        restore_user_data(1, backup)
    assert snapshot(1)["reserve_balance"] > 0


def test_excel_has_income_and_reserve_sheets(client):
    import io

    from openpyxl import load_workbook

    from app.excel_export import build_workbook

    _rich_history(client)
    wb = load_workbook(io.BytesIO(build_workbook(1)))
    assert {"Доходы", "Резерв на непредвиденное"} <= set(wb.sheetnames)
    rows = list(wb["Доходы"].iter_rows(min_row=2, values_only=True))
    first = rows[0]
    # K was changed to 20 days: only the internet bill of 10 October is in the
    # window of the 3 October income now (K and N are not frozen per income).
    assert (first[1], first[2], first[5], first[6], first[7]) == (20000, "ремонт", 2000, 1000, 17000)
    reserve = list(wb["Резерв на непредвиденное"].iter_rows(min_row=2, values_only=True))
    assert [r[2] for r in reserve] == [2000, -300, 700, 100]
    assert reserve[-1][3] == 2500
    assert reserve[1][4] == "Лекарства"


def test_morning_report_in_irregular_mode(client):
    from datetime import datetime

    from app.morning_reports import pending_morning_reports
    from bot import morning_report_text

    at(date(2026, 10, 3))
    income(client, 20000, "2026-10-03")
    at(date(2026, 10, 20))
    report = pending_morning_reports(datetime(2026, 10, 20, 10, 0))[0]
    assert report["mode"] == "irregular"
    assert report["no_income_warning"] is True
    text = morning_report_text(report)
    assert "Можно сегодня" in text and "растягиваем до 30.10" in text
    assert "Резерв на непредвиденное: 2 000,00 ₽" in text
    # The internet bill of 10 October is still unpaid: it is the nearest one.
    assert "Ближайший обязательный платёж: Интернет — 1 000,00 ₽, 10.10 (отложено 1 000,00 ₽)" in text
    assert "На обязательные не хватает 1 000,00 ₽ к 10.11" in text
    assert "Дохода не было 17 дн." in text


def test_deleting_expense_paid_from_reserve_returns_the_money(client):
    at(date(2026, 10, 3))
    income(client, 20000, "2026-10-03")
    category = client.get("/api/bootstrap").json()["categories"][0]["id"]
    client.post("/api/irregular/reserve/withdraw", json={
        "amount": 800, "purpose": "pay_expense", "reason": "Ветеринар", "category_id": category})
    assert client.get("/api/irregular").json()["reserve_balance"] == 1200
    expense = next(t for t in client.get("/api/transactions").json() if t["type"] == "expense")
    assert client.delete(f"/api/transactions/{expense['id']}").status_code == 200
    assert client.get("/api/irregular").json()["reserve_balance"] == 2000
    assert client.get("/api/irregular/reserve").json()["movements"][0]["kind"] == "income"
