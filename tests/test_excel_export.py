import io

from openpyxl import load_workbook

from app.db import connect, ensure_user, init_db
from app.excel_export import build_workbook


def test_excel_export_contains_operations_pivot_and_piggy(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_PATH", str(tmp_path / "budget.sqlite3"))
    init_db()
    ensure_user(1, "Alex")
    with connect() as con:
        category_id = con.execute(
            "INSERT INTO categories(user_id,title,emoji) VALUES(1,'Тестовая кофейня','☕')"
        ).lastrowid
        con.executemany(
            "INSERT INTO transactions(user_id,type,amount,tx_date,category_id,note) VALUES(1,?,?,?,?,?)",
            [
                ("expense", 300, "2026-08-10", category_id, "кофе"),
                ("expense", 700, "2026-09-02", category_id, ""),
                ("income", 5000, "2026-09-03", None, "подарок"),
            ],
        )
        con.execute(
            "INSERT INTO piggy_bank_movements(user_id,direction,amount,movement_date,note) VALUES(1,'deposit',1000,'2026-09-01','')"
        )
        con.execute(
            "INSERT INTO piggy_bank_movements(user_id,direction,amount,movement_date,note) VALUES(1,'withdraw',400,'2026-09-05','')"
        )

    wb = load_workbook(io.BytesIO(build_workbook(1)))
    assert {"Операции", "Расходы по месяцам", "Копилка", "Обязательные платежи", "Доходы"} <= set(wb.sheetnames)

    ops = list(wb["Операции"].iter_rows(min_row=2, values_only=True))
    assert [(r[1], r[2], r[3]) for r in ops] == [
        ("Расход", "Тестовая кофейня", -300), ("Расход", "Тестовая кофейня", -700), ("Доход", "Дневной бюджет", 5000),
    ]
    assert ops[0][0].isoformat()[:10] == "2026-08-10"

    pivot = list(wb["Расходы по месяцам"].iter_rows(values_only=True))
    assert pivot[0] == ("Категория", "08.2026", "09.2026", "Всего")
    assert pivot[1] == ("Тестовая кофейня", 300, 700, 1000)
    assert pivot[-1] == ("Итого", 300, 700, 1000)

    piggy = list(wb["Копилка"].iter_rows(min_row=2, values_only=True))
    assert [(r[1], r[2], r[3]) for r in piggy] == [("Пополнение", 1000, 1000), ("Снятие", -400, 600)]
