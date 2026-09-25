from __future__ import annotations

import io
import os
from collections import defaultdict
from datetime import datetime, timezone
from typing import Any

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from telegram import Bot, InputFile
from telegram.error import TelegramError

from .auth import TelegramUser, current_user
from .db import connect, ensure_user


router = APIRouter(prefix="/api/export", tags=["export"])
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4F6BED")
TOTAL_FONT = Font(bold=True)
MONEY_FORMAT = "#,##0.00"
DATE_FORMAT = "DD.MM.YYYY"

INCOME_DESTINATIONS = {"daily": "Дневной бюджет", "buffer": "Буфер", "piggy": "Копилка"}
INCOME_KINDS = {"salary": "Зарплата", "advance": "Аванс", "other": "Другой доход"}


def _date(value: Any):
    if not value:
        return None
    try:
        return datetime.fromisoformat(str(value)[:10]).date()
    except ValueError:
        return str(value)


def _piggy_title(row: dict) -> str:
    if row["direction"] == "deposit":
        if row.get("bill_payment_id"):
            return "Остаток обязательного платежа"
        if row.get("income_transaction_id"):
            return "Доход в копилку"
        return "Из остатка дня" if row["source"] == "daily_budget" else "Пополнение"
    if row["source"] == "daily_budget":
        return "На карту · по дням периода" if row.get("purpose") == "transfer" else "На карту · на сегодня"
    return "Снятие"


def _write_table(
    ws: Worksheet,
    headers: list[str],
    rows: list[list[Any]],
    money_columns: set[int] = frozenset(),
    date_columns: set[int] = frozenset(),
    widths: list[int] | None = None,
) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    for index in range(1, len(headers) + 1):
        letter = get_column_letter(index)
        if index in money_columns:
            for cell in ws[letter][1:]:
                cell.number_format = MONEY_FORMAT
        if index in date_columns:
            for cell in ws[letter][1:]:
                cell.number_format = DATE_FORMAT
        width = widths[index - 1] if widths else max(12, len(headers[index - 1]) + 2)
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions


def _query(con: Any, sql: str, params: tuple) -> list[dict]:
    return [dict(row) for row in con.execute(sql, params).fetchall()]


def build_workbook(user_id: int) -> bytes:
    with connect() as con:
        con.execute("BEGIN")
        settings = con.execute("SELECT currency FROM settings WHERE user_id=?", (user_id,)).fetchone()
        if settings is None:
            raise ValueError("Настройки пользователя не найдены")
        transactions = _query(
            con,
            "SELECT t.*, c.title AS category_title, c.emoji AS category_emoji, b.title AS bill_title "
            "FROM transactions t LEFT JOIN categories c ON c.id=t.category_id AND c.user_id=t.user_id "
            "LEFT JOIN bill_rules b ON b.id=t.bill_rule_id AND b.user_id=t.user_id "
            "WHERE t.user_id=? ORDER BY t.tx_date,t.id",
            (user_id,),
        )
        piggy = _query(
            con,
            "SELECT * FROM piggy_bank_movements WHERE user_id=? ORDER BY movement_date,id",
            (user_id,),
        )
        bills = _query(
            con,
            "SELECT b.*, c.title AS category_title FROM bill_rules b "
            "LEFT JOIN categories c ON c.id=b.category_id AND c.user_id=b.user_id "
            "WHERE b.user_id=? AND b.archived=0 ORDER BY b.day_of_month,b.id",
            (user_id,),
        )
        incomes = _query(
            con,
            "SELECT * FROM income_rules WHERE user_id=? AND archived=0 ORDER BY day_of_month,id",
            (user_id,),
        )
        vacations = _query(
            con, "SELECT * FROM vacations WHERE user_id=? ORDER BY start_date,id", (user_id,)
        )
    currency = settings["currency"] or "RUB"

    wb = Workbook()

    ws = wb.active
    ws.title = "Операции"
    op_rows = []
    for t in transactions:
        if t["bill_rule_id"]:
            kind, category = "Обязательный платёж", t["bill_title"] or t["category_title"] or ""
        elif t["type"] == "expense":
            kind, category = "Расход", t["category_title"] or "Без категории"
        else:
            kind, category = "Доход", INCOME_DESTINATIONS.get(t["income_destination"] or "daily", "")
        signed = -float(t["amount"]) if t["type"] == "expense" else float(t["amount"])
        op_rows.append([_date(t["tx_date"]), kind, category, signed, t["note"] or ""])
    _write_table(
        ws, ["Дата", "Тип", "Категория / назначение", f"Сумма, {currency}", "Комментарий"], op_rows,
        money_columns={4}, date_columns={1}, widths=[12, 22, 28, 16, 40],
    )

    ws = wb.create_sheet("Расходы по месяцам")
    totals: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    for t in transactions:
        if t["type"] != "expense":
            continue
        category = "Обязательные платежи" if t["bill_rule_id"] else (t["category_title"] or "Без категории")
        totals[category][str(t["tx_date"])[:7]] += float(t["amount"])
    months = sorted({month for by_month in totals.values() for month in by_month})
    categories = sorted(totals, key=lambda name: -sum(totals[name].values()))
    pivot_rows = [
        [name, *[round(totals[name].get(month, 0.0), 2) or None for month in months],
         round(sum(totals[name].values()), 2)]
        for name in categories
    ]
    if pivot_rows:
        pivot_rows.append([
            "Итого",
            *[round(sum(totals[name].get(month, 0.0) for name in categories), 2) for month in months],
            round(sum(sum(v.values()) for v in totals.values()), 2),
        ])
    month_labels = [f"{m[5:7]}.{m[:4]}" for m in months]
    _write_table(
        ws, ["Категория", *month_labels, "Всего"], pivot_rows,
        money_columns=set(range(2, len(months) + 3)), widths=[28, *[14] * len(months), 16],
    )
    if pivot_rows:
        for cell in ws[ws.max_row]:
            cell.font = TOTAL_FONT

    ws = wb.create_sheet("Копилка")
    balance = 0.0
    piggy_rows = []
    for m in piggy:
        signed = float(m["amount"]) if m["direction"] == "deposit" else -float(m["amount"])
        balance = round(balance + signed, 2)
        piggy_rows.append([_date(m["movement_date"]), _piggy_title(m), signed, balance, m["note"] or ""])
    _write_table(
        ws, ["Дата", "Операция", f"Сумма, {currency}", "Остаток", "Комментарий"], piggy_rows,
        money_columns={3, 4}, date_columns={1}, widths=[12, 30, 16, 16, 40],
    )

    ws = wb.create_sheet("Обязательные платежи")
    _write_table(
        ws, ["Название", "День месяца", f"Сумма, {currency}", "Категория", "Активен"],
        [[b["title"], b["day_of_month"], float(b["amount"]), b["category_title"] or "", "Да" if b["active"] else "Нет"]
         for b in bills],
        money_columns={3}, widths=[30, 14, 16, 24, 10],
    )

    ws = wb.create_sheet("Доходы")
    _write_table(
        ws, ["Название", "Вид", "День месяца", f"Сумма, {currency}", "Активен"],
        [[i["title"], INCOME_KINDS.get(i["kind"], i["kind"]), i["day_of_month"], float(i["amount"]),
          "Да" if i["active"] else "Нет"] for i in incomes],
        money_columns={4}, widths=[30, 16, 14, 16, 10],
    )

    if vacations:
        ws = wb.create_sheet("Отпуска")
        _write_table(
            ws, ["Начало", "Окончание", "Дата выплаты", f"Отпускные, {currency}", "Комментарий"],
            [[_date(v["start_date"]), _date(v["end_date"]), _date(v["payment_date"]), float(v["amount"]), v["note"] or ""]
             for v in vacations],
            money_columns={4}, date_columns={1, 2, 3}, widths=[12, 12, 14, 18, 40],
        )

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _filename() -> str:
    return f"budget-{datetime.now(timezone.utc).date().isoformat()}.xlsx"


def _workbook_for(user: TelegramUser) -> bytes:
    ensure_user(user.id, user.first_name, user.username)
    try:
        return build_workbook(user.id)
    except ValueError as exc:
        raise HTTPException(404, str(exc)) from exc


@router.get("/xlsx")
def download_xlsx(user: TelegramUser = Depends(current_user)) -> Response:
    return Response(
        _workbook_for(user),
        media_type=XLSX_MIME,
        headers={
            "Content-Disposition": f'attachment; filename="{_filename()}"',
            "Cache-Control": "no-store",
        },
    )


@router.post("/xlsx/send")
async def send_xlsx_to_chat(user: TelegramUser = Depends(current_user)) -> dict:
    token = os.getenv("BOT_TOKEN", "")
    if not token:
        raise HTTPException(503, "BOT_TOKEN не настроен на сервере")
    content = _workbook_for(user)
    filename = _filename()
    try:
        async with Bot(token=token) as bot:
            message = await bot.send_document(
                chat_id=user.id,
                document=InputFile(io.BytesIO(content), filename=filename),
                caption="Выгрузка бюджета в Excel: операции, расходы по месяцам, копилка и правила.",
            )
    except TelegramError as exc:
        raise HTTPException(502, "Не удалось отправить файл в чат с ботом") from exc
    return {"ok": True, "message_id": message.message_id, "filename": filename}
